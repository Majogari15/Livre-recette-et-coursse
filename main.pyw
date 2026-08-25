import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
import copy
import difflib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import urllib.error
import urllib.parse
import urllib.request
import uuid
import zipfile
from datetime import datetime, timedelta

# Pillow est nécessaire pour afficher les photos des recettes.
try:
    from PIL import Image, ImageTk
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False

# reportlab est nécessaire pour l'export PDF de la liste de courses.
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen import canvas as pdf_canvas
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

# openpyxl est nécessaire pour l'export Excel de la liste de courses.
try:
    from openpyxl import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

# qrcode est nécessaire pour exporter une recette sous forme de QR code.
try:
    import qrcode
    QRCODE_AVAILABLE = True
except ImportError:
    QRCODE_AVAILABLE = False

# pytesseract est nécessaire pour importer une recette depuis une photo (OCR).
# Il ne suffit pas de l'installer via pip : il nécessite aussi le programme
# Tesseract OCR installé séparément sur le système (voir le LISEZ-MOI).
try:
    import pytesseract
    PYTESSERACT_AVAILABLE = True
except ImportError:
    PYTESSERACT_AVAILABLE = False

if getattr(sys, "frozen", False):
    # Application compilée avec PyInstaller (.exe) : __file__ pointe vers un
    # dossier temporaire d'extraction interne à Windows (pas vers le dossier
    # où se trouve le .exe), donc on utilise plutôt l'emplacement réel de
    # l'exécutable pour que les fichiers de données (recettes, ingrédients,
    # photos...) soient bien lus/écrits à côté du .exe, là où l'utilisateur
    # les voit et les place.
    BASE_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DATA_FILE = os.path.join(BASE_DIR, "recipes.json")
INGREDIENTS_FILE = os.path.join(BASE_DIR, "ingredients.json")
DEFAULT_INGREDIENTS_FILE = os.path.join(BASE_DIR, "ingredients_par_defaut.json")
NUTRITION_DATA_FILE = os.path.join(BASE_DIR, "valeurs_nutritionnelles.json")
INGREDIENT_ALLERGENS_FILE = os.path.join(BASE_DIR, "ingredient_allergenes.json")
INGREDIENT_OVERRIDES_FILE = os.path.join(BASE_DIR, "ingredient_custom_data.json")
INGREDIENT_PRICES_FILE = os.path.join(BASE_DIR, "ingredient_prices.json")
IMAGES_DIR = os.path.join(BASE_DIR, "images")
WEEKLY_PLAN_FILE = os.path.join(BASE_DIR, "weekly_plan.json")
MENUS_FILE = os.path.join(BASE_DIR, "menus.json")
TRASH_FILE = os.path.join(BASE_DIR, "trash.json")
BACKUPS_DIR = os.path.join(BASE_DIR, "backups")
RECENT_VIEWS_FILE = os.path.join(BASE_DIR, "recent_views.json")
SETTINGS_FILE = os.path.join(BASE_DIR, "settings.json")

os.makedirs(IMAGES_DIR, exist_ok=True)


def load_default_ingredients():
    """Charge la liste des ~1000 ingrédients de cuisine les plus courants,
    fournie avec l'application (fichier ingredients_par_defaut.json)."""
    if os.path.exists(DEFAULT_INGREDIENTS_FILE):
        try:
            with open(DEFAULT_INGREDIENTS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list):
                    return data
        except Exception:
            return []
    return []


def normalize_oe(text):
    """Remplace les ligatures œ/Œ par 'oe'/'Oe' afin qu'un mot comme « Œuf »
    s'affiche et se classe avec les autres mots commençant par « o »."""
    return text.replace("œ", "oe").replace("Œ", "Oe")


def ingredient_sort_key(text):
    """Clé de tri qui ignore les accents (é, è, ê, à, ç...) afin qu'un mot
    comme « Échalote » se classe avec les autres mots en « e », et qui
    convertit d'abord œ/Œ en oe/Oe pour un classement cohérent avec « o »."""
    normalized = normalize_oe(text)
    decomposed = unicodedata.normalize("NFD", normalized)
    stripped = "".join(c for c in decomposed if unicodedata.category(c) != "Mn")
    return stripped.lower()


def print_file(path):
    """Tente d'envoyer un fichier directement à l'imprimante par défaut du
    système. Si ce n'est pas possible (aucune application n'ayant enregistré
    de « impression silencieuse » pour ce type de fichier — un cas fréquent
    sur Windows selon le lecteur PDF installé), ouvre le fichier dans
    l'application par défaut à la place, pour que l'utilisateur puisse
    l'imprimer d'un clic depuis celle-ci.

    Retourne "printed" (envoyé directement à l'imprimante), "opened" (ouvert
    dans l'application par défaut, à imprimer manuellement), ou None (échec
    complet)."""
    if os.name == "nt":
        try:
            os.startfile(path, "print")
            return "printed"
        except OSError:
            try:
                os.startfile(path)
                return "opened"
            except OSError:
                return None
    else:
        try:
            subprocess.run(["lp", path], check=True)
            return "printed"
        except Exception:
            try:
                subprocess.run(["xdg-open", path], check=True)
                return "opened"
            except Exception:
                return None


def report_print_result(result, temp_path, subject):
    """Affiche le message adapté selon le résultat de print_file()."""
    if result == "printed":
        messagebox.showinfo("Impression", f"Envoyé à l'imprimante : {subject}.")
    elif result == "opened":
        messagebox.showinfo(
            "Ouvert pour impression",
            f"Impossible d'envoyer directement à l'imprimante depuis l'application.\n\n"
            f"Le PDF ({subject}) a été ouvert dans votre lecteur habituel : "
            "utilisez Ctrl+P (ou le bouton Imprimer) pour l'imprimer depuis là."
        )
    else:
        messagebox.showerror(
            "Impression impossible",
            f"Impossible d'ouvrir ou d'imprimer automatiquement.\n"
            f"Le PDF a tout de même été généré ici, vous pouvez l'ouvrir et l'imprimer manuellement :\n{temp_path}"
        )


def get_temp_pdf_path(prefix):
    """Crée un chemin unique dans le dossier temporaire du système, utilisé
    pour générer un PDF juste avant de l'envoyer à l'imprimante."""
    return os.path.join(tempfile.gettempdir(), f"{prefix}_{uuid.uuid4().hex}.pdf")


# Classement des ingrédients par rayon de magasin, utilisé pour regrouper la
# liste de courses. Chaque rayon est associé à une liste de mots-clés ; le
# premier rayon dont un mot-clé apparaît dans le nom de l'ingrédient est
# retenu. L'ordre de la liste correspond à un parcours de magasin classique.
RAYON_KEYWORDS = [
    ("Fruits & Légumes", [
        "ail", "oignon", "echalote", "poireau", "carotte", "celeri", "panais",
        "navet", "betterave", "radis", "pomme de terre", "patate", "topinambour",
        "tomate", "concombre", "courgette", "aubergine", "poivron", "piment",
        "chou", "brocoli", "epinard", "blette", "oseille", "roquette", "mache",
        "laitue", "batavia", "endive", "chicoree", "cresson", "pissenlit",
        "fenouil", "artichaut", "asperge", "petit pois", "haricot vert",
        "haricot beurre", "fève", "mais", "courge", "potiron", "citrouille",
        "champignon", "girolle", "cepe", "truffe", "igname", "manioc", "gombo",
        "salsifis", "cardon", "crosne", "chayotte", "cornichon", "avocat",
        "pomme", "poire", "banane", "orange", "clementine", "mandarine",
        "pamplemousse", "pomelo", "citron", "kiwi", "fraise", "framboise",
        "myrtille", "mure", "groseille", "cassis", "cerise", "griotte",
        "abricot", "peche", "nectarine", "prune", "mirabelle", "reine-claude",
        "raisin", "melon", "pasteque", "ananas", "mangue", "papaye",
        "fruit de la passion", "litchi", "grenade", "figue", "datte", "coing",
        "kaki", "rhubarbe", "noix de coco", "kumquat", "goyave", "carambole",
        "persil", "basilic", "thym", "romarin", "origan",
        "marjolaine", "sauge", "laurier", "menthe", "ciboulette", "cerfeuil",
        "estragon", "aneth", "gingembre frais", "curcuma frais", "citronnelle",
        "germe de soja", "pousse", "algue",
    ]),
    ("Viandes & Poissons", [
        "boeuf", "bœuf", "veau", "porc", "lard", "bacon", "jambon", "saucisse",
        "chorizo", "andouille", "boudin", "saucisson", "pancetta", "agneau",
        "mouton", "poulet", "coq", "dinde", "canard", "magret", "confit de",
        "foie gras", "oie", "pintade", "caille", "pigeon", "lapin", "gibier",
        "chevreuil", "sanglier", "cheval", "steak", "viande", "kefta",
        "saumon", "truite", "cabillaud", "morue", "merlu", "colin", "lieu",
        "bar", "loup de mer", "dorade", "daurade", "thon", "espadon",
        "maquereau", "sardine", "anchois", "hareng", "sole", "turbot",
        "flétan", "raie", "rouget", "saint-pierre", "lotte", "baudroie",
        "congre", "anguille", "carpe", "brochet", "perche", "sandre",
        "tilapia", "panga", "poisson", "surimi", "caviar", "tarama",
        "crevette", "gambas", "langoustine", "homard", "langouste", "crabe",
        "tourteau", "etrille", "moule", "huitre", "huître", "palourde",
        "coque", "praire", "bulot", "bigorneau", "couteau", "petoncle",
        "coquille saint-jacques", "calamar", "encornet", "seiche", "poulpe",
        "oursin", "ormeau", "escargot", "grenouille",
    ]),
    ("Crèmerie", [
        "lait", "creme", "crème", "beurre", "margarine", "yaourt",
        "fromage", "faisselle", "petit-suisse", "mascarpone", "ricotta",
        "cottage", "emmental", "gruyere", "comte", "beaufort", "cantal",
        "reblochon", "morbier", "tomme", "camembert", "brie", "coulommiers",
        "munster", "epoisses", "maroilles", "livarot", "pont-l'eveque",
        "roquefort", "bleu", "fourme", "gorgonzola", "chevre", "chèvre",
        "crottin", "feta", "halloumi", "mozzarella", "burrata", "parmesan",
        "pecorino", "grana padano", "provolone", "gouda", "edam", "cheddar",
        "raclette", "fondue", "babeurre", "kefir", "skyr", "oeuf", "œuf",
    ]),
    ("Boulangerie & Pâtisserie", [
        "farine", "levure", "bicarbonate", "chocolat", "cacao", "praline",
        "praliné", "nougat", "caramel", "gelatine", "agar-agar", "pectine",
        "pate feuilletee", "pate brisee", "pate sablee", "pate a choux",
        "pate a pizza", "pate filo", "pate a crepes", "pate a gaufres",
        "genoise", "biscuit", "boudoir", "speculoos", "meringue",
        "poudre d'amande", "poudre de noisette", "amande effilee",
        "fruits confits", "nappage", "glacage", "fondant", "marron glace",
        "chataigne", "châtaigne", "pain", "baguette", "chapelure", "croutons",
        "biscotte", "cracker", "sucre", "cassonade", "vergeoise", "miel",
        "sirop", "melasse", "vanille",
    ]),
    ("Épicerie", [
        "riz", "semoule", "couscous", "boulgour", "quinoa", "epeautre",
        "orge", "sarrasin", "avoine", "ble", "blé", "fecule", "maizena",
        "tapioca", "pate", "spaghetti", "penne", "fusilli", "tagliatelle",
        "macaroni", "lasagne", "nouille", "vermicelle", "lentille",
        "pois chiche", "soja", "haricot rouge", "haricot blanc",
        "haricot noir", "huile", "graisse", "saindoux", "ghee",
        "moutarde", "ketchup", "mayonnaise", "vinaigre", "sauce",
        "concentre de tomate", "coulis", "pesto", "tapenade", "houmous",
        "tahini", "confiture", "marmelade", "gelee", "chutney", "pickles",
        "cornichons au vinaigre", "câpres", "capres", "olive", "raifort",
        "wasabi", "bouillon", "fond de veau", "fumet", "tofu", "seitan",
        "tempeh", "conserve", "boite", "boîte", "chips", "nachos",
        "pop-corn", "biscuit apero", "cacahuete grillee",
    ]),
    ("Herbes & Épices", [
        "poivre", "sel", "paprika", "cumin", "coriandre en", "cannelle",
        "muscade", "girofle", "cardamome", "anis", "curry", "garam masala",
        "ras el hanout", "za'atar", "sumac", "safran", "vanille en poudre",
        "reglisse", "genievre", "moutarde en poudre", "herbes de provence",
        "bouquet garni", "quatre epices", "piment de cayenne",
        "piment d'espelette", "chili en poudre", "epices", "épices",
        "curcuma en poudre", "gingembre en poudre", "sesame", "sésame",
        "graines de", "colorant",
    ]),
    ("Boissons", [
        "vin", "champagne", "porto", "madere", "marsala", "vermouth",
        "cognac", "armagnac", "calvados", "rhum", "whisky", "bourbon", "gin",
        "vodka", "grand marnier", "cointreau", "amaretto", "kirsch",
        "eau de vie", "biere", "bière", "cidre", "cafe", "café", "the", "thé",
        "eau gazeuse", "jus de", "sirop de grenadine", "sirop de menthe",
    ]),
]


def get_ingredient_rayon(name):
    """Retourne le rayon de magasin associé à un ingrédient, en se basant sur
    des mots-clés (recherche insensible aux accents et à la casse). Renvoie
    'Autre' si aucun mot-clé ne correspond."""
    key = ingredient_sort_key(name)
    for rayon, keywords in RAYON_KEYWORDS:
        for kw in keywords:
            if ingredient_sort_key(kw) in key:
                return rayon
    return "Autre"


RAYON_ORDER = [
    "Fruits & Légumes", "Viandes & Poissons", "Crèmerie",
    "Boulangerie & Pâtisserie", "Épicerie", "Herbes & Épices", "Boissons", "Autre",
]


def rating_stars(n):
    """Retourne une représentation textuelle d'une note sur 5 étoiles,
    ex: rating_stars(3) -> '★★★☆☆'."""
    try:
        n = int(n or 0)
    except (TypeError, ValueError):
        n = 0
    n = max(0, min(5, n))
    return "★" * n + "☆" * (5 - n)


def recipe_matches_search(recipe, search_key):
    """Vérifie si une recette correspond à une recherche (déjà normalisée via
    ingredient_sort_key), en comparant le nom ET les étiquettes."""
    if not search_key:
        return True
    if search_key in ingredient_sort_key(recipe["name"]):
        return True
    for tag in recipe.get("tags", []):
        if search_key in ingredient_sort_key(tag):
            return True
    return False


RECIPE_SORT_OPTIONS = ["Nom (A-Z)", "Temps de préparation", "Difficulté", "Note", "Ajoutées récemment"]
_DIFFICULTY_ORDER = {"Facile": 1, "Moyen": 2, "Difficile": 3}


def recipe_sort_key(recipe, option):
    """Retourne une clé de tri pour une recette selon l'option choisie parmi
    RECIPE_SORT_OPTIONS."""
    if option == "Temps de préparation":
        try:
            return float(recipe.get("prep_time") or 0)
        except (TypeError, ValueError):
            return 0.0
    if option == "Difficulté":
        return _DIFFICULTY_ORDER.get(recipe.get("difficulty"), 0)
    if option == "Note":
        return -int(recipe.get("rating", 0) or 0)  # négatif : meilleure note en premier
    if option == "Ajoutées récemment":
        return recipe.get("created_at") or ""
    return ingredient_sort_key(recipe["name"])


def format_recipe_list_label(recipe):
    """Construit le libellé affiché pour une recette dans les listes de
    sélection : favori, catégorie, nom, note, et — pour qu'on les voie tout
    de suite lors d'un tri — le temps total, la difficulté et les
    allergènes éventuels."""
    cat = recipe.get("category", "Autre")
    star = "⭐ " if recipe.get("favorite") else ""
    rating = recipe.get("rating", 0)
    rating_suffix = f" {rating_stars(rating)}" if rating else ""
    label = f"{star}[{cat}] {recipe['name']}{rating_suffix}"

    info_bits = []
    prep = recipe.get("prep_time")
    cook = recipe.get("cook_time")
    if prep or cook:
        try:
            total = float(prep or 0) + float(cook or 0)
            total_display = int(total) if total == int(total) else total
        except (TypeError, ValueError):
            total_display = None
        if total_display is not None:
            info_bits.append(f"{total_display} min")
    difficulty = recipe.get("difficulty")
    if difficulty:
        info_bits.append(difficulty)
    allergens = recipe.get("allergens") or []
    if allergens:
        info_bits.append(f"⚠ {', '.join(allergens)}")
    if info_bits:
        label += f"  ({' · '.join(info_bits)})"
    return label


def load_recipes():
    """Charge les recettes depuis le fichier JSON local."""
    if os.path.exists(DATA_FILE):
        try:
            with open(DATA_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return []
    return []


def save_recipes(recipes):
    """Sauvegarde la liste des recettes dans le fichier JSON local."""
    with open(DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(recipes, f, ensure_ascii=False, indent=2)


def load_ingredients():
    """Charge la liste des ingrédients connus (triée, sans doublons)."""
    if os.path.exists(INGREDIENTS_FILE):
        try:
            with open(INGREDIENTS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list):
                    return sorted(dict.fromkeys(normalize_oe(i) for i in data), key=ingredient_sort_key)
        except Exception:
            return []
    return []


def save_ingredients(ingredients):
    """Sauvegarde la liste des ingrédients connus (triée, sans doublons)."""
    cleaned_map = {}
    for i in ingredients:
        name = normalize_oe(i.strip())
        if name:
            cleaned_map.setdefault(name.lower(), name)
    cleaned = sorted(cleaned_map.values(), key=ingredient_sort_key)
    with open(INGREDIENTS_FILE, "w", encoding="utf-8") as f:
        json.dump(cleaned, f, ensure_ascii=False, indent=2)
    return cleaned


def sync_ingredients_from_recipes():
    """S'assure que tout ingrédient déjà utilisé dans une recette existante
    figure bien dans la liste des ingrédients connus (utile lors de la
    première utilisation de cette fonctionnalité, ou après import de
    données). Lors du tout premier lancement (aucun ingredients.json), la
    liste des ~1000 ingrédients courants fournis avec l'application est
    utilisée comme point de départ."""
    first_run = not os.path.exists(INGREDIENTS_FILE)
    known = load_default_ingredients() if first_run else load_ingredients()
    known_lower = {i.lower() for i in known}
    changed = first_run
    for recipe in load_recipes():
        for ing in recipe.get("ingredients", []):
            name = ing.get("name", "").strip()
            if name and name.lower() not in known_lower:
                known.append(name)
                known_lower.add(name.lower())
                changed = True
    if changed:
        return save_ingredients(known)
    return known


def merge_default_ingredients():
    """Ajoute à la liste actuelle les ingrédients courants fournis avec
    l'application qui ne seraient pas déjà présents. Retourne le nombre
    d'ingrédients réellement ajoutés."""
    current = load_ingredients()
    current_lower = {i.lower() for i in current}
    defaults = load_default_ingredients()
    added = 0
    for name in defaults:
        if name.lower() not in current_lower:
            current.append(name)
            current_lower.add(name.lower())
            added += 1
    save_ingredients(current)
    return added


def rename_ingredient_everywhere(old_name, new_name):
    """Renomme un ingrédient dans toutes les recettes qui l'utilisent."""
    recipes = load_recipes()
    changed = False
    for recipe in recipes:
        for ing in recipe.get("ingredients", []):
            if ing.get("name", "").strip().lower() == old_name.strip().lower():
                ing["name"] = new_name
                changed = True
    if changed:
        save_recipes(recipes)


def count_ingredient_usage(name):
    """Compte le nombre de recettes utilisant cet ingrédient."""
    count = 0
    for recipe in load_recipes():
        for ing in recipe.get("ingredients", []):
            if ing.get("name", "").strip().lower() == name.strip().lower():
                count += 1
                break
    return count


# ---------------------------------------------------------------------------
# Estimation du coût d'une recette : les prix sont renseignés par
# l'utilisateur (aucune source de prix en ligne n'est disponible/fiable pour
# une application locale), un ingrédient par un, dans "Gérer les prix".
# ---------------------------------------------------------------------------

PRICE_UNIT_OPTIONS = ["kg", "L", "pièce", "cuillère à soupe", "cuillère à café"]


def load_ingredient_prices():
    """Retourne {nom_ingredient_en_minuscules: {"name":, "price":, "unit":}}."""
    if os.path.exists(INGREDIENT_PRICES_FILE):
        try:
            with open(INGREDIENT_PRICES_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, dict):
                    return data
        except Exception:
            return {}
    return {}


def save_ingredient_prices(prices):
    with open(INGREDIENT_PRICES_FILE, "w", encoding="utf-8") as f:
        json.dump(prices, f, ensure_ascii=False, indent=2)


def get_ingredient_price(name):
    return load_ingredient_prices().get(name.strip().lower())


def set_ingredient_price(name, price, unit):
    """price=None efface le prix enregistré pour cet ingrédient."""
    prices = load_ingredient_prices()
    key = name.strip().lower()
    if price is None:
        prices.pop(key, None)
    else:
        prices[key] = {"name": name.strip(), "price": price, "unit": unit}
    save_ingredient_prices(prices)


def compute_recipe_cost(recipe, persons):
    """Retourne (coût_total_estimé, nb_ingrédients_avec_prix_connu,
    nb_ingrédients_total). Un ingrédient sans prix renseigné, ou dont
    l'unité ne peut pas être convertie vers celle du prix, est ignoré
    (le coût rendu est donc une estimation partielle si des prix manquent)."""
    prices = load_ingredient_prices()
    total = 0.0
    known = 0
    total_count = len(recipe["ingredients"])
    for ing in recipe["ingredients"]:
        price_info = prices.get(ing["name"].strip().lower())
        if not price_info:
            continue
        qty = ing["quantity"] * persons
        unit_lower = ing["unit"].strip().lower()
        price = price_info["price"]
        price_unit = price_info["unit"]
        contribution = None
        if price_unit == "kg" and unit_lower in ("gr", "g", "gramme", "grammes"):
            contribution = (qty / 1000.0) * price
        elif price_unit == "L" and unit_lower == "cl":
            contribution = (qty / 100.0) * price
        elif price_unit.lower() == unit_lower:
            contribution = qty * price
        if contribution is not None:
            total += contribution
            known += 1
    return total, known, total_count


# ---------------------------------------------------------------------------
# Estimation des valeurs nutritionnelles d'une recette, à partir d'une base
# de valeurs typiques par ingrédient (kcal / protéines / glucides / lipides
# pour 100 g ou 100 ml), fournie avec l'application.
# ---------------------------------------------------------------------------

_nutrition_cache = None

# Équivalence approximative de chaque unité de recette vers des grammes (ou
# millilitres, assimilés à des grammes pour les liquides — approximation
# standard). Les unités "pièce" et "autre" ne sont pas converties : le poids
# d'une "pièce" dépend trop de l'ingrédient pour être généralisé.
UNIT_TO_GRAMS = {
    "gr": 1.0, "g": 1.0, "gramme": 1.0, "grammes": 1.0,
    "cl": 10.0,
    "cuillère à soupe": 15.0,
    "cuillère à café": 5.0,
}


def load_nutrition_data():
    """Charge (une seule fois, en cache) la base de valeurs nutritionnelles
    fournie avec l'application."""
    global _nutrition_cache
    if _nutrition_cache is not None:
        return _nutrition_cache
    data = {}
    if os.path.exists(NUTRITION_DATA_FILE):
        try:
            with open(NUTRITION_DATA_FILE, "r", encoding="utf-8") as f:
                raw = json.load(f)
                if isinstance(raw, dict):
                    data = {k.strip().lower(): v for k, v in raw.items()}
        except Exception:
            data = {}
    _nutrition_cache = data
    return _nutrition_cache


def get_ingredient_nutrition(name):
    """Retourne le dict nutrition {kcal, protein_g, carbs_g, fat_g} pour un
    ingrédient (pour 100 g/100 ml), ou None si inconnu. Une surcharge
    personnelle est prioritaire sur la base fournie."""
    override = get_ingredient_override(name)
    if override and "nutrition" in override and override["nutrition"]:
        return override["nutrition"]
    return load_nutrition_data().get(name.strip().lower())


def compute_recipe_nutrition(recipe, persons):
    """Retourne (totaux {kcal, protein_g, carbs_g, fat_g}, nb_ingrédients pris
    en compte, nb_ingrédients_total). Les ingrédients inconnus de la base, ou
    exprimés en "pièce"/unité personnalisée, sont exclus du total (estimation
    partielle dans ce cas)."""
    totals = {"kcal": 0.0, "protein_g": 0.0, "carbs_g": 0.0, "fat_g": 0.0}
    converted = 0
    total_count = len(recipe["ingredients"])
    for ing in recipe["ingredients"]:
        nutri = get_ingredient_nutrition(ing["name"])
        if not nutri:
            continue
        grams_per_unit = UNIT_TO_GRAMS.get(ing["unit"].strip().lower())
        if grams_per_unit is None:
            continue
        grams = ing["quantity"] * persons * grams_per_unit
        factor = grams / 100.0
        totals["kcal"] += nutri.get("kcal", 0) * factor
        totals["protein_g"] += nutri.get("protein_g", 0) * factor
        totals["carbs_g"] += nutri.get("carbs_g", 0) * factor
        totals["fat_g"] += nutri.get("fat_g", 0) * factor
        converted += 1
    return totals, converted, total_count


# ---------------------------------------------------------------------------
# Allergènes présents dans chaque ingrédient, à partir d'une base fournie
# avec l'application (les ~1000 ingrédients courants). Sert à détecter
# automatiquement les allergènes d'une recette à partir de ses ingrédients.
# ---------------------------------------------------------------------------

_ingredient_allergens_cache = None


def load_ingredient_allergens():
    """Charge (une seule fois, en cache) la base des allergènes par
    ingrédient fournie avec l'application."""
    global _ingredient_allergens_cache
    if _ingredient_allergens_cache is not None:
        return _ingredient_allergens_cache
    data = {}
    if os.path.exists(INGREDIENT_ALLERGENS_FILE):
        try:
            with open(INGREDIENT_ALLERGENS_FILE, "r", encoding="utf-8") as f:
                raw = json.load(f)
                if isinstance(raw, dict):
                    data = {k.strip().lower(): v for k, v in raw.items()}
        except Exception:
            data = {}
    _ingredient_allergens_cache = data
    return _ingredient_allergens_cache


# ---------------------------------------------------------------------------
# Surcharges personnelles par ingrédient (allergènes et/ou valeurs
# nutritionnelles modifiés par l'utilisateur, ou définis pour un ingrédient
# créé par lui). Prioritaires sur les bases fournies avec l'application.
# ---------------------------------------------------------------------------

def load_ingredient_overrides():
    if os.path.exists(INGREDIENT_OVERRIDES_FILE):
        try:
            with open(INGREDIENT_OVERRIDES_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, dict):
                    return data
        except Exception:
            return {}
    return {}


def save_ingredient_overrides(overrides):
    with open(INGREDIENT_OVERRIDES_FILE, "w", encoding="utf-8") as f:
        json.dump(overrides, f, ensure_ascii=False, indent=2)


def get_ingredient_override(name):
    return load_ingredient_overrides().get(name.strip().lower())


def set_ingredient_override(name, allergens=None, nutrition=None):
    """allergens : liste (peut être vide) ou None pour ne pas y toucher.
    nutrition : dict {kcal, protein_g, carbs_g, fat_g} ou None pour ne pas
    y toucher (passer un dict vide {} pour effacer la surcharge nutrition)."""
    overrides = load_ingredient_overrides()
    key = name.strip().lower()
    entry = overrides.get(key, {})
    if allergens is not None:
        entry["allergens"] = allergens
    if nutrition is not None:
        entry["nutrition"] = nutrition
    if entry:
        overrides[key] = entry
    else:
        overrides.pop(key, None)
    save_ingredient_overrides(overrides)


def rename_ingredient_override(old_name, new_name):
    """Transfère une éventuelle surcharge de l'ancien nom vers le nouveau,
    lors d'un renommage d'ingrédient."""
    overrides = load_ingredient_overrides()
    old_key = old_name.strip().lower()
    new_key = new_name.strip().lower()
    if old_key in overrides and old_key != new_key:
        overrides[new_key] = overrides.pop(old_key)
        save_ingredient_overrides(overrides)


def get_ingredient_allergens(name):
    """Retourne la liste des allergènes connus pour un ingrédient donné
    (liste vide si l'ingrédient n'est pas reconnu ou n'en contient aucun).
    Une surcharge personnelle est prioritaire sur la base fournie."""
    override = get_ingredient_override(name)
    if override and "allergens" in override:
        return override["allergens"]
    return load_ingredient_allergens().get(name.strip().lower(), [])


def compute_recipe_allergens(ingredients):
    """Retourne l'ensemble des allergènes détectés à partir d'une liste
    d'ingrédients de recette, dans l'ordre de la liste ALLERGENS."""
    detected = set()
    for ing in ingredients:
        detected.update(get_ingredient_allergens(ing["name"]))
    return [a for a in ALLERGENS if a in detected]


def find_similar_ingredient_pairs(names, threshold=0.82):
    """Retourne une liste de paires (nom_a, nom_b, score) d'ingrédients dont
    les noms se ressemblent fortement (accents/casse ignorés), pouvant
    indiquer un doublon ou une faute de frappe (ex. "Tomate" / "Tomates",
    "Echalotte" / "Échalote"). Les ingrédients sont d'abord regroupés par
    leurs deux premières lettres pour limiter le nombre de comparaisons sur
    de grandes listes."""
    normalized = [(n, ingredient_sort_key(n)) for n in names]
    buckets = {}
    for name, key in normalized:
        prefix = key[:2] if len(key) >= 2 else key
        buckets.setdefault(prefix, []).append((name, key))

    pairs = []
    for bucket in buckets.values():
        for i in range(len(bucket)):
            name_a, key_a = bucket[i]
            for j in range(i + 1, len(bucket)):
                name_b, key_b = bucket[j]
                if key_a == key_b:
                    continue
                is_plural_variant = (
                    key_a in (key_b + "s", key_b + "x") or key_b in (key_a + "s", key_a + "x")
                )
                ratio = difflib.SequenceMatcher(None, key_a, key_b).ratio()
                if is_plural_variant or ratio >= threshold:
                    pairs.append((name_a, name_b, ratio if not is_plural_variant else max(ratio, 0.9)))

    pairs.sort(key=lambda t: -t[2])
    return pairs


def copy_image_to_store(source_path):
    """Copie une image choisie par l'utilisateur dans le dossier images/
    et retourne le nom de fichier généré (à stocker dans la recette)."""
    ext = os.path.splitext(source_path)[1].lower()
    if ext not in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"):
        ext = ".png"
    new_filename = f"{uuid.uuid4().hex}{ext}"
    dest_path = os.path.join(IMAGES_DIR, new_filename)
    shutil.copy2(source_path, dest_path)
    return new_filename


def delete_image_file(image_filename):
    """Supprime un fichier image du dossier images/ (si présent)."""
    if not image_filename:
        return
    path = os.path.join(IMAGES_DIR, image_filename)
    if os.path.exists(path):
        try:
            os.remove(path)
        except OSError:
            pass


def load_thumbnail(image_filename, size=(240, 180)):
    """Retourne un objet ImageTk.PhotoImage pour affichage, ou None si
    l'image n'existe pas ou si Pillow n'est pas installé."""
    if not image_filename or not PIL_AVAILABLE:
        return None
    path = os.path.join(IMAGES_DIR, image_filename)
    if not os.path.exists(path):
        return None
    try:
        img = Image.open(path)
        img.thumbnail(size)
        return ImageTk.PhotoImage(img)
    except Exception:
        return None


def get_recipe_images(recipe):
    """Retourne la liste des noms de fichiers image d'une recette, en gérant
    la compatibilité avec l'ancien format à une seule photo (clé 'image')."""
    images = recipe.get("images")
    if images:
        return list(images)
    legacy = recipe.get("image")
    return [legacy] if legacy else []


def delete_recipe_images(recipe):
    """Supprime tous les fichiers photo associés à une recette."""
    for fname in get_recipe_images(recipe):
        delete_image_file(fname)


def duplicate_recipe_images(recipe):
    """Copie physiquement tous les fichiers photo d'une recette sous de
    nouveaux noms, pour qu'une recette dupliquée ait ses propres fichiers
    indépendants (supprimer l'une ne doit pas affecter l'autre)."""
    new_names = []
    for fname in get_recipe_images(recipe):
        src = os.path.join(IMAGES_DIR, fname)
        if os.path.exists(src):
            new_names.append(copy_image_to_store(src))
    return new_names


# ---------------------------------------------------------------------------
# Fonctions partagées de calcul / export de liste de courses, réutilisées par
# "Voir toutes les recettes", le planificateur de repas et les menus.
# ---------------------------------------------------------------------------

def format_quantity_with_unit(qty, unit):
    """Convertit automatiquement une grande quantité vers une unité plus
    parlante pour un total de liste de courses : les grammes passent en
    kilogrammes au-delà de 1000 g (1 kg = 1000 g), et les centilitres
    passent en litres au-delà de 100 cl (1 L = 100 cl)."""
    unit_lower = (unit or "").strip().lower()
    if unit_lower in ("gr", "g", "gramme", "grammes") and qty >= 1000:
        qty = qty / 1000
        unit = "kg"
    elif unit_lower == "cl" and qty >= 100:
        qty = qty / 100
        unit = "L"
    qty = round(qty, 2)
    if qty == int(qty):
        qty = int(qty)
    return qty, unit


def compute_grouped_totals(recipe_persons_pairs):
    """À partir d'une liste de (recette, nombre_de_personnes), calcule le
    total de chaque ingrédient nécessaire, regroupé par rayon de magasin
    (dans l'ordre RAYON_ORDER). Retourne grouped_totals :
    [(rayon, [(nom, quantité, unité), ...]), ...]. Les grandes quantités
    (≥ 1000 g, ≥ 100 cl) sont automatiquement affichées en kg / L."""
    totals = {}
    for recipe, persons in recipe_persons_pairs:
        for ing in recipe["ingredients"]:
            key = (ing["name"].strip().lower(), ing["unit"].strip().lower())
            totals[key] = totals.get(key, 0) + ing["quantity"] * persons

    by_rayon = {}
    for (name, unit), qty in totals.items():
        display_qty, display_unit = format_quantity_with_unit(qty, unit)
        rayon = get_ingredient_rayon(name)
        by_rayon.setdefault(rayon, []).append((name.capitalize(), display_qty, display_unit))

    grouped_totals = []
    for rayon in RAYON_ORDER:
        if rayon in by_rayon:
            items = sorted(by_rayon[rayon], key=lambda x: ingredient_sort_key(x[0]))
            grouped_totals.append((rayon, items))
    return grouped_totals


def write_shopping_list_txt(path, title, chosen_recipes, grouped_totals):
    """Écrit une liste de courses au format texte. chosen_recipes est une
    liste de (libellé_affiché, personnes)."""
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"=== {title} ===\n\n")
        f.write(datetime.now().strftime("Générée le %d/%m/%Y à %H:%M") + "\n\n")
        f.write("Recettes sélectionnées :\n")
        for label, persons in chosen_recipes:
            f.write(f"- {label} ({persons} pers.)\n")
        for rayon, items in grouped_totals:
            f.write(f"\n{rayon} :\n")
            for name, qty, unit in items:
                unit_display = f" {unit}" if unit else ""
                f.write(f"- {name} : {qty}{unit_display}\n")


def build_shopping_list_workbook(chosen_recipes, grouped_totals):
    """Construit un classeur Excel (openpyxl) pour une liste de courses.
    Nécessite que OPENPYXL_AVAILABLE soit vrai."""
    wb = Workbook()
    ws_recipes = wb.active
    ws_recipes.title = "Recettes"
    ws_recipes.append(["Recette", "Nombre de personnes"])
    for label, persons in chosen_recipes:
        ws_recipes.append([label, persons])

    ws_ing = wb.create_sheet("Ingrédients")
    ws_ing.append(["Rayon", "Ingrédient", "Quantité totale", "Unité"])
    for rayon, items in grouped_totals:
        for name, qty, unit in items:
            ws_ing.append([rayon, name, qty, unit])
    return wb


def build_shopping_list_pdf(path, title, chosen_recipes, grouped_totals):
    """Construit un PDF pour une liste de courses. Nécessite que
    REPORTLAB_AVAILABLE soit vrai."""
    c = pdf_canvas.Canvas(path, pagesize=A4)
    width, height = A4
    y = height - 2 * cm

    c.setFont("Helvetica-Bold", 18)
    c.drawString(2 * cm, y, title)
    y -= 1 * cm

    c.setFont("Helvetica", 10)
    c.drawString(2 * cm, y, datetime.now().strftime("Générée le %d/%m/%Y à %H:%M"))
    y -= 1 * cm

    c.setFont("Helvetica-Bold", 12)
    c.drawString(2 * cm, y, "Recettes sélectionnées :")
    y -= 0.6 * cm
    c.setFont("Helvetica", 10)
    for label, persons in chosen_recipes:
        c.drawString(2.3 * cm, y, f"- {label} ({persons} pers.)")
        y -= 0.5 * cm
        if y < 3 * cm:
            c.showPage()
            y = height - 2 * cm

    for rayon, items in grouped_totals:
        y -= 0.4 * cm
        if y < 3.5 * cm:
            c.showPage()
            y = height - 2 * cm
        c.setFont("Helvetica-Bold", 12)
        c.drawString(2 * cm, y, f"{rayon} :")
        y -= 0.6 * cm
        c.setFont("Helvetica", 10)
        for name, qty, unit in items:
            unit_display = f" {unit}" if unit else ""
            c.drawString(2.3 * cm, y, f"- {name} : {qty}{unit_display}")
            y -= 0.5 * cm
            if y < 2 * cm:
                c.showPage()
                y = height - 2 * cm

    c.save()


ALLERGENS = ["Gluten", "Lactose", "Œufs", "Arachides", "Fruits à coque",
             "Soja", "Poisson", "Crustacés", "Sésame", "Céleri", "Moutarde",
             "Sulfites", "Lupin", "Mollusques"]

# Ingrédients de base qu'on a presque toujours sous la main, pré-cochés par
# défaut dans "Que puis-je cuisiner ?" (l'utilisateur reste libre de les
# retirer au cas par cas).
PANTRY_STAPLES = [
    "Sel", "Poivre", "Huile de tournesol", "Huile d'olive", "Beurre",
    "Farine", "Sucre", "Vinaigre", "Moutarde", "Riz", "Pâtes", "Lait",
]


def draw_recipe_content(c, recipe, persons, width, height):
    """Dessine le contenu complet d'une recette (titre, infos, photo,
    ingrédients, description, notes) sur un canevas reportlab déjà créé, à
    partir du haut d'une page. Retourne la position verticale (y) atteinte à
    la fin, pour pouvoir enchaîner d'autres contenus sur la même page si
    besoin (sinon l'appelant peut faire c.showPage() lui-même)."""
    y = height - 2 * cm

    star = "⭐ " if recipe.get("favorite") else ""
    c.setFont("Helvetica-Bold", 18)
    c.drawString(2 * cm, y, f"{star}{recipe['name']}")
    y -= 0.9 * cm

    cat = recipe.get("category", "Autre")
    c.setFont("Helvetica", 10)
    c.drawString(2 * cm, y, f"Catégorie : {cat}    Pour {persons} personne(s)")
    y -= 0.6 * cm

    rating = recipe.get("rating", 0)
    if rating:
        c.drawString(2 * cm, y, f"Note : {rating_stars(rating)}")
        y -= 0.6 * cm

    info_bits = []
    if recipe.get("prep_time"):
        info_bits.append(f"Préparation : {recipe['prep_time']} min")
    if recipe.get("cook_time"):
        info_bits.append(f"Cuisson : {recipe['cook_time']} min")
    if recipe.get("difficulty"):
        info_bits.append(f"Difficulté : {recipe['difficulty']}")
    if info_bits:
        c.drawString(2 * cm, y, "   |   ".join(info_bits))
        y -= 0.6 * cm

    allergens = recipe.get("allergens") or []
    if allergens:
        c.setFillColorRGB(0.7, 0.2, 0.2)
        c.drawString(2 * cm, y, f"⚠ Allergènes : {', '.join(allergens)}")
        c.setFillColorRGB(0, 0, 0)
        y -= 0.6 * cm

    y -= 0.2 * cm

    # Photo (la première disponible)
    images = get_recipe_images(recipe)
    if images:
        img_path = os.path.join(IMAGES_DIR, images[0])
        if os.path.exists(img_path):
            try:
                img_reader = ImageReader(img_path)
                iw, ih = img_reader.getSize()
                max_w = 8 * cm
                scale = max_w / iw
                draw_w = max_w
                draw_h = ih * scale
                if y - draw_h < 3 * cm:
                    c.showPage()
                    y = height - 2 * cm
                c.drawImage(img_reader, 2 * cm, y - draw_h, width=draw_w, height=draw_h,
                            preserveAspectRatio=True, mask="auto")
                y -= draw_h + 0.6 * cm
            except Exception:
                pass

    if y < 4 * cm:
        c.showPage()
        y = height - 2 * cm

    c.setFont("Helvetica-Bold", 12)
    c.drawString(2 * cm, y, "Ingrédients :")
    y -= 0.6 * cm
    c.setFont("Helvetica", 10)
    for ing in recipe["ingredients"]:
        qty = round(ing["quantity"] * persons, 2)
        unit = f" {ing['unit']}" if ing["unit"] else ""
        c.drawString(2.3 * cm, y, f"- {ing['name'].capitalize()} : {qty}{unit}")
        y -= 0.5 * cm
        if y < 3 * cm:
            c.showPage()
            y = height - 2 * cm

    cost, cost_known, cost_total = compute_recipe_cost(recipe, persons)
    nutrition, nutri_known, nutri_total = compute_recipe_nutrition(recipe, persons)
    if cost_known or nutri_known:
        y -= 0.3 * cm
        if y < 3 * cm:
            c.showPage()
            y = height - 2 * cm
        c.setFont("Helvetica-Oblique", 9)
        if cost_known:
            partial = "" if cost_known == cost_total else f" (partiel, {cost_known}/{cost_total})"
            c.drawString(2 * cm, y, f"Coût estimé : {cost:.2f} €{partial}")
            y -= 0.45 * cm
        if nutri_known:
            partial = "" if nutri_known == nutri_total else f" (partiel, {nutri_known}/{nutri_total})"
            c.drawString(
                2 * cm, y,
                f"Nutrition estimée{partial} : {nutrition['kcal']:.0f} kcal · "
                f"{nutrition['protein_g']:.0f}g prot. · {nutrition['carbs_g']:.0f}g gluc. · "
                f"{nutrition['fat_g']:.0f}g lip."
            )
            y -= 0.45 * cm
        c.setFont("Helvetica", 10)

    def draw_wrapped_section(title, text, y):
        y -= 0.4 * cm
        if y < 4 * cm:
            c.showPage()
            y = height - 2 * cm
        c.setFont("Helvetica-Bold", 12)
        c.drawString(2 * cm, y, title)
        y -= 0.6 * cm
        c.setFont("Helvetica", 10)
        for paragraph in text.split("\n"):
            words = paragraph.split(" ")
            line = ""
            for word in words:
                test_line = f"{line} {word}".strip()
                if c.stringWidth(test_line, "Helvetica", 10) > (width - 4 * cm):
                    c.drawString(2 * cm, y, line)
                    y -= 0.5 * cm
                    if y < 2 * cm:
                        c.showPage()
                        y = height - 2 * cm
                    line = word
                else:
                    line = test_line
            if line:
                c.drawString(2 * cm, y, line)
                y -= 0.5 * cm
                if y < 2 * cm:
                    c.showPage()
                    y = height - 2 * cm
        return y

    description = recipe.get("description", "").strip()
    if description:
        y = draw_wrapped_section("Description :", description, y)

    personal_notes = recipe.get("personal_notes", "").strip()
    if personal_notes:
        y = draw_wrapped_section("Notes personnelles :", personal_notes, y)

    return y


def build_cookbook_pdf(path, recipes_with_persons):
    """Construit un PDF regroupant plusieurs recettes à la suite (une page de
    titre listant le sommaire, puis une recette par page)."""
    c = pdf_canvas.Canvas(path, pagesize=A4)
    width, height = A4

    # Page de titre / sommaire
    y = height - 3 * cm
    c.setFont("Helvetica-Bold", 24)
    c.drawString(2 * cm, y, "Mon Livre de Recettes")
    y -= 1 * cm
    c.setFont("Helvetica", 10)
    c.drawString(2 * cm, y, datetime.now().strftime("Généré le %d/%m/%Y"))
    y -= 1.2 * cm
    c.setFont("Helvetica-Bold", 13)
    c.drawString(2 * cm, y, "Sommaire")
    y -= 0.7 * cm
    c.setFont("Helvetica", 10)
    for recipe, persons in recipes_with_persons:
        cat = recipe.get("category", "Autre")
        c.drawString(2.3 * cm, y, f"- [{cat}] {recipe['name']}")
        y -= 0.5 * cm
        if y < 2 * cm:
            c.showPage()
            y = height - 2 * cm

    for recipe, persons in recipes_with_persons:
        c.showPage()
        draw_recipe_content(c, recipe, persons, width, height)

    c.save()


WEEKDAYS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]


def load_weekly_plan():
    """Charge le planning de la semaine : {jour: {'recipe_name':.., 'persons':..}}."""
    if os.path.exists(WEEKLY_PLAN_FILE):
        try:
            with open(WEEKLY_PLAN_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, dict):
                    return data
        except Exception:
            return {}
    return {}


def save_weekly_plan(plan):
    with open(WEEKLY_PLAN_FILE, "w", encoding="utf-8") as f:
        json.dump(plan, f, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Export du planning de la semaine vers un fichier .ics (iCalendar), lisible
# par Google Agenda, Outlook, Apple Calendrier, etc.
# ---------------------------------------------------------------------------

_ICS_MEAL_TIMES = {
    "Petit-déjeuner": ("0800", "0830"),
    "Déjeuner": ("1230", "1330"),
    "Dîner": ("1930", "2030"),
}
_ICS_MEAL_GROUPS = {
    "Petit-déjeuner": ["Petit-déjeuner"],
    "Déjeuner": ["Déjeuner — Entrée", "Déjeuner — Plat", "Déjeuner — Dessert"],
    "Dîner": ["Dîner — Entrée", "Dîner — Plat", "Dîner — Dessert"],
}


def _escape_ics_text(text):
    return (text.replace("\\", "\\\\")
                .replace(";", "\\;")
                .replace(",", "\\,")
                .replace("\n", "\\n"))


def _fold_ics_line(line):
    """Replie une ligne selon la limite de 75 octets recommandée par la
    norme iCalendar (RFC 5545), pour une compatibilité maximale."""
    if len(line.encode("utf-8")) <= 75:
        return line
    parts = []
    current = ""
    for ch in line:
        if len((current + ch).encode("utf-8")) > 74:
            parts.append(current)
            current = " " + ch
        else:
            current += ch
    parts.append(current)
    return "\r\n".join(parts)


def build_weekly_plan_ics(plan):
    """Construit le contenu d'un fichier .ics à partir du planning de la
    semaine : un évènement hebdomadaire récurrent par repas (petit-déjeuner,
    déjeuner, dîner) et par jour où quelque chose est prévu."""
    weekday_index = {d: i for i, d in enumerate(WEEKDAYS)}  # Lundi=0 ... Dimanche=6
    today = datetime.now().date()
    now_stamp = datetime.now().strftime("%Y%m%dT%H%M%S")

    lines = ["BEGIN:VCALENDAR", "VERSION:2.0",
             "PRODID:-//Mon Livre de Recettes//FR", "CALSCALE:GREGORIAN"]
    event_count = 0

    for day, slots_data in plan.items():
        if day not in weekday_index or not slots_data:
            continue
        delta_days = (weekday_index[day] - today.weekday()) % 7
        event_date = today + timedelta(days=delta_days)
        date_str = event_date.strftime("%Y%m%d")

        for meal_period, slot_names in _ICS_MEAL_GROUPS.items():
            components = []
            for slot in slot_names:
                info = slots_data.get(slot)
                if info and info.get("recipe_name"):
                    label = slot.split(" — ", 1)[1] if " — " in slot else None
                    components.append(f"{label} : {info['recipe_name']}" if label else info["recipe_name"])
            if not components:
                continue

            start_time, end_time = _ICS_MEAL_TIMES[meal_period]
            summary_names = [c.split(" : ", 1)[1] if " : " in c else c for c in components]
            summary = f"{meal_period} : " + ", ".join(summary_names)
            description = "\\n".join(components)

            event_count += 1
            uid = f"recette-{event_count}-{uuid.uuid4().hex}@monlivrederecettes"

            lines.append("BEGIN:VEVENT")
            lines.append(_fold_ics_line(f"UID:{uid}"))
            lines.append(f"DTSTAMP:{now_stamp}")
            lines.append(f"DTSTART:{date_str}T{start_time}00")
            lines.append(f"DTEND:{date_str}T{end_time}00")
            lines.append("RRULE:FREQ=WEEKLY")
            lines.append(_fold_ics_line(f"SUMMARY:{_escape_ics_text(summary)}"))
            lines.append(_fold_ics_line(f"DESCRIPTION:{_escape_ics_text(description)}"))
            lines.append("END:VEVENT")

    lines.append("END:VCALENDAR")
    return "\r\n".join(lines) + "\r\n"


def load_menus():
    """Charge la liste des menus enregistrés :
    [{'name':.., 'items': [{'recipe_name':.., 'persons':..}, ...]}, ...]."""
    if os.path.exists(MENUS_FILE):
        try:
            with open(MENUS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list):
                    return data
        except Exception:
            return []
    return []


def save_menus(menus):
    with open(MENUS_FILE, "w", encoding="utf-8") as f:
        json.dump(menus, f, ensure_ascii=False, indent=2)


def find_recipe_by_name(recipes, name):
    return next((r for r in recipes if r.get("name") == name), None)


# ---------------------------------------------------------------------------
# Import d'une recette depuis un lien internet (données structurées
# Schema.org "Recipe", utilisées par la plupart des sites de cuisine).
# ---------------------------------------------------------------------------

_FRACTION_MAP = {"½": 0.5, "¼": 0.25, "¾": 0.75, "⅓": 1 / 3, "⅔": 2 / 3}

_UNIT_ALIASES = {
    "g": "Gr", "gr": "Gr", "gramme": "Gr", "grammes": "Gr",
    "cl": "cl",
    "cas": "cuillère à soupe", "c.a.s": "cuillère à soupe",
    "cuillere": "cuillère à soupe", "cuillères": "cuillère à soupe", "cuillère": "cuillère à soupe",
    "cuillere a soupe": "cuillère à soupe", "cuillère à soupe": "cuillère à soupe",
    "cuillères à soupe": "cuillère à soupe", "cuilleres a soupe": "cuillère à soupe",
    "cac": "cuillère à café",
    "cuillere a cafe": "cuillère à café", "cuillère à café": "cuillère à café",
    "cuillères à café": "cuillère à café", "cuilleres a cafe": "cuillère à café",
    "piece": "pièce", "pièce": "pièce", "pièces": "pièce", "pieces": "pièce",
}


def parse_quantity_token(token):
    token = token.strip()
    if token in _FRACTION_MAP:
        return _FRACTION_MAP[token]
    m = re.match(r"^(\d+)\s*/\s*(\d+)$", token)
    if m:
        return int(m.group(1)) / int(m.group(2))
    token = token.replace(",", ".")
    try:
        return float(token)
    except ValueError:
        return None


def _match_unit_tokens(tokens, start):
    """Essaie de faire correspondre 3, puis 2, puis 1 token(s) à partir de
    `start` à une unité connue (ex. « cuillères à soupe »), en testant la
    séquence la plus longue en premier. Insensible aux accents/à la casse."""
    for length in (3, 2, 1):
        end = start + length
        if end <= len(tokens):
            candidate = " ".join(t.strip(".,") for t in tokens[start:end])
            candidate_key = ingredient_sort_key(candidate)
            for alias, unit in _UNIT_ALIASES.items():
                if ingredient_sort_key(alias) == candidate_key:
                    return unit, end
    return None, start


def parse_ingredient_line(line):
    """Tente d'extraire (nom, quantité, unité) d'une ligne d'ingrédient en
    texte libre, telle que fournie par un site de recettes. Retourne un
    résultat raisonnable même quand le texte ne suit pas un format standard
    (l'utilisateur pourra toujours corriger à la main après import)."""
    original = line.strip()
    if not original:
        return None
    tokens = original.split()
    if not tokens:
        return None

    qty = parse_quantity_token(tokens[0])
    unit = None
    rest_start = 0
    if qty is not None:
        rest_start = 1
        if len(tokens) > 1:
            unit, rest_start = _match_unit_tokens(tokens, 1)

    name = " ".join(tokens[rest_start:]).strip()
    name = re.sub(r"^(de |d['’]|of )", "", name, flags=re.IGNORECASE).strip()

    if not name:
        name = original
        qty = qty or 1
        unit = unit or "pièce"

    return {"name": name.capitalize(), "quantity": qty or 1, "unit": unit or "pièce"}


def parse_iso8601_duration_minutes(duration):
    """Convertit une durée ISO 8601 (ex. 'PT1H30M') en nombre de minutes."""
    if not duration:
        return None
    match = re.search(r"PT(?:(\d+)H)?(?:(\d+)M)?(?:(\d+)S)?", str(duration))
    if not match:
        return None
    hours = int(match.group(1) or 0)
    minutes = int(match.group(2) or 0)
    seconds = int(match.group(3) or 0)
    total = hours * 60 + minutes + (1 if seconds >= 30 else 0)
    return total if total > 0 else None


def _find_recipe_jsonld(data):
    """Cherche récursivement un objet Schema.org de type Recipe dans une
    structure JSON-LD (qui peut être un objet, une liste, ou contenir
    '@graph')."""
    if isinstance(data, dict):
        type_value = data.get("@type")
        types = type_value if isinstance(type_value, list) else [type_value]
        if any(t and "recipe" in str(t).lower() for t in types):
            return data
        if "@graph" in data:
            found = _find_recipe_jsonld(data["@graph"])
            if found:
                return found
        for value in data.values():
            found = _find_recipe_jsonld(value)
            if found:
                return found
    elif isinstance(data, list):
        for item in data:
            found = _find_recipe_jsonld(item)
            if found:
                return found
    return None


def fetch_recipe_from_url(url):
    """Télécharge une page de recette et tente d'en extraire le contenu à
    partir des données structurées Schema.org (JSON-LD), un format utilisé
    par la grande majorité des sites de recettes. Lève une exception avec un
    message clair en cas d'échec."""
    request = urllib.request.Request(
        url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
    )
    try:
        with urllib.request.urlopen(request, timeout=15) as response:
            raw = response.read()
            charset = response.headers.get_content_charset() or "utf-8"
            page_html = raw.decode(charset, errors="replace")
    except urllib.error.URLError as e:
        raise RuntimeError(f"Impossible d'accéder à cette adresse : {e}")
    except Exception as e:
        raise RuntimeError(f"Erreur lors du téléchargement de la page : {e}")

    scripts = re.findall(
        r'<script[^>]+type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
        page_html, flags=re.DOTALL | re.IGNORECASE
    )

    recipe_data = None
    for script in scripts:
        try:
            parsed = json.loads(script.strip())
        except (json.JSONDecodeError, ValueError):
            continue
        found = _find_recipe_jsonld(parsed)
        if found:
            recipe_data = found
            break

    if recipe_data is None:
        raise RuntimeError(
            "Aucune recette structurée n'a été trouvée sur cette page.\n\n"
            "Cet import fonctionne avec les sites qui utilisent le format "
            "standard « Schema.org Recipe » (la plupart des grands sites de "
            "cuisine). Vous pouvez toujours créer la recette manuellement."
        )

    def clean_text(value):
        if isinstance(value, list):
            value = " ".join(str(v) for v in value)
        return html.unescape(re.sub(r"<[^>]+>", " ", str(value or ""))).strip()

    name = clean_text(recipe_data.get("name", "")) or "Recette importée"

    raw_ingredients = recipe_data.get("recipeIngredient") or recipe_data.get("ingredients") or []
    if isinstance(raw_ingredients, str):
        raw_ingredients = [raw_ingredients]
    ingredients = []
    for line in raw_ingredients:
        parsed_ing = parse_ingredient_line(clean_text(line))
        if parsed_ing:
            ingredients.append(parsed_ing)

    instructions = recipe_data.get("recipeInstructions") or []
    steps = []
    if isinstance(instructions, str):
        steps = [instructions]
    elif isinstance(instructions, list):
        for item in instructions:
            if isinstance(item, dict):
                text = item.get("text") or item.get("name") or ""
                steps.append(clean_text(text))
            else:
                steps.append(clean_text(item))
    steps = [s for s in steps if s]
    description = "\n".join(f"{i + 1}. {s}" for i, s in enumerate(steps))

    prep_time = parse_iso8601_duration_minutes(recipe_data.get("prepTime"))
    cook_time = parse_iso8601_duration_minutes(recipe_data.get("cookTime"))

    yield_value = recipe_data.get("recipeYield")
    default_persons = None
    if yield_value:
        if isinstance(yield_value, list):
            yield_value = yield_value[0] if yield_value else None
        match = re.search(r"\d+", str(yield_value or ""))
        if match:
            default_persons = int(match.group())

    if not ingredients:
        raise RuntimeError(
            "Une recette a été trouvée sur cette page, mais aucun ingrédient "
            "n'a pu en être extrait. Vous pouvez créer la recette "
            "manuellement à la place."
        )

    # Récupération de la photo de la recette, si le site en indique une.
    images = []
    image_url = _extract_recipe_image_url(recipe_data.get("image"))
    if image_url:
        image_url = urllib.parse.urljoin(url, image_url)  # gère les URLs relatives
        downloaded = download_image_to_store(image_url)
        if downloaded:
            images.append(downloaded)

    return {
        "name": name[:100],
        "description": description[:2056],
        "ingredients": ingredients,
        "prep_time": str(prep_time) if prep_time else "",
        "cook_time": str(cook_time) if cook_time else "",
        "default_persons": default_persons or 4,
        "images": images,
        "source_url": url,
    }


def _extract_recipe_image_url(image):
    """Extrait une URL d'image utilisable depuis le champ 'image' d'une
    donnée structurée Schema.org, qui peut prendre plusieurs formes :
    chaîne, liste de chaînes, objet ImageObject, ou liste d'objets."""
    if not image:
        return None
    if isinstance(image, list):
        image = image[0] if image else None
    if isinstance(image, dict):
        return image.get("url") or image.get("@id")
    if isinstance(image, str):
        return image.strip() or None
    return None


def download_image_to_store(image_url, timeout=15):
    """Télécharge une image depuis une URL et l'enregistre dans le dossier
    images/. Retourne le nom de fichier généré, ou None en cas d'échec (page
    introuvable, ce n'est pas une image, pas de connexion...)."""
    try:
        request = urllib.request.Request(
            image_url, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            content_type = (response.headers.get_content_type() or "").lower()
            if content_type and not content_type.startswith("image/"):
                return None
            data = response.read()
    except Exception:
        return None

    if not data:
        return None

    ext_map = {
        "image/jpeg": ".jpg", "image/jpg": ".jpg", "image/png": ".png",
        "image/webp": ".webp", "image/gif": ".gif",
    }
    ext = ext_map.get(content_type)
    if not ext:
        guessed = os.path.splitext(image_url.split("?")[0])[1].lower()
        ext = guessed if guessed in (".jpg", ".jpeg", ".png", ".webp", ".gif") else ".jpg"

    new_filename = f"{uuid.uuid4().hex}{ext}"
    dest_path = os.path.join(IMAGES_DIR, new_filename)
    try:
        with open(dest_path, "wb") as f:
            f.write(data)
    except OSError:
        return None
    return new_filename


# ---------------------------------------------------------------------------
# Corbeille : les recettes supprimées y sont déplacées (avec leurs photos
# conservées) au lieu d'être effacées immédiatement, pour pouvoir les
# récupérer en cas d'erreur.
# ---------------------------------------------------------------------------

def load_trash():
    if os.path.exists(TRASH_FILE):
        try:
            with open(TRASH_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list):
                    return data
        except Exception:
            return []
    return []


def save_trash(trash):
    with open(TRASH_FILE, "w", encoding="utf-8") as f:
        json.dump(trash, f, ensure_ascii=False, indent=2)


def move_recipe_to_trash(recipe):
    """Déplace une recette vers la corbeille (ses photos ne sont PAS
    supprimées, pour pouvoir tout restaurer intact)."""
    trash = load_trash()
    trash.insert(0, {"recipe": recipe, "deleted_at": datetime.now().isoformat()})
    save_trash(trash)


# ---------------------------------------------------------------------------
# Historique des recettes récemment consultées (affiché sur la page d'accueil)
# ---------------------------------------------------------------------------

RECENT_VIEWS_MAX = 8


def load_recent_view_names():
    if os.path.exists(RECENT_VIEWS_FILE):
        try:
            with open(RECENT_VIEWS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, list):
                    return data
        except Exception:
            return []
    return []


def record_recipe_view(name):
    names = load_recent_view_names()
    names = [n for n in names if n != name]
    names.insert(0, name)
    names = names[:RECENT_VIEWS_MAX]
    with open(RECENT_VIEWS_FILE, "w", encoding="utf-8") as f:
        json.dump(names, f, ensure_ascii=False, indent=2)


# ---------------------------------------------------------------------------
# Palette et style visuel de l'application — une identité chaleureuse et
# gourmande plutôt que le gris par défaut de Windows, appliquée d'un coup à
# toute l'application (fenêtre principale et toutes les fenêtres ouvertes
# ensuite), sans avoir à retoucher chaque écran un par un. Deux palettes
# (clair/sombre) sont disponibles, basculables depuis la page d'accueil.
# ---------------------------------------------------------------------------

LIGHT_PALETTE = {
    "BG": "#FBF6EF",           # fond général, blanc cassé chaleureux
    "CARD": "#FFFFFF",         # fond des zones "carte"
    "BORDER": "#E8DCC8",       # bordures discrètes
    "TEXT": "#332B22",         # texte principal, brun très foncé
    "TEXT_MUTED": "#8A7D68",   # texte secondaire
    "ACCENT": "#D97B3F",       # orange terracotta — actions principales
    "ACCENT_DARK": "#B8622C",  # survol / pression
    "ACCENT_LIGHT": "#F3D9C4",  # fonds légers accentués
    "GREEN": "#5C8A57",        # vert sauge — validations, favoris, succès
    "ERROR": "#B4483A",        # rouge terracotta foncé — erreurs/avertissements
}

DARK_PALETTE = {
    "BG": "#211F1C",           # fond général, brun-noir doux
    "CARD": "#2C2A25",         # fond des zones "carte"
    "BORDER": "#43403A",       # bordures discrètes
    "TEXT": "#EFE8DB",         # texte principal, crème clair
    "TEXT_MUTED": "#AB9F8C",   # texte secondaire
    "ACCENT": "#E08A4F",       # orange terracotta, plus lumineux sur fond sombre
    "ACCENT_DARK": "#F0A868",  # survol / pression (plus clair que ACCENT en sombre)
    "ACCENT_LIGHT": "#4A3B2C",  # fonds légers accentués
    "GREEN": "#7CB273",        # vert sauge, plus lumineux
    "ERROR": "#E0897A",        # rouge corail, plus lumineux
}

COLOR_BG = LIGHT_PALETTE["BG"]
COLOR_CARD = LIGHT_PALETTE["CARD"]
COLOR_BORDER = LIGHT_PALETTE["BORDER"]
COLOR_TEXT = LIGHT_PALETTE["TEXT"]
COLOR_TEXT_MUTED = LIGHT_PALETTE["TEXT_MUTED"]
COLOR_ACCENT = LIGHT_PALETTE["ACCENT"]
COLOR_ACCENT_DARK = LIGHT_PALETTE["ACCENT_DARK"]
COLOR_ACCENT_LIGHT = LIGHT_PALETTE["ACCENT_LIGHT"]
COLOR_GREEN = LIGHT_PALETTE["GREEN"]
COLOR_ERROR = LIGHT_PALETTE["ERROR"]


def apply_palette(dark):
    """Met à jour les constantes de couleur globales selon le thème choisi.
    Doit être suivi d'un appel à configure_app_style() pour que les styles
    ttk et la base d'options Tk reflètent les nouvelles couleurs."""
    global COLOR_BG, COLOR_CARD, COLOR_BORDER, COLOR_TEXT, COLOR_TEXT_MUTED
    global COLOR_ACCENT, COLOR_ACCENT_DARK, COLOR_ACCENT_LIGHT, COLOR_GREEN, COLOR_ERROR
    palette = DARK_PALETTE if dark else LIGHT_PALETTE
    COLOR_BG = palette["BG"]
    COLOR_CARD = palette["CARD"]
    COLOR_BORDER = palette["BORDER"]
    COLOR_TEXT = palette["TEXT"]
    COLOR_TEXT_MUTED = palette["TEXT_MUTED"]
    COLOR_ACCENT = palette["ACCENT"]
    COLOR_ACCENT_DARK = palette["ACCENT_DARK"]
    COLOR_ACCENT_LIGHT = palette["ACCENT_LIGHT"]
    COLOR_GREEN = palette["GREEN"]
    COLOR_ERROR = palette["ERROR"]


def get_dark_mode_preference():
    return bool(load_settings().get("dark_mode", False))


def set_dark_mode_preference(value):
    settings = load_settings()
    settings["dark_mode"] = bool(value)
    save_settings(settings)


def configure_app_style(root):
    """Configure l'apparence de toute l'application : couleurs ttk (boutons,
    champs, onglets...) ainsi que les widgets Tkinter bruts (Listbox, Text,
    Canvas...) via la base d'options de Tk, qui s'applique automatiquement à
    tous les widgets créés ensuite dans n'importe quelle fenêtre."""
    style = ttk.Style(root)
    try:
        style.theme_use("clam")
    except tk.TclError:
        pass

    base_font = ("Segoe UI", 10)

    style.configure(".", background=COLOR_BG, foreground=COLOR_TEXT, font=base_font)
    style.configure("TFrame", background=COLOR_BG)
    style.configure("TLabel", background=COLOR_BG, foreground=COLOR_TEXT)
    style.configure("TCheckbutton", background=COLOR_BG, foreground=COLOR_TEXT)
    style.map("TCheckbutton", background=[("active", COLOR_BG)])

    style.configure("TButton", background=COLOR_ACCENT, foreground="white",
                     font=base_font, padding=(12, 7), borderwidth=0, relief="flat")
    style.map("TButton",
              background=[("active", COLOR_ACCENT_DARK), ("disabled", "#D8CBB8")],
              foreground=[("disabled", "#F4EEE3")])

    style.configure("TEntry", fieldbackground=COLOR_CARD, foreground=COLOR_TEXT,
                     bordercolor=COLOR_BORDER, lightcolor=COLOR_BORDER, darkcolor=COLOR_BORDER,
                     padding=5)
    style.configure("TCombobox", fieldbackground=COLOR_CARD, foreground=COLOR_TEXT,
                     bordercolor=COLOR_BORDER, arrowcolor=COLOR_ACCENT_DARK, padding=5)
    style.map("TCombobox", fieldbackground=[("readonly", COLOR_CARD)])

    style.configure("TLabelframe", background=COLOR_BG, bordercolor=COLOR_BORDER)
    style.configure("TLabelframe.Label", background=COLOR_BG, foreground=COLOR_ACCENT_DARK,
                     font=("Segoe UI", 10, "bold"))

    style.configure("TNotebook", background=COLOR_BG, bordercolor=COLOR_BORDER)
    style.configure("TNotebook.Tab", background=COLOR_ACCENT_LIGHT, foreground=COLOR_TEXT,
                     padding=(12, 6), font=base_font)
    style.map("TNotebook.Tab",
              background=[("selected", COLOR_ACCENT)],
              foreground=[("selected", "white")])

    style.configure("TScrollbar", background=COLOR_ACCENT, troughcolor=COLOR_BG,
                     bordercolor=COLOR_BG, arrowcolor="white")
    style.map("TScrollbar", background=[("active", COLOR_ACCENT_DARK)])

    style.configure("TSeparator", background=COLOR_BORDER)

    style.configure("TPanedwindow", background=COLOR_BG)

    # Boutons secondaires plus discrets (utilisés pour des actions annexes) :
    # à activer au cas par cas avec style="Secondary.TButton" si besoin plus tard.
    style.configure("Secondary.TButton", background=COLOR_CARD, foreground=COLOR_ACCENT_DARK,
                     padding=(12, 7))
    style.map("Secondary.TButton", background=[("active", COLOR_ACCENT_LIGHT)])

    # Variantes "carte" (fond blanc) pour les encadrés mis en valeur sur la
    # page d'accueil, afin que les widgets ttk placés dedans (Label, Button)
    # aient le même fond que la carte plutôt que le fond général de la page.
    style.configure("Card.TFrame", background=COLOR_CARD)
    style.configure("Card.TLabel", background=COLOR_CARD, foreground=COLOR_TEXT)

    # ---- Widgets Tkinter bruts (non gérés par ttk) : Listbox, Text, Canvas,
    # Button (celles en tk.Button, ex. mode cuisine), via la base d'options.
    # S'applique à toute fenêtre créée dans l'application, sans exception. ----
    root.option_add("*Font", "{Segoe UI} 10")
    root.option_add("*Background", COLOR_BG)
    root.option_add("*Foreground", COLOR_TEXT)

    root.option_add("*Listbox.Background", COLOR_CARD)
    root.option_add("*Listbox.Foreground", COLOR_TEXT)
    root.option_add("*Listbox.selectBackground", COLOR_ACCENT)
    root.option_add("*Listbox.selectForeground", "white")
    root.option_add("*Listbox.borderWidth", 1)
    root.option_add("*Listbox.relief", "solid")
    root.option_add("*Listbox.highlightThickness", 0)

    root.option_add("*Text.Background", COLOR_CARD)
    root.option_add("*Text.Foreground", COLOR_TEXT)
    root.option_add("*Text.borderWidth", 1)
    root.option_add("*Text.relief", "solid")
    root.option_add("*Text.highlightThickness", 0)

    root.option_add("*Canvas.Background", COLOR_BG)
    root.option_add("*Canvas.highlightThickness", 0)

    root.option_add("*Button.Background", COLOR_ACCENT)
    root.option_add("*Button.Foreground", "white")
    root.option_add("*Button.activeBackground", COLOR_ACCENT_DARK)
    root.option_add("*Button.activeForeground", "white")
    root.option_add("*Button.relief", "flat")
    root.option_add("*Button.borderWidth", 0)
    root.option_add("*Button.padX", 10)
    root.option_add("*Button.padY", 5)

    root.option_add("*Entry.Background", COLOR_CARD)
    root.option_add("*Entry.Foreground", COLOR_TEXT)
    root.option_add("*Entry.relief", "solid")
    root.option_add("*Entry.borderWidth", 1)

    return style


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Mon Livre de Recettes")
        # La hauteur de la fenêtre correspond à la hauteur de l'écran sur
        # lequel l'application est lancée, pour profiter de tout l'espace
        # vertical disponible dès le démarrage.
        screen_height = self.winfo_screenheight()
        self.geometry(f"560x{screen_height}+40+0")
        self.minsize(480, 500)
        self.resizable(True, True)

        self.dark_mode = get_dark_mode_preference()
        apply_palette(self.dark_mode)
        configure_app_style(self)
        self.configure(background=COLOR_BG)

        self.recipes = load_recipes()
        self.ingredient_names = sync_ingredients_from_recipes()
        self.shopping_selection = {}  # sélection en cours pour la liste de courses (nom -> personnes)
        self.timers_window = None  # fenêtre unique des minuteurs, créée à la demande

        # Sauvegarde automatique périodique (silencieuse, ne bloque jamais le démarrage)
        try:
            maybe_create_auto_backup()
        except Exception:
            pass

        self._build_home_ui()

    def toggle_dark_mode(self):
        """Bascule entre thème clair et sombre : met à jour la palette, les
        styles ttk (effet immédiat sur toutes les fenêtres déjà ouvertes),
        puis reconstruit entièrement la page d'accueil pour que ses widgets
        Tkinter bruts (bannière, cartes...) reflètent aussi les nouvelles
        couleurs. Les fenêtres secondaires déjà ouvertes devront être
        refermées puis rouvertes pour refléter pleinement le nouveau thème."""
        self.dark_mode = not self.dark_mode
        set_dark_mode_preference(self.dark_mode)
        apply_palette(self.dark_mode)
        configure_app_style(self)
        self.configure(background=COLOR_BG)
        try:
            self.unbind_all("<MouseWheel>")
        except tk.TclError:
            pass
        # Ne détruit que les widgets propres à la page d'accueil : les
        # fenêtres secondaires (Toplevel) déjà ouvertes — comme les
        # minuteurs en cours — ne doivent surtout pas être affectées.
        for child in self.winfo_children():
            if not isinstance(child, tk.Toplevel):
                child.destroy()
        self._build_home_ui()

    def _build_home_ui(self):
        # ---- Barre supérieure fixe (hors zone de défilement) avec le
        # bouton de bascule de thème, toujours visible en haut à droite. ----
        top_bar = tk.Frame(self, background=COLOR_BG)
        top_bar.pack(fill="x")
        toggle_text = "☀️ Thème clair" if self.dark_mode else "🌙 Thème sombre"
        ttk.Button(top_bar, text=toggle_text, style="Secondary.TButton",
                   command=self.toggle_dark_mode).pack(side="right", padx=10, pady=8)

        # ---- Conteneur scrollable pour toute la page d'accueil : ainsi, quel
        # que soit le nombre de boutons ou la taille de la fenêtre, rien ne
        # peut jamais être coupé ou invisible. Le pied de page reste fixé en
        # bas, en dehors de la zone qui défile. ----
        outer = ttk.Frame(self)
        outer.pack(fill="both", expand=True)
        canvas = tk.Canvas(outer, highlightthickness=0)
        scrollbar = ttk.Scrollbar(outer, orient="vertical", command=canvas.yview)
        content = ttk.Frame(canvas)
        content.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=content, anchor="n")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        def _bind_mousewheel(event):
            canvas.bind_all("<MouseWheel>", _on_mousewheel)

        def _unbind_mousewheel(event):
            canvas.unbind_all("<MouseWheel>")

        canvas.bind("<Enter>", _bind_mousewheel)
        canvas.bind("<Leave>", _unbind_mousewheel)

        banner = tk.Frame(content, background=COLOR_ACCENT)
        banner.pack(fill="x")
        tk.Label(banner, text="👨‍🍳 Mon Livre de Recettes", font=("Segoe UI", 21, "bold"),
                 background=COLOR_ACCENT, foreground="white").pack(pady=(18, 2))
        tk.Label(banner, text="Toutes vos recettes, à portée de main",
                 font=("Segoe UI", 10), background=COLOR_ACCENT,
                 foreground=COLOR_ACCENT_LIGHT).pack(pady=(0, 16))

        # ---- Filtres rapides : accès direct à une liste déjà filtrée ----
        quick_filters_frame = ttk.Frame(content)
        quick_filters_frame.pack(padx=20, pady=(15, 0), fill="x")
        ttk.Button(quick_filters_frame, text="⭐ Favoris", style="Secondary.TButton",
                   command=lambda: self.open_manage_recipes(quick_filter="favoris")).pack(
            side="left", expand=True, fill="x", padx=(0, 4))
        ttk.Button(quick_filters_frame, text="⏱️ Rapide (≤ 30 min)", style="Secondary.TButton",
                   command=lambda: self.open_manage_recipes(quick_filter="rapide")).pack(
            side="left", expand=True, fill="x", padx=4)
        ttk.Button(quick_filters_frame, text="🥗 Végétarien", style="Secondary.TButton",
                   command=lambda: self.open_manage_recipes(quick_filter="vegetarien")).pack(
            side="left", expand=True, fill="x", padx=(4, 0))

        grid_frame = ttk.Frame(content)
        grid_frame.pack(padx=20, pady=(15, 0), fill="x")
        grid_frame.columnconfigure(0, weight=1)

        buttons = [
            ("➕  Ajouter une recette", self.open_add_recipe),
            ("🌐  Importer une recette depuis un lien", self.open_import_from_url),
            ("📷  Importer une recette depuis une photo", self.open_import_from_photo),
            ("🧾  Voir toutes les recettes (liste de courses)", self.open_all_recipes),
            ("🍽️  Voir une recette précise", self.open_one_recipe),
            ("✏️  Modifier / Supprimer une recette", self.open_manage_recipes),
            ("⚖️  Comparer deux recettes", self.open_compare_recipes),
            ("🥕  Gérer les ingrédients", self.open_manage_ingredients),
            ("🔎  Recherche par ingrédient", self.open_ingredient_search),
            ("🧊  Que puis-je cuisiner ?", self.open_what_can_i_cook),
            ("📅  Planning de la semaine", self.open_weekly_plan),
            ("📋  Mes menus", self.open_menus),
            ("📊  Statistiques", self.open_statistics),
            ("📖  Exporter le livre de recettes", self.open_cookbook_export),
            ("💾  Importer / Exporter les données", self.open_import_export),
            ("🗑️  Corbeille", self.open_trash),
        ]
        for i, (text, command) in enumerate(buttons):
            ttk.Button(grid_frame, text=text, command=command).grid(
                row=i, column=0, padx=4, pady=4, sticky="ew"
            )

        # ---- Repas du jour ----
        today_card = tk.Frame(content, background=COLOR_CARD, highlightbackground=COLOR_BORDER,
                               highlightthickness=1)
        today_card.pack(padx=15, pady=(20, 0), fill="x")
        ttk.Label(today_card, text="📅 Aujourd'hui", font=("Segoe UI", 12, "bold"),
                  style="Card.TLabel", foreground=COLOR_ACCENT_DARK).pack(anchor="w", padx=12, pady=(10, 4))
        self.today_frame = ttk.Frame(today_card, style="Card.TFrame")
        self.today_frame.pack(padx=12, pady=(0, 12), fill="x")
        self._refresh_today_meals()

        # ---- Récemment consultées ----
        recent_card = tk.Frame(content, background=COLOR_CARD, highlightbackground=COLOR_BORDER,
                                highlightthickness=1)
        recent_card.pack(padx=15, pady=(12, 0), fill="x")
        ttk.Label(recent_card, text="🕘 Récemment consultées", font=("Segoe UI", 12, "bold"),
                  style="Card.TLabel", foreground=COLOR_ACCENT_DARK).pack(anchor="w", padx=12, pady=(10, 4))
        recent_frame = ttk.Frame(recent_card, style="Card.TFrame")
        recent_frame.pack(padx=12, pady=(0, 12), fill="x")
        self.recent_listbox = tk.Listbox(recent_frame, height=5)
        self.recent_listbox.pack(side="left", fill="x", expand=True)
        self.recent_listbox.bind("<Double-Button-1>", lambda e: self.open_recent_selected())
        ttk.Button(recent_frame, text="👁 Ouvrir", command=self.open_recent_selected).pack(side="left", padx=(8, 0))
        self._refresh_recent_views()

        warnings = []
        if not PIL_AVAILABLE:
            warnings.append("Pillow non installé : les photos ne s'afficheront pas (pip install pillow)")
        if not REPORTLAB_AVAILABLE:
            warnings.append("reportlab non installé : export PDF indisponible (pip install reportlab)")
        if not OPENPYXL_AVAILABLE:
            warnings.append("openpyxl non installé : export Excel indisponible (pip install openpyxl)")
        if not QRCODE_AVAILABLE:
            warnings.append("qrcode non installé : export QR code indisponible (pip install qrcode)")
        if not PYTESSERACT_AVAILABLE:
            warnings.append("pytesseract non installé : import depuis une photo indisponible (pip install pytesseract, + Tesseract OCR)")
        if warnings:
            ttk.Label(content, text="\n".join(warnings), foreground=COLOR_ERROR, font=("Segoe UI", 8),
                      justify="center").pack(pady=(10, 10))
        else:
            ttk.Label(content, text="", font=("Segoe UI", 4)).pack(pady=(5, 5))

        self.footer = ttk.Label(self, text=f"{len(self.recipes)} recette(s) enregistrée(s)",
                                 font=("Segoe UI", 9))
        self.footer.pack(side="bottom", pady=15)

    def refresh_recipes(self):
        self.recipes = load_recipes()
        self.footer.config(text=f"{len(self.recipes)} recette(s) enregistrée(s)")

    def _refresh_today_meals(self):
        for child in self.today_frame.winfo_children():
            child.destroy()

        today_name = WEEKDAYS[datetime.now().weekday()]  # 0=Lundi ... 6=Dimanche
        plan = load_weekly_plan()
        day_data = plan.get(today_name) or {}

        entries = []
        for slot in WeeklyPlanWindow.MEAL_SLOTS:
            slot_data = day_data.get(slot)
            if slot_data and slot_data.get("recipe_name"):
                entries.append((slot, slot_data["recipe_name"], slot_data.get("persons", 1)))

        if not entries:
            ttk.Label(
                self.today_frame,
                text=f"Rien de planifié pour {today_name.lower()}. "
                     "Remplissez le « 📅 Planning de la semaine » pour le voir ici.",
                font=("Segoe UI", 9), foreground=COLOR_TEXT_MUTED, wraplength=560, justify="left",
                style="Card.TLabel"
            ).pack(anchor="w", pady=3)
            return

        ttk.Label(self.today_frame, text=today_name, font=("Segoe UI", 9, "bold"),
                  foreground=COLOR_TEXT_MUTED, style="Card.TLabel").pack(anchor="w")
        for slot, recipe_name, persons in entries:
            row = ttk.Frame(self.today_frame, style="Card.TFrame")
            row.pack(fill="x", pady=1)
            ttk.Label(row, text=f"{slot} :", width=18, anchor="w", style="Card.TLabel").pack(side="left")
            ttk.Label(row, text=f"{recipe_name} ({persons} pers.)", anchor="w",
                      style="Card.TLabel").pack(side="left")
            ttk.Button(row, text="👁", width=3,
                       command=lambda n=recipe_name: self._open_today_recipe(n)).pack(side="left", padx=5)

    def _open_today_recipe(self, recipe_name):
        win = OneRecipeWindow(self, initial_recipe_name=recipe_name)
        self.wait_window(win)
        self._refresh_recent_views()

    def _refresh_recent_views(self):
        self.recent_listbox.delete(0, tk.END)
        names = load_recent_view_names()
        self._recent_recipes = []
        for name in names:
            recipe = find_recipe_by_name(self.recipes, name)
            if recipe is None:
                continue  # la recette a été supprimée depuis
            self._recent_recipes.append(recipe)
            self.recent_listbox.insert(tk.END, format_recipe_list_label(recipe))
        if not self._recent_recipes:
            self.recent_listbox.insert(tk.END, "Aucune recette consultée pour le moment.")

    def open_recent_selected(self):
        sel = self.recent_listbox.curselection()
        if not sel or not getattr(self, "_recent_recipes", None):
            return
        recipe = self._recent_recipes[sel[0]]
        win = OneRecipeWindow(self, initial_recipe_name=recipe["name"])
        self.wait_window(win)
        self._refresh_recent_views()

    def refresh_ingredients(self):
        self.ingredient_names = load_ingredients()

    # ---------- Ajouter une recette ----------
    def open_add_recipe(self):
        RecipeFormWindow(self, recipe_index=None)

    # ---------- Voir toutes les recettes ----------
    def open_all_recipes(self):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        AllRecipesWindow(self)

    # ---------- Voir une recette précise ----------
    def open_one_recipe(self):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        win = OneRecipeWindow(self)
        self.wait_window(win)
        self._refresh_recent_views()

    # ---------- Modifier / Supprimer une recette ----------
    def open_manage_recipes(self, quick_filter=None):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        win = ManageRecipesWindow(self, quick_filter=quick_filter)
        self.wait_window(win)
        self._refresh_recent_views()

    # ---------- Gérer les ingrédients ----------
    def open_manage_ingredients(self):
        ManageIngredientsWindow(self)

    # ---------- Corbeille ----------
    def open_trash(self):
        win = TrashWindow(self)
        self.wait_window(win)
        self._refresh_recent_views()

    # ---------- Recherche par ingrédient ----------
    def open_ingredient_search(self):
        IngredientSearchWindow(self)

    # ---------- Comparer deux recettes ----------
    def open_compare_recipes(self):
        if len(self.recipes) < 2:
            messagebox.showinfo("Info", "Il faut au moins 2 recettes enregistrées pour pouvoir comparer.")
            return
        CompareRecipesWindow(self)

    # ---------- Importer une recette depuis un lien ----------
    def open_import_from_url(self):
        ImportFromUrlWindow(self)

    # ---------- Importer une recette depuis une photo ----------
    def open_import_from_photo(self):
        ImportFromPhotoWindow(self)

    # ---------- Importer / Exporter les données ----------
    def open_import_export(self):
        ImportExportWindow(self)

    # ---------- Que puis-je cuisiner ? ----------
    def open_what_can_i_cook(self):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        WhatCanICookWindow(self)

    # ---------- Planning de la semaine ----------
    def open_weekly_plan(self):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        win = WeeklyPlanWindow(self)
        self.wait_window(win)
        self._refresh_today_meals()

    # ---------- Mes menus ----------
    def open_menus(self):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        MenuManagerWindow(self)

    # ---------- Statistiques ----------
    def open_statistics(self):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        StatisticsWindow(self)

    # ---------- Exporter le livre de recettes ----------
    def open_cookbook_export(self):
        if not self.recipes:
            messagebox.showinfo("Info", "Aucune recette enregistrée pour le moment.")
            return
        CookbookExportWindow(self)


class RecipeFormWindow(tk.Toplevel):
    """Fenêtre d'ajout OU de modification d'une recette (nom, catégorie, photo,
    description, ingrédients pour 1 personne). Si recipe_index est fourni, la
    fenêtre s'ouvre en mode modification, pré-remplie avec la recette
    existante."""

    CATEGORY_OPTIONS = ["Entrée", "Plat", "Dessert", "Apéro", "Boisson", "Sauce", "Autre"]
    DIFFICULTY_OPTIONS = ["Facile", "Moyen", "Difficile"]
    MAX_DESC_LEN = 2056
    MAX_NOTES_LEN = 500

    def __init__(self, app, recipe_index=None, prefill=None):
        super().__init__(app)
        self.app = app
        self.recipe_index = recipe_index
        self.editing = recipe_index is not None
        self.existing_recipe = app.recipes[recipe_index] if self.editing else None
        self.prefill = prefill if not self.editing else None
        self.ingredient_names = load_ingredients()

        # Galerie de photos : chaque élément est ("existing", nom_de_fichier)
        # pour une photo déjà enregistrée, ou ("new", chemin_source) pour une
        # photo qui vient d'être choisie et sera copiée lors de l'enregistrement.
        self.gallery_items = []
        if self.editing:
            self.gallery_items = [("existing", fname) for fname in get_recipe_images(self.existing_recipe)]
        elif self.prefill:
            # La photo a déjà été téléchargée et enregistrée dans images/ par
            # fetch_recipe_from_url ; on la référence donc comme "existing".
            self.gallery_items = [("existing", fname) for fname in self.prefill.get("images", [])]
        self._gallery_thumb_refs = []  # garder une référence pour éviter le garbage collector

        self.title("Modifier la recette" if self.editing else "Ajouter une recette")
        self.geometry("560x760")
        self.minsize(480, 400)
        self.grab_set()

        # ---- Conteneur scrollable pour tout le formulaire ----
        # Tout le contenu (nom, catégorie, photo, description, ingrédients,
        # boutons) est placé dans ce conteneur qui défile d'un bloc : ainsi,
        # les boutons "+ Ajouter un ingrédient" / "Enregistrer" restent
        # toujours juste après le dernier ingrédient, où qu'il soit.
        outer = ttk.Frame(self)
        outer.pack(fill="both", expand=True)
        self.canvas = tk.Canvas(outer, highlightthickness=0)
        scrollbar = ttk.Scrollbar(outer, orient="vertical", command=self.canvas.yview)
        self.content_frame = ttk.Frame(self.canvas)
        self.content_frame.bind(
            "<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all"))
        )
        self.canvas.create_window((0, 0), window=self.content_frame, anchor="nw")
        self.canvas.configure(yscrollcommand=scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        ttk.Label(self.content_frame, text="Nom de la recette :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(15, 5))
        self.name_entry = ttk.Entry(self.content_frame, width=42)
        self.name_entry.pack()
        if self.editing:
            self.name_entry.insert(0, self.existing_recipe["name"])
        elif self.prefill:
            self.name_entry.insert(0, self.prefill.get("name", ""))

        self.favorite_var = tk.BooleanVar(
            value=self.existing_recipe.get("favorite", False) if self.editing else False
        )
        ttk.Checkbutton(self.content_frame, text="⭐ Marquer comme recette favorite",
                         variable=self.favorite_var).pack(pady=(8, 0))

        # ---- Note personnelle (1 à 5 étoiles, cliquables) ----
        self.rating_value = self.existing_recipe.get("rating", 0) if self.editing else 0
        rating_frame = ttk.Frame(self.content_frame)
        rating_frame.pack(pady=(8, 0))
        ttk.Label(rating_frame, text="Ma note :").pack(side="left", padx=(0, 5))
        self.rating_star_labels = []
        for i in range(1, 6):
            lbl = ttk.Label(rating_frame, text="☆", font=("Segoe UI", 14), cursor="hand2")
            lbl.pack(side="left")
            lbl.bind("<Button-1>", lambda e, i=i: self._set_rating(i))
            self.rating_star_labels.append(lbl)
        self._refresh_rating_stars()

        ttk.Label(self.content_frame, text="Catégorie :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(15, 5))
        self.category_combo = ttk.Combobox(self.content_frame, values=self.CATEGORY_OPTIONS,
                                            state="readonly", width=20)
        self.category_combo.set(
            self.existing_recipe.get("category", "Plat") if self.editing else "Plat"
        )
        self.category_combo.pack()

        # ---- Temps de préparation / cuisson / difficulté ----
        times_frame = ttk.Frame(self.content_frame)
        times_frame.pack(pady=(15, 5))
        ttk.Label(times_frame, text="Préparation (min) :").grid(row=0, column=0, padx=3, sticky="e")
        self.prep_time_entry = ttk.Entry(times_frame, width=6)
        self.prep_time_entry.grid(row=0, column=1, padx=3)
        ttk.Label(times_frame, text="Cuisson (min) :").grid(row=0, column=2, padx=3, sticky="e")
        self.cook_time_entry = ttk.Entry(times_frame, width=6)
        self.cook_time_entry.grid(row=0, column=3, padx=3)
        if self.editing:
            self.prep_time_entry.insert(0, str(self.existing_recipe.get("prep_time", "") or ""))
            self.cook_time_entry.insert(0, str(self.existing_recipe.get("cook_time", "") or ""))
        elif self.prefill:
            self.prep_time_entry.insert(0, str(self.prefill.get("prep_time", "") or ""))
            self.cook_time_entry.insert(0, str(self.prefill.get("cook_time", "") or ""))

        difficulty_frame = ttk.Frame(self.content_frame)
        difficulty_frame.pack(pady=(5, 5))
        ttk.Label(difficulty_frame, text="Difficulté :").pack(side="left", padx=3)
        self.difficulty_combo = ttk.Combobox(difficulty_frame, values=self.DIFFICULTY_OPTIONS,
                                              state="readonly", width=15)
        self.difficulty_combo.set(
            self.existing_recipe.get("difficulty", "Facile") if self.editing else "Facile"
        )
        self.difficulty_combo.pack(side="left", padx=3)
        ttk.Label(difficulty_frame, text="   Personnes par défaut :").pack(side="left", padx=(10, 3))
        self.default_persons_entry = ttk.Entry(difficulty_frame, width=5)
        if self.editing:
            default_persons_value = str(self.existing_recipe.get("default_persons", 4))
        elif self.prefill:
            default_persons_value = str(self.prefill.get("default_persons", 4))
        else:
            default_persons_value = "4"
        self.default_persons_entry.insert(0, default_persons_value)
        self.default_persons_entry.pack(side="left")

        # ---- Étiquettes libres ----
        ttk.Label(self.content_frame, text="Étiquettes (séparées par des virgules) :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(15, 5))
        self.tags_entry = ttk.Entry(self.content_frame, width=42)
        self.tags_entry.pack()
        if self.editing and self.existing_recipe.get("tags"):
            self.tags_entry.insert(0, ", ".join(self.existing_recipe["tags"]))
        ttk.Label(self.content_frame, text="ex. végétarien, sans gluten, rapide, économique",
                  font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED).pack()

        # ---- Allergènes ----
        allergens_header = ttk.Frame(self.content_frame)
        allergens_header.pack(fill="x", padx=10, pady=(15, 5))
        ttk.Label(allergens_header, text="Allergènes présents :",
                  font=("Segoe UI", 11, "bold")).pack(side="left")
        ttk.Button(allergens_header, text="🔍 Détecter automatiquement",
                   command=self.detect_allergens_from_ingredients).pack(side="right")
        allergens_frame = ttk.Frame(self.content_frame)
        allergens_frame.pack()
        if self.editing:
            existing_allergens = set(self.existing_recipe.get("allergens", []))
        elif self.prefill and self.prefill.get("ingredients"):
            # Recette importée depuis un lien : détection automatique dès l'ouverture
            existing_allergens = set(compute_recipe_allergens(self.prefill["ingredients"]))
        else:
            existing_allergens = set()
        self.allergen_vars = {}
        for i, allergen in enumerate(ALLERGENS):
            var = tk.BooleanVar(value=allergen in existing_allergens)
            self.allergen_vars[allergen] = var
            ttk.Checkbutton(allergens_frame, text=allergen, variable=var).grid(
                row=i // 3, column=i % 3, sticky="w", padx=8, pady=2
            )
        # Sert à ne jamais décocher un allergène que l'utilisateur aurait
        # coché lui-même sans lien avec un ingrédient détecté : on ne
        # décoche automatiquement que ce que la détection a elle-même coché.
        if self.editing:
            self._auto_detected_allergens = set(
                compute_recipe_allergens(self.existing_recipe.get("ingredients", []))
            )
        elif self.prefill and self.prefill.get("ingredients"):
            self._auto_detected_allergens = set(compute_recipe_allergens(self.prefill["ingredients"]))
        else:
            self._auto_detected_allergens = set()
        ttk.Label(
            self.content_frame,
            text="La détection automatique se base sur les ingrédients de la\n"
                 "recette déjà saisis ci-dessous : elle coche et décoche les\n"
                 "cases en fonction, sans jamais toucher à celles que vous\n"
                 "auriez cochées vous-même sans lien avec un ingrédient détecté.",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center"
        ).pack(pady=(0, 5))

        # ---- Photos (galerie) ----
        ttk.Label(self.content_frame, text="Photos :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(15, 5))
        gallery_outer = ttk.Frame(self.content_frame)
        gallery_outer.pack(fill="x", padx=10)
        self.gallery_canvas = tk.Canvas(gallery_outer, height=130, highlightthickness=0)
        gallery_scrollbar = ttk.Scrollbar(gallery_outer, orient="horizontal",
                                           command=self.gallery_canvas.xview)
        self.gallery_frame = ttk.Frame(self.gallery_canvas)
        self.gallery_frame.bind(
            "<Configure>", lambda e: self.gallery_canvas.configure(scrollregion=self.gallery_canvas.bbox("all"))
        )
        self.gallery_canvas.create_window((0, 0), window=self.gallery_frame, anchor="nw")
        self.gallery_canvas.configure(xscrollcommand=gallery_scrollbar.set)
        self.gallery_canvas.pack(fill="x")
        gallery_scrollbar.pack(fill="x")
        ttk.Button(self.content_frame, text="📷 Ajouter une photo",
                   command=self.choose_images).pack(pady=5)
        self._refresh_gallery()

        # ---- Description ----
        ttk.Label(self.content_frame, text="Description (informations, étapes, astuces...) :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(15, 5))
        desc_frame = ttk.Frame(self.content_frame)
        desc_frame.pack(padx=10)
        self.description_text = tk.Text(desc_frame, height=8, width=48, wrap="word")
        self.description_text.pack()
        if self.editing and self.existing_recipe.get("description"):
            self.description_text.insert("1.0", self.existing_recipe["description"])
        elif self.prefill and self.prefill.get("description"):
            self.description_text.insert("1.0", self.prefill["description"])
        self.desc_counter_label = ttk.Label(self.content_frame, text="", font=("Segoe UI", 8),
                                             foreground=COLOR_TEXT_MUTED)
        self.desc_counter_label.pack(pady=(2, 0))
        self.description_text.bind("<<Modified>>", self._on_description_modified)
        self._on_description_modified()

        # ---- Notes personnelles ----
        ttk.Label(self.content_frame, text="Notes personnelles (avis, ajustements pour la prochaine fois...) :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(15, 5))
        notes_frame = ttk.Frame(self.content_frame)
        notes_frame.pack(padx=10)
        self.notes_text = tk.Text(notes_frame, height=4, width=48, wrap="word")
        self.notes_text.pack()
        if self.editing and self.existing_recipe.get("personal_notes"):
            self.notes_text.insert("1.0", self.existing_recipe["personal_notes"])
        self.notes_counter_label = ttk.Label(self.content_frame, text="", font=("Segoe UI", 8),
                                              foreground=COLOR_TEXT_MUTED)
        self.notes_counter_label.pack(pady=(2, 0))
        self.notes_text.bind("<<Modified>>", self._on_notes_modified)
        self._on_notes_modified()

        # ---- Ingrédients ----
        ing_header_frame = ttk.Frame(self.content_frame)
        ing_header_frame.pack(fill="x", padx=10, pady=(15, 5))
        ttk.Label(ing_header_frame, text="Ingrédients (quantité pour 1 personne) :",
                  font=("Segoe UI", 11, "bold")).pack(side="left")
        ttk.Button(ing_header_frame, text="🥕 Nouvel ingrédient",
                   command=self.add_new_ingredient_global).pack(side="right")

        if not self.ingredient_names:
            ttk.Label(self.content_frame,
                      text="Aucun ingrédient enregistré. Cliquez sur « 🥕 Nouvel ingrédient »\npour en créer un premier.",
                      font=("Segoe UI", 8), foreground=COLOR_ERROR, justify="center").pack()

        self.rows_frame = ttk.Frame(self.content_frame)
        self.rows_frame.pack(fill="x", padx=10)

        header = ttk.Frame(self.rows_frame)
        header.pack(fill="x", pady=2)
        ttk.Label(header, text="Ingrédient", width=17, font=("Segoe UI", 9, "bold")).grid(row=0, column=0)
        ttk.Label(header, text="Quantité", width=9, font=("Segoe UI", 9, "bold")).grid(row=0, column=1)
        ttk.Label(header, text="Unité", width=15, font=("Segoe UI", 9, "bold")).grid(row=0, column=2)
        ttk.Label(header, text="(si autre)", width=10, font=("Segoe UI", 9, "bold")).grid(row=0, column=3)

        self.ingredient_rows = []
        if self.editing and self.existing_recipe["ingredients"]:
            for ing in self.existing_recipe["ingredients"]:
                self.add_ingredient_row(ing["name"], ing["quantity"], ing["unit"])
        elif self.prefill and self.prefill.get("ingredients"):
            for ing in self.prefill["ingredients"]:
                self.add_ingredient_row(ing["name"], ing["quantity"], ing["unit"])
        else:
            self.add_ingredient_row()

        # Ces deux éléments sont recréés/déplacés à chaque ajout de ligne afin
        # de toujours rester juste en dessous du dernier ingrédient.
        self.add_ingredient_button = ttk.Button(
            self.rows_frame, text="+ Ajouter un ingrédient", command=lambda: self.add_ingredient_row()
        )
        self.add_ingredient_button.pack(pady=10)

        self.bottom_actions_frame = ttk.Frame(self.rows_frame)
        self.bottom_actions_frame.pack(pady=(0, 20))
        ttk.Button(self.bottom_actions_frame, text="Enregistrer",
                   command=self.save_recipe).grid(row=0, column=0, padx=5)
        if self.editing:
            ttk.Button(self.bottom_actions_frame, text="Supprimer cette recette",
                       command=self.delete_recipe).grid(row=0, column=1, padx=5)

    def _on_description_modified(self, event=None):
        self.description_text.edit_modified(False)
        content = self.description_text.get("1.0", "end-1c")
        if len(content) > self.MAX_DESC_LEN:
            content = content[: self.MAX_DESC_LEN]
            self.description_text.delete("1.0", "end")
            self.description_text.insert("1.0", content)
            self.description_text.edit_modified(False)
        self.desc_counter_label.config(text=f"{len(content)} / {self.MAX_DESC_LEN} caractères")

    def _on_notes_modified(self, event=None):
        self.notes_text.edit_modified(False)
        content = self.notes_text.get("1.0", "end-1c")
        if len(content) > self.MAX_NOTES_LEN:
            content = content[: self.MAX_NOTES_LEN]
            self.notes_text.delete("1.0", "end")
            self.notes_text.insert("1.0", content)
            self.notes_text.edit_modified(False)
        self.notes_counter_label.config(text=f"{len(content)} / {self.MAX_NOTES_LEN} caractères")

    def _set_rating(self, value):
        self.rating_value = 0 if self.rating_value == value else value
        self._refresh_rating_stars()

    def _refresh_rating_stars(self):
        for i, lbl in enumerate(self.rating_star_labels, start=1):
            lbl.config(text="★" if i <= self.rating_value else "☆")

    def detect_allergens_from_ingredients(self):
        current_ingredients = []
        for name_e, qty_e, unit_e, custom_e in self.ingredient_rows:
            ing_name = name_e.get().strip()
            if ing_name:
                current_ingredients.append({"name": ing_name})
        if not current_ingredients:
            messagebox.showinfo("Info", "Ajoutez d'abord des ingrédients à la recette.")
            return

        before = {a for a, var in self.allergen_vars.items() if var.get()}
        self._sync_allergens_from_ingredients()
        after = {a for a, var in self.allergen_vars.items() if var.get()}

        added = sorted(after - before)
        removed = sorted(before - after)
        if added or removed:
            parts = []
            if added:
                parts.append(f"ajouté(s) : {', '.join(added)}")
            if removed:
                parts.append(f"retiré(s) : {', '.join(removed)}")
            messagebox.showinfo("Allergènes mis à jour", "Allergène(s) " + " ; ".join(parts) + ".")
        else:
            messagebox.showinfo("Info", "Aucun changement : les allergènes cochés correspondent déjà aux ingrédients.")

    def choose_images(self):
        paths = filedialog.askopenfilenames(
            title="Choisir une ou plusieurs photos",
            filetypes=[("Images", "*.png *.jpg *.jpeg *.gif *.bmp *.webp")]
        )
        for path in paths:
            self.gallery_items.append(("new", path))
        if paths:
            self._refresh_gallery()

    def _remove_gallery_item(self, index):
        if 0 <= index < len(self.gallery_items):
            self.gallery_items.pop(index)
            self._refresh_gallery()

    def _refresh_gallery(self):
        for child in self.gallery_frame.winfo_children():
            child.destroy()
        self._gallery_thumb_refs = []

        if not self.gallery_items:
            ttk.Label(self.gallery_frame, text="(aucune photo)").pack(side="left", padx=10, pady=10)
            return

        for idx, (kind, ref) in enumerate(self.gallery_items):
            cell = ttk.Frame(self.gallery_frame)
            cell.pack(side="left", padx=5, pady=5)

            thumb = None
            if PIL_AVAILABLE:
                try:
                    if kind == "existing":
                        thumb = load_thumbnail(ref, size=(110, 90))
                    else:
                        img = Image.open(ref)
                        img.thumbnail((110, 90))
                        thumb = ImageTk.PhotoImage(img)
                except Exception:
                    thumb = None

            if thumb is not None:
                self._gallery_thumb_refs.append(thumb)
                ttk.Label(cell, image=thumb).pack()
            else:
                ttk.Label(cell, text="(aperçu\nindisponible)", justify="center").pack()

            ttk.Button(cell, text="🗑 Retirer", width=10,
                       command=lambda i=idx: self._remove_gallery_item(i)).pack(pady=(3, 0))

    UNIT_OPTIONS = ["Gr", "cl", "pièce", "cuillère à soupe", "cuillère à café", "autre"]

    @staticmethod
    def _map_unit_for_edit(unit):
        """Convertit une unité stockée (éventuellement héritée d'une ancienne
        version) vers (valeur du menu déroulant, texte personnalisé)."""
        u = (unit or "").strip()
        u_lower = u.lower()
        if u_lower in ("gr", "g", "gramme", "grammes"):
            return "Gr", ""
        if u_lower == "cl":
            return "cl", ""
        if u_lower in ("pièce", "piece", "pieces", "pièces"):
            return "pièce", ""
        if u_lower in ("cuillère à soupe", "cuillere a soupe", "c. à soupe", "cas"):
            return "cuillère à soupe", ""
        if u_lower in ("cuillère à café", "cuillere a café", "cuillere a cafe", "c. à café", "cac"):
            return "cuillère à café", ""
        if u == "":
            return "Gr", ""
        if u == "autre":
            return "autre", ""
        return "autre", u

    @staticmethod
    def _filter_ingredients(full_values, typed):
        if not typed:
            return full_values
        typed_key = ingredient_sort_key(typed)
        filtered = [v for v in full_values if ingredient_sort_key(v).startswith(typed_key)]
        if not filtered:
            filtered = [v for v in full_values if typed_key in ingredient_sort_key(v)]
        return filtered

    def _hide_suggestions(self, entry):
        popup = getattr(entry, "_suggestion_popup", None)
        if popup is not None:
            try:
                popup.destroy()
            except tk.TclError:
                pass
            entry._suggestion_popup = None
            entry._suggestion_listbox = None

    def _show_suggestions(self, entry, filtered):
        self._hide_suggestions(entry)
        if not filtered:
            return
        popup = tk.Toplevel(entry)
        popup.wm_overrideredirect(True)
        try:
            popup.wm_attributes("-topmost", True)
        except tk.TclError:
            pass
        x = entry.winfo_rootx()
        y = entry.winfo_rooty() + entry.winfo_height()
        width = max(entry.winfo_width(), 160)
        height = min(6, len(filtered)) * 20
        popup.wm_geometry(f"{width}x{height}+{x}+{y}")

        listbox = tk.Listbox(popup, height=min(6, len(filtered)), exportselection=False)
        listbox.pack(fill="both", expand=True)
        for v in filtered:
            listbox.insert(tk.END, v)

        def choose(event=None):
            sel = listbox.curselection()
            if sel:
                value = listbox.get(sel[0])
                entry.delete(0, tk.END)
                entry.insert(0, value)
                self._sync_allergens_from_ingredients()
            self._hide_suggestions(entry)
            entry.focus_set()

        listbox.bind("<ButtonRelease-1>", choose)
        listbox.bind("<Return>", choose)
        entry._suggestion_popup = popup
        entry._suggestion_listbox = listbox

    def _sync_allergens_from_ingredients(self):
        """Recalcule les allergènes détectés à partir de TOUS les ingrédients
        actuellement saisis, et coche/décoche les cases en conséquence. Ne
        décoche jamais une case pour un allergène que l'utilisateur aurait
        cochée lui-même sans qu'un ingrédient détecté ne soit à l'origine
        (seuls les allergènes que la détection a elle-même cochés peuvent
        être décochés automatiquement par la suite)."""
        if not hasattr(self, "allergen_vars"):
            return
        current_ingredients = []
        for name_e, qty_e, unit_e, custom_e in self.ingredient_rows:
            ing_name = name_e.get().strip()
            if ing_name:
                current_ingredients.append({"name": ing_name})
        detected = set(compute_recipe_allergens(current_ingredients))

        for allergen in self._auto_detected_allergens - detected:
            if allergen in self.allergen_vars:
                self.allergen_vars[allergen].set(False)
        for allergen in detected:
            if allergen in self.allergen_vars:
                self.allergen_vars[allergen].set(True)

        self._auto_detected_allergens = detected

    def _on_ingredient_keyrelease(self, event, entry):
        if event.keysym == "Down":
            listbox = getattr(entry, "_suggestion_listbox", None)
            if listbox is not None:
                listbox.focus_set()
                listbox.selection_set(0)
            return
        if event.keysym in ("Escape",):
            self._hide_suggestions(entry)
            return
        if event.keysym in ("Return", "Tab", "Shift_L", "Shift_R", "Control_L", "Control_R",
                              "Caps_Lock", "Alt_L", "Alt_R", "Left", "Right"):
            return
        full_values = getattr(entry, "full_values", [])
        typed = entry.get()
        filtered = self._filter_ingredients(full_values, typed)
        if filtered:
            self._show_suggestions(entry, filtered)
        else:
            self._hide_suggestions(entry)

    def _on_ingredient_focus_in(self, event, entry):
        full_values = getattr(entry, "full_values", [])
        typed = entry.get()
        filtered = self._filter_ingredients(full_values, typed)
        if filtered:
            self._show_suggestions(entry, filtered)

    def _on_ingredient_focus_out(self, event, entry):
        # Recalcule les allergènes à partir de l'ensemble des ingrédients
        # actuellement saisis (coche ou décoche selon le résultat).
        self._sync_allergens_from_ingredients()
        # Petit délai pour laisser le temps à un clic dans la liste de
        # suggestions de s'exécuter avant qu'elle ne soit masquée.
        entry.after(200, lambda: self._hide_suggestions(entry))

    def add_ingredient_row(self, name="", qty="", unit=""):
        # On détache temporairement les boutons de fin de liste pour que la
        # nouvelle ligne s'insère bien avant eux, puis on les replace après.
        has_controls = hasattr(self, "add_ingredient_button")
        if has_controls:
            self.add_ingredient_button.pack_forget()
            self.bottom_actions_frame.pack_forget()

        row = ttk.Frame(self.rows_frame)
        row.pack(fill="x", pady=2)
        values = list(self.ingredient_names)
        if name and name not in values:
            values = sorted(values + [name], key=ingredient_sort_key)
        name_e = ttk.Entry(row, width=17)  # champ libre avec suggestions déroulantes personnalisées
        name_e.full_values = values
        name_e._suggestion_popup = None
        name_e._suggestion_listbox = None
        if name:
            name_e.insert(0, name)
        name_e.grid(row=0, column=0, padx=2)
        name_e.bind("<KeyRelease>", lambda e, ent=name_e: self._on_ingredient_keyrelease(e, ent))
        name_e.bind("<FocusIn>", lambda e, ent=name_e: self._on_ingredient_focus_in(e, ent))
        name_e.bind("<FocusOut>", lambda e, ent=name_e: self._on_ingredient_focus_out(e, ent))
        qty_e = ttk.Entry(row, width=9)
        qty_e.insert(0, "" if qty == "" else str(qty))
        qty_e.grid(row=0, column=1, padx=2)

        combo_value, custom_text = self._map_unit_for_edit(unit)
        unit_e = ttk.Combobox(row, width=15, state="readonly", values=self.UNIT_OPTIONS)
        unit_e.set(combo_value)
        unit_e.grid(row=0, column=2, padx=2)

        custom_e = ttk.Entry(row, width=10)
        custom_e.insert(0, custom_text)
        custom_e.grid(row=0, column=3, padx=2)
        if combo_value != "autre":
            custom_e.grid_remove()

        def on_unit_change(event, u=unit_e, c=custom_e):
            if u.get() == "autre":
                c.grid()
            else:
                c.delete(0, tk.END)
                c.grid_remove()

        unit_e.bind("<<ComboboxSelected>>", on_unit_change)
        self.ingredient_rows.append((name_e, qty_e, unit_e, custom_e))

        if has_controls:
            self.add_ingredient_button.pack(pady=10)
            self.bottom_actions_frame.pack(pady=(0, 20))
            # Fait défiler la fenêtre pour amener la nouvelle ligne en vue
            self.update_idletasks()
            self.canvas.yview_moveto(1.0)

    def add_new_ingredient_global(self):
        new_name = simpledialog.askstring("Nouvel ingrédient", "Nom du nouvel ingrédient :", parent=self)
        if not new_name:
            return
        new_name = normalize_oe(new_name.strip())
        if not new_name:
            return
        existing_lower = [n.lower() for n in self.ingredient_names]
        if new_name.lower() in existing_lower:
            messagebox.showinfo("Info", f"L'ingrédient « {new_name} » existe déjà.")
            return

        ingredients = load_ingredients()
        ingredients.append(new_name)
        self.ingredient_names = save_ingredients(ingredients)
        self.app.refresh_ingredients()

        # Met à jour les suggestions déjà affichées dans cette fenêtre
        for name_e, qty_e, unit_e, custom_e in self.ingredient_rows:
            name_e.full_values = self.ingredient_names

        messagebox.showinfo("Ajouté", f"L'ingrédient « {new_name} » a été ajouté.\n"
                                       "Sélectionnez-le dans une des listes déroulantes.")

    @staticmethod
    def _validate_time_field(raw_value, label):
        """Valide un champ de temps optionnel (minutes). Retourne
        (valeur_normalisée, ok). Une valeur vide est acceptée (temps non
        renseigné)."""
        raw_value = raw_value.strip()
        if not raw_value:
            return "", True
        try:
            value = float(raw_value.replace(",", "."))
            if value < 0:
                raise ValueError
        except ValueError:
            return None, False
        if value == int(value):
            return str(int(value)), True
        return str(value), True

    def save_recipe(self):
        name = self.name_entry.get().strip()
        if not name:
            messagebox.showerror("Erreur", "Merci d'indiquer un nom de recette.")
            return

        prep_time, ok_prep = self._validate_time_field(self.prep_time_entry.get(), "préparation")
        if not ok_prep:
            messagebox.showerror("Erreur", "Le temps de préparation doit être un nombre positif (ou vide).")
            return
        cook_time, ok_cook = self._validate_time_field(self.cook_time_entry.get(), "cuisson")
        if not ok_cook:
            messagebox.showerror("Erreur", "Le temps de cuisson doit être un nombre positif (ou vide).")
            return

        ingredients = []
        for name_e, qty_e, unit_e, custom_e in self.ingredient_rows:
            ing_name_raw = name_e.get().strip()
            qty_str = qty_e.get().strip().replace(",", ".")
            unit_choice = unit_e.get().strip()
            if not ing_name_raw:
                continue
            canonical = next(
                (n for n in self.ingredient_names if n.lower() == ing_name_raw.lower()), None
            )
            if canonical is None:
                messagebox.showerror(
                    "Ingrédient inconnu",
                    f"« {ing_name_raw} » ne correspond à aucun ingrédient enregistré.\n"
                    "Choisissez-en un dans la liste déroulante, ou cliquez sur "
                    "« 🥕 Nouvel ingrédient » pour l'ajouter d'abord."
                )
                return
            ing_name = canonical
            try:
                qty = float(qty_str) if qty_str else 0
            except ValueError:
                messagebox.showerror("Erreur", f"Quantité invalide pour '{ing_name}'.")
                return
            if unit_choice == "autre":
                unit = custom_e.get().strip()
                if not unit:
                    messagebox.showerror("Erreur", f"Précisez l'unité personnalisée pour '{ing_name}'.")
                    return
            else:
                unit = unit_choice
            ingredients.append({"name": ing_name, "quantity": qty, "unit": unit})

        if not ingredients:
            messagebox.showerror("Erreur", "Ajoutez au moins un ingrédient valide.")
            return

        # Construit la liste finale des photos : celles déjà enregistrées qui
        # n'ont pas été retirées, plus celles nouvellement choisies (copiées
        # sur le disque à ce moment-là).
        final_images = []
        for kind, ref in self.gallery_items:
            if kind == "existing":
                final_images.append(ref)
            else:
                final_images.append(copy_image_to_store(ref))

        if self.editing:
            old_images = get_recipe_images(self.existing_recipe)
            for fname in old_images:
                if fname not in final_images:
                    delete_image_file(fname)

        try:
            default_persons = float(self.default_persons_entry.get().strip().replace(",", "."))
            if default_persons <= 0:
                raise ValueError
        except ValueError:
            default_persons = 4
        if default_persons == int(default_persons):
            default_persons = int(default_persons)

        tags = [t.strip() for t in self.tags_entry.get().split(",") if t.strip()]

        recipes = load_recipes()
        recipe_data = {
            "name": name,
            "category": self.category_combo.get(),
            "favorite": self.favorite_var.get(),
            "rating": self.rating_value,
            "prep_time": prep_time,
            "cook_time": cook_time,
            "difficulty": self.difficulty_combo.get(),
            "default_persons": default_persons,
            "tags": tags,
            "allergens": [a for a, var in self.allergen_vars.items() if var.get()],
            "description": self.description_text.get("1.0", "end-1c").strip()[: self.MAX_DESC_LEN],
            "personal_notes": self.notes_text.get("1.0", "end-1c").strip()[: self.MAX_NOTES_LEN],
            "ingredients": ingredients,
            "images": final_images,
            "created_at": self.existing_recipe.get("created_at") if self.editing else datetime.now().isoformat(),
            "times_cooked": self.existing_recipe.get("times_cooked", 0) if self.editing else 0,
            "cooked_dates": self.existing_recipe.get("cooked_dates", []) if self.editing else [],
        }
        if not recipe_data["created_at"]:
            recipe_data["created_at"] = datetime.now().isoformat()

        if self.editing:
            recipes[self.recipe_index] = recipe_data
        else:
            recipes.append(recipe_data)

        save_recipes(recipes)
        self.app.refresh_recipes()
        messagebox.showinfo("Succès", f"La recette « {name} » a été enregistrée.")
        self.destroy()

    def delete_recipe(self):
        if not messagebox.askyesno(
            "Confirmer",
            f"Envoyer la recette « {self.existing_recipe['name']} » à la corbeille ?\n\n"
            "Vous pourrez la restaurer plus tard depuis le bouton « 🗑️ Corbeille »."
        ):
            return
        recipes = load_recipes()
        removed = recipes.pop(self.recipe_index)
        save_recipes(recipes)
        move_recipe_to_trash(removed)
        self.app.refresh_recipes()
        messagebox.showinfo("Envoyée à la corbeille", "La recette a été déplacée vers la corbeille.")
        self.destroy()



class ManageRecipesWindow(tk.Toplevel):
    """Fenêtre listant les recettes pour choisir laquelle modifier, dupliquer
    ou supprimer."""

    QUICK_FILTER_LABELS = {
        "favoris": "⭐ Favoris uniquement",
        "rapide": "⏱️ Recettes rapides (≤ 30 min) uniquement",
        "vegetarien": "🥗 Recettes végétariennes uniquement",
    }

    def __init__(self, app, quick_filter=None):
        super().__init__(app)
        self.app = app
        self.title("Modifier / Supprimer une recette")
        self.geometry("560x580")
        self.grab_set()
        self.filtered_indices = []  # correspondance ligne affichée -> index réel dans app.recipes
        self.quick_filter = quick_filter

        ttk.Label(self, text="Sélectionnez une recette :", font=("Segoe UI", 11, "bold")).pack(pady=(10, 5))

        if quick_filter in self.QUICK_FILTER_LABELS:
            filter_bar = ttk.Frame(self)
            filter_bar.pack(fill="x", padx=15, pady=(0, 5))
            ttk.Label(filter_bar, text=self.QUICK_FILTER_LABELS[quick_filter],
                      foreground=COLOR_ACCENT_DARK, font=("Segoe UI", 9, "bold")).pack(side="left")
            ttk.Button(filter_bar, text="✕ Retirer le filtre",
                       command=self._clear_quick_filter).pack(side="right")

        top_frame = ttk.Frame(self)
        top_frame.pack(pady=(0, 5), fill="x", padx=15)
        ttk.Label(top_frame, text="🔍 Rechercher :").pack(side="left")
        self.search_entry = ttk.Entry(top_frame, width=18)
        self.search_entry.pack(side="left", padx=5, fill="x", expand=True)
        self.search_entry.bind("<KeyRelease>", lambda e: self._populate())

        sort_frame = ttk.Frame(self)
        sort_frame.pack(pady=(0, 5), fill="x", padx=15)
        ttk.Label(sort_frame, text="Trier par :").pack(side="left")
        self.sort_combo = ttk.Combobox(sort_frame, values=RECIPE_SORT_OPTIONS, state="readonly", width=22)
        self.sort_combo.set(RECIPE_SORT_OPTIONS[0])
        self.sort_combo.pack(side="left", padx=5)
        self.sort_combo.bind("<<ComboboxSelected>>", lambda e: self._populate())

        list_frame = ttk.Frame(self)
        list_frame.pack(pady=5, padx=15, fill="both", expand=True)
        self.listbox = tk.Listbox(list_frame, width=56, height=15)
        list_scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=list_scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        list_scrollbar.pack(side="right", fill="y")
        self._populate()

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=15)
        ttk.Button(btn_frame, text="✏️ Modifier", command=self.edit_selected).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="📋 Dupliquer", command=self.duplicate_selected).grid(row=0, column=1, padx=5)
        ttk.Button(btn_frame, text="🗑️ Supprimer", command=self.delete_selected).grid(row=0, column=2, padx=5)

    def _clear_quick_filter(self):
        self.quick_filter = None
        self.destroy()
        ManageRecipesWindow(self.app)

    def _matches_quick_filter(self, recipe):
        if self.quick_filter == "favoris":
            return bool(recipe.get("favorite"))
        if self.quick_filter == "rapide":
            try:
                total = float(recipe.get("prep_time") or 0) + float(recipe.get("cook_time") or 0)
            except (TypeError, ValueError):
                return False
            return 0 < total <= 30
        if self.quick_filter == "vegetarien":
            tag_keys = {ingredient_sort_key(t) for t in recipe.get("tags", [])}
            return "vegetarien" in tag_keys or "vegetarienne" in tag_keys
        return True

    def _populate(self):
        self.listbox.delete(0, tk.END)
        self.filtered_indices = []
        search = self.search_entry.get().strip() if hasattr(self, "search_entry") else ""
        search_key = ingredient_sort_key(search) if search else ""
        option = self.sort_combo.get()
        indexed = list(enumerate(self.app.recipes))
        indexed = [pair for pair in indexed if recipe_matches_search(pair[1], search_key)]
        indexed = [pair for pair in indexed if self._matches_quick_filter(pair[1])]
        reverse = option in ("Ajoutées récemment",)
        indexed.sort(key=lambda pair: recipe_sort_key(pair[1], option), reverse=reverse)
        for idx, recipe in indexed:
            self.listbox.insert(tk.END, format_recipe_list_label(recipe))
            self.filtered_indices.append(idx)

    def _selected_index(self):
        sel = self.listbox.curselection()
        if not sel:
            messagebox.showinfo("Info", "Sélectionnez une recette dans la liste.")
            return None
        return self.filtered_indices[sel[0]]

    def edit_selected(self):
        idx = self._selected_index()
        if idx is None:
            return
        self.destroy()
        RecipeFormWindow(self.app, recipe_index=idx)

    def duplicate_selected(self):
        idx = self._selected_index()
        if idx is None:
            return
        recipes = load_recipes()
        original = recipes[idx]
        new_recipe = copy.deepcopy(original)
        new_recipe["name"] = f"{original['name']} (copie)"
        new_recipe["images"] = duplicate_recipe_images(original)
        new_recipe.pop("image", None)
        recipes.append(new_recipe)
        save_recipes(recipes)
        self.app.refresh_recipes()
        self._populate()
        messagebox.showinfo("Dupliquée",
                             f"« {original['name']} » a été dupliquée sous le nom « {new_recipe['name']} ».")

    def delete_selected(self):
        idx = self._selected_index()
        if idx is None:
            return
        recipe = self.app.recipes[idx]
        if not messagebox.askyesno(
            "Confirmer",
            f"Envoyer la recette « {recipe['name']} » à la corbeille ?\n\n"
            "Vous pourrez la restaurer plus tard depuis le bouton « 🗑️ Corbeille »."
        ):
            return
        recipes = load_recipes()
        removed = recipes.pop(idx)
        save_recipes(recipes)
        move_recipe_to_trash(removed)
        self.app.refresh_recipes()
        self._populate()
        messagebox.showinfo("Envoyée à la corbeille", "La recette a été déplacée vers la corbeille.")


class TrashWindow(tk.Toplevel):
    """Corbeille : liste les recettes supprimées récemment, avec la
    possibilité de les restaurer ou de les effacer définitivement."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Corbeille")
        self.geometry("560x520")
        self.grab_set()

        ttk.Label(self, text="🗑️ Recettes supprimées", font=("Segoe UI", 13, "bold")).pack(pady=(15, 5))
        ttk.Label(
            self,
            text="Les photos des recettes de la corbeille sont conservées\n"
                 "jusqu'à leur suppression définitive.",
            justify="center", font=("Segoe UI", 9), foreground=COLOR_TEXT_MUTED
        ).pack(pady=(0, 10))

        list_frame = ttk.Frame(self)
        list_frame.pack(padx=15, pady=5, fill="both", expand=True)
        self.listbox = tk.Listbox(list_frame, width=56, height=14)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self._populate()

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=12)
        ttk.Button(btn_frame, text="♻️ Restaurer", command=self.restore_selected).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="🗑️ Supprimer définitivement",
                   command=self.delete_selected_forever).grid(row=0, column=1, padx=5)
        ttk.Button(btn_frame, text="🧹 Vider la corbeille",
                   command=self.empty_trash).grid(row=0, column=2, padx=5)

    def _populate(self):
        self.listbox.delete(0, tk.END)
        self.trash = load_trash()
        for entry in self.trash:
            recipe = entry.get("recipe", {})
            try:
                deleted_at = datetime.fromisoformat(entry.get("deleted_at", ""))
                date_display = deleted_at.strftime("%d/%m/%Y à %H:%M")
            except ValueError:
                date_display = "date inconnue"
            self.listbox.insert(tk.END, f"{recipe.get('name', '(sans nom)')}  —  supprimée le {date_display}")
        if not self.trash:
            self.listbox.insert(tk.END, "La corbeille est vide.")

    def _selected_index(self):
        sel = self.listbox.curselection()
        if not sel or not self.trash:
            messagebox.showinfo("Info", "Sélectionnez une recette dans la corbeille.")
            return None
        return sel[0]

    def restore_selected(self):
        idx = self._selected_index()
        if idx is None:
            return
        entry = self.trash[idx]
        recipe = entry["recipe"]

        recipes = load_recipes()
        existing_names_lower = {r["name"].strip().lower() for r in recipes}
        if recipe.get("name", "").strip().lower() in existing_names_lower:
            recipe["name"] = f"{recipe['name']} (restaurée)"
        recipes.append(recipe)
        save_recipes(recipes)

        trash = load_trash()
        trash.pop(idx)
        save_trash(trash)

        self.app.refresh_recipes()
        self._populate()
        messagebox.showinfo("Restaurée", f"« {recipe['name']} » a été restaurée.")

    def delete_selected_forever(self):
        idx = self._selected_index()
        if idx is None:
            return
        entry = self.trash[idx]
        recipe = entry["recipe"]
        if not messagebox.askyesno(
            "Confirmer",
            f"Supprimer définitivement « {recipe['name']} » ?\n\nCette action est irréversible."
        ):
            return
        delete_recipe_images(recipe)
        trash = load_trash()
        trash.pop(idx)
        save_trash(trash)
        self._populate()
        messagebox.showinfo("Supprimée", "La recette a été définitivement supprimée.")

    def empty_trash(self):
        trash = load_trash()
        if not trash:
            messagebox.showinfo("Info", "La corbeille est déjà vide.")
            return
        if not messagebox.askyesno(
            "Confirmer",
            f"Supprimer définitivement les {len(trash)} recette(s) de la corbeille ?\n\n"
            "Cette action est irréversible."
        ):
            return
        for entry in trash:
            delete_recipe_images(entry["recipe"])
        save_trash([])
        self._populate()
        messagebox.showinfo("Corbeille vidée", "La corbeille a été vidée.")


class ManageIngredientsWindow(tk.Toplevel):
    """Fenêtre pour ajouter, renommer ou supprimer un ingrédient de la liste
    réutilisable dans les recettes."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Gérer les ingrédients")
        self.geometry("420x680")
        self.grab_set()

        ttk.Label(self, text="Liste des ingrédients enregistrés :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(10, 5))

        search_frame = ttk.Frame(self)
        search_frame.pack(pady=(0, 5), fill="x", padx=15)
        ttk.Label(search_frame, text="🔍 Rechercher :").pack(side="left")
        self.search_entry = ttk.Entry(search_frame, width=28)
        self.search_entry.pack(side="left", padx=5, fill="x", expand=True)
        self.search_entry.bind("<KeyRelease>", lambda e: self._populate())

        list_frame = ttk.Frame(self)
        list_frame.pack(pady=5, padx=15, fill="both", expand=True)
        self.listbox = tk.Listbox(list_frame, width=36, height=14)
        list_scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=list_scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        list_scrollbar.pack(side="right", fill="y")
        self._populate()

        add_frame = ttk.Frame(self)
        add_frame.pack(pady=10)
        self.new_entry = ttk.Entry(add_frame, width=25)
        self.new_entry.grid(row=0, column=0, padx=5)
        ttk.Button(add_frame, text="➕ Ajouter", command=self.add_ingredient).grid(row=0, column=1)
        self.new_entry.bind("<Return>", lambda e: self.add_ingredient())

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="✏️ Modifier", command=self.edit_selected).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="🗑️ Supprimer", command=self.delete_selected).grid(row=0, column=1, padx=5)

        ttk.Button(self, text="📚 Charger les ~1000 ingrédients courants",
                   command=self.load_defaults).pack(pady=(5, 5))
        ttk.Button(self, text="🔤 Vérifier les doublons / fautes de frappe",
                   command=self.open_spell_check).pack(pady=(0, 5))
        ttk.Button(self, text="💰 Gérer les prix (pour le coût des recettes)",
                   command=self.open_prices).pack(pady=(0, 10))

        ttk.Label(self, text="\"Modifier\" permet de changer le nom (mis à jour\n"
                              "partout où l'ingrédient est utilisé), ses allergènes,\n"
                              "ses valeurs nutritionnelles et son prix.",
                  font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center").pack(pady=(0, 10))

    def _populate(self):
        self.listbox.delete(0, tk.END)
        search = self.search_entry.get().strip() if hasattr(self, "search_entry") else ""
        if search:
            search_key = ingredient_sort_key(search)
            names = [n for n in self.app.ingredient_names if search_key in ingredient_sort_key(n)]
        else:
            names = self.app.ingredient_names
        for name in names:
            self.listbox.insert(tk.END, name)

    def _selected_name(self):
        sel = self.listbox.curselection()
        if not sel:
            messagebox.showinfo("Info", "Sélectionnez un ingrédient dans la liste.")
            return None
        return self.listbox.get(sel[0])

    def add_ingredient(self):
        prefill_name = normalize_oe(self.new_entry.get().strip())
        self.new_entry.delete(0, tk.END)
        IngredientEditWindow(self.app, manage_window=self, existing_name=None,
                              prefill_name=prefill_name)

    def edit_selected(self):
        name = self._selected_name()
        if name is None:
            return
        IngredientEditWindow(self.app, manage_window=self, existing_name=name)

    def delete_selected(self):
        name = self._selected_name()
        if name is None:
            return
        usage = count_ingredient_usage(name)
        message = f"Supprimer « {name} » de la liste des ingrédients ?"
        if usage:
            message += (f"\n\nAttention : il est utilisé dans {usage} recette(s). "
                        "Ces recettes conserveront cet ingrédient, mais il ne sera "
                        "plus proposé dans le menu déroulant, sauf si vous le rajoutez.")
        if not messagebox.askyesno("Confirmer", message):
            return
        ingredients = [n for n in load_ingredients() if n.lower() != name.lower()]
        self.app.ingredient_names = save_ingredients(ingredients)
        self._populate()

    def load_defaults(self):
        if not os.path.exists(DEFAULT_INGREDIENTS_FILE):
            messagebox.showerror(
                "Fichier manquant",
                "Le fichier ingredients_par_defaut.json est introuvable.\n"
                "Assurez-vous qu'il se trouve dans le même dossier que main.py."
            )
            return
        added = merge_default_ingredients()
        self.app.ingredient_names = load_ingredients()
        self._populate()
        if added:
            messagebox.showinfo(
                "Terminé",
                f"{added} nouvel(aux) ingrédient(s) ajouté(s) à partir de la liste courante."
            )
        else:
            messagebox.showinfo("Terminé", "Tous les ingrédients courants étaient déjà présents.")

    def open_spell_check(self):
        IngredientSpellCheckWindow(self.app, manage_window=self)

    def open_prices(self):
        IngredientPricesWindow(self.app)


class IngredientPricesWindow(tk.Toplevel):
    """Fenêtre pour renseigner le prix de vos ingrédients, utilisé ensuite
    pour estimer le coût de vos recettes. Les prix sont saisis par vous —
    aucune source de prix en ligne n'est utilisée."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Gérer les prix des ingrédients")
        self.geometry("480x620")
        self.grab_set()

        ttk.Label(self, text="💰 Prix des ingrédients", font=("Segoe UI", 13, "bold")).pack(pady=(15, 5))
        ttk.Label(
            self,
            text="Renseignez un prix pour les ingrédients qui vous\n"
                 "intéressent — inutile de tous les faire. Le coût\n"
                 "d'une recette est estimé à partir de ces prix.",
            justify="center", font=("Segoe UI", 9)
        ).pack(pady=(0, 10))

        search_frame = ttk.Frame(self)
        search_frame.pack(pady=(0, 5), fill="x", padx=15)
        ttk.Label(search_frame, text="🔍 Rechercher :").pack(side="left")
        self.search_entry = ttk.Entry(search_frame, width=28)
        self.search_entry.pack(side="left", padx=5, fill="x", expand=True)
        self.search_entry.bind("<KeyRelease>", lambda e: self._populate())

        list_frame = ttk.Frame(self)
        list_frame.pack(pady=5, padx=15, fill="both", expand=True)
        self.listbox = tk.Listbox(list_frame, width=48, height=13)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self.listbox.bind("<<ListboxSelect>>", lambda e: self._load_selected_price())
        self._populate()

        edit_frame = ttk.Frame(self)
        edit_frame.pack(pady=10, padx=15, fill="x")
        ttk.Label(edit_frame, text="Prix (€) :").grid(row=0, column=0, padx=3)
        self.price_entry = ttk.Entry(edit_frame, width=8)
        self.price_entry.grid(row=0, column=1, padx=3)
        ttk.Label(edit_frame, text="pour 1").grid(row=0, column=2, padx=3)
        self.unit_combo = ttk.Combobox(edit_frame, values=PRICE_UNIT_OPTIONS,
                                        state="readonly", width=15)
        self.unit_combo.set(PRICE_UNIT_OPTIONS[0])
        self.unit_combo.grid(row=0, column=3, padx=3)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="💾 Enregistrer le prix",
                   command=self.save_price).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="🗑 Effacer le prix",
                   command=self.clear_price).grid(row=0, column=1, padx=5)

        ttk.Label(
            self,
            text="kg ↔ recettes en Gr   ·   L ↔ recettes en cl   ·   les prix\n"
                 "en pièce/cuillère s'appliquent tels quels.",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center"
        ).pack(pady=(0, 10))

    def _populate(self):
        self.listbox.delete(0, tk.END)
        search = self.search_entry.get().strip()
        search_key = ingredient_sort_key(search) if search else ""
        self._names = []
        for name in self.app.ingredient_names:
            if search_key and search_key not in ingredient_sort_key(name):
                continue
            self._names.append(name)
            price_info = get_ingredient_price(name)
            if price_info:
                suffix = f"  —  {price_info['price']:.2f} € / {price_info['unit']}"
            else:
                suffix = "  —  (prix non renseigné)"
            self.listbox.insert(tk.END, f"{name}{suffix}")

    def _selected_name(self):
        sel = self.listbox.curselection()
        if not sel or not self._names:
            return None
        return self._names[sel[0]]

    def _load_selected_price(self):
        name = self._selected_name()
        if name is None:
            return
        price_info = get_ingredient_price(name)
        self.price_entry.delete(0, tk.END)
        if price_info:
            self.price_entry.insert(0, str(price_info["price"]))
            self.unit_combo.set(price_info["unit"])
        else:
            self.unit_combo.set(PRICE_UNIT_OPTIONS[0])

    def save_price(self):
        name = self._selected_name()
        if name is None:
            messagebox.showinfo("Info", "Sélectionnez un ingrédient dans la liste.")
            return
        raw_price = self.price_entry.get().strip().replace(",", ".")
        try:
            price = float(raw_price)
            if price < 0:
                raise ValueError
        except ValueError:
            messagebox.showerror("Erreur", "Entrez un prix valide (nombre positif).")
            return
        set_ingredient_price(name, price, self.unit_combo.get())
        self._populate()
        messagebox.showinfo("Enregistré", f"Prix enregistré pour « {name} ».")

    def clear_price(self):
        name = self._selected_name()
        if name is None:
            messagebox.showinfo("Info", "Sélectionnez un ingrédient dans la liste.")
            return
        set_ingredient_price(name, None, None)
        self.price_entry.delete(0, tk.END)
        self._populate()


class IngredientEditWindow(tk.Toplevel):
    """Fenêtre unifiée pour ajouter un nouvel ingrédient ou modifier un
    ingrédient existant : nom, allergènes, valeurs nutritionnelles et prix."""

    def __init__(self, app, manage_window=None, existing_name=None, prefill_name=""):
        super().__init__(app)
        self.app = app
        self.manage_window = manage_window
        self.existing_name = existing_name
        self.editing = existing_name is not None
        self.title("Modifier un ingrédient" if self.editing else "Nouvel ingrédient")
        self.geometry("480x680")
        self.grab_set()

        ttk.Label(self, text="✏️ Modifier l'ingrédient" if self.editing else "➕ Nouvel ingrédient",
                  font=("Segoe UI", 13, "bold")).pack(pady=(15, 10))

        ttk.Label(self, text="Nom :", font=("Segoe UI", 10, "bold")).pack()
        self.name_entry = ttk.Entry(self, width=40)
        self.name_entry.pack(pady=(2, 10))
        self.name_entry.insert(0, existing_name if self.editing else prefill_name)

        ttk.Label(self, text="Allergènes présents :", font=("Segoe UI", 10, "bold")).pack(pady=(5, 5))
        allergens_frame = ttk.Frame(self)
        allergens_frame.pack()
        existing_allergens = set(get_ingredient_allergens(existing_name)) if self.editing else set()
        self.allergen_vars = {}
        for i, allergen in enumerate(ALLERGENS):
            var = tk.BooleanVar(value=allergen in existing_allergens)
            self.allergen_vars[allergen] = var
            ttk.Checkbutton(allergens_frame, text=allergen, variable=var).grid(
                row=i // 3, column=i % 3, sticky="w", padx=8, pady=2
            )

        ttk.Label(self, text="Valeurs nutritionnelles (pour 100 g / 100 ml) :",
                  font=("Segoe UI", 10, "bold")).pack(pady=(15, 5))
        nutri_frame = ttk.Frame(self)
        nutri_frame.pack()
        existing_nutri = (get_ingredient_nutrition(existing_name) or {}) if self.editing else {}
        nutri_labels = [("kcal", "Calories (kcal)"), ("protein_g", "Protéines (g)"),
                         ("carbs_g", "Glucides (g)"), ("fat_g", "Lipides (g)")]
        self.nutri_entries = {}
        for i, (key, label) in enumerate(nutri_labels):
            ttk.Label(nutri_frame, text=f"{label} :").grid(row=i, column=0, sticky="e", padx=5, pady=3)
            entry = ttk.Entry(nutri_frame, width=10)
            if key in existing_nutri:
                entry.insert(0, str(existing_nutri[key]))
            entry.grid(row=i, column=1, sticky="w", padx=5, pady=3)
            self.nutri_entries[key] = entry
        ttk.Label(self, text="Laissez vide si vous ne connaissez pas ces valeurs.",
                  font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED).pack()

        ttk.Label(self, text="Prix :", font=("Segoe UI", 10, "bold")).pack(pady=(15, 5))
        price_frame = ttk.Frame(self)
        price_frame.pack()
        existing_price = get_ingredient_price(existing_name) if self.editing else None
        ttk.Label(price_frame, text="Prix (€) :").grid(row=0, column=0, padx=3)
        self.price_entry = ttk.Entry(price_frame, width=8)
        if existing_price:
            self.price_entry.insert(0, str(existing_price["price"]))
        self.price_entry.grid(row=0, column=1, padx=3)
        ttk.Label(price_frame, text="pour 1").grid(row=0, column=2, padx=3)
        self.unit_combo = ttk.Combobox(price_frame, values=PRICE_UNIT_OPTIONS, state="readonly", width=15)
        self.unit_combo.set(existing_price["unit"] if existing_price else PRICE_UNIT_OPTIONS[0])
        self.unit_combo.grid(row=0, column=3, padx=3)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=15)
        ttk.Button(btn_frame, text="💾 Enregistrer", command=self.save).grid(row=0, column=0, padx=5)
        if self.editing:
            ttk.Button(btn_frame, text="🗑️ Supprimer cet ingrédient",
                       command=self.delete_ingredient).grid(row=0, column=1, padx=5)

    def _parse_float_or_none(self, entry, field_label):
        raw = entry.get().strip().replace(",", ".")
        if not raw:
            return None, True
        try:
            value = float(raw)
            if value < 0:
                raise ValueError
            return value, True
        except ValueError:
            messagebox.showerror("Erreur", f"« {field_label} » doit être un nombre positif (ou vide).")
            return None, False

    def save(self):
        new_name = normalize_oe(self.name_entry.get().strip())
        if not new_name:
            messagebox.showerror("Erreur", "Merci d'indiquer un nom d'ingrédient.")
            return

        existing_lower = [
            n.lower() for n in self.app.ingredient_names
            if not (self.editing and n.lower() == self.existing_name.lower())
        ]
        if new_name.lower() in existing_lower:
            messagebox.showerror("Erreur", f"L'ingrédient « {new_name} » existe déjà.")
            return

        nutrition = {}
        nutri_field_labels = {"kcal": "Calories", "protein_g": "Protéines",
                               "carbs_g": "Glucides", "fat_g": "Lipides"}
        for key, entry in self.nutri_entries.items():
            value, ok = self._parse_float_or_none(entry, nutri_field_labels[key])
            if not ok:
                return
            if value is not None:
                nutrition[key] = value

        raw_price = self.price_entry.get().strip().replace(",", ".")
        price = None
        if raw_price:
            try:
                price = float(raw_price)
                if price < 0:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Erreur", "Le prix doit être un nombre positif (ou vide).")
                return

        allergens = [a for a, var in self.allergen_vars.items() if var.get()]

        if self.editing:
            old_name = self.existing_name
            if new_name.lower() != old_name.lower():
                ingredients = [new_name if n.lower() == old_name.lower() else n
                               for n in load_ingredients()]
                self.app.ingredient_names = save_ingredients(ingredients)
                rename_ingredient_everywhere(old_name, new_name)
                rename_ingredient_override(old_name, new_name)
                old_price = get_ingredient_price(old_name)
                if old_price:
                    set_ingredient_price(old_name, None, None)
                    set_ingredient_price(new_name, old_price["price"], old_price["unit"])
                self.app.refresh_recipes()
        else:
            ingredients = load_ingredients()
            ingredients.append(new_name)
            self.app.ingredient_names = save_ingredients(ingredients)

        set_ingredient_override(new_name, allergens=allergens, nutrition=nutrition)
        if raw_price:
            set_ingredient_price(new_name, price, self.unit_combo.get())
        elif self.editing:
            set_ingredient_price(new_name, None, None)

        if self.manage_window is not None:
            self.manage_window.app.ingredient_names = self.app.ingredient_names
            self.manage_window._populate()

        messagebox.showinfo("Enregistré", f"« {new_name} » a été enregistré.")
        self.destroy()

    def delete_ingredient(self):
        name = self.existing_name
        usage = count_ingredient_usage(name)
        message = f"Supprimer « {name} » de la liste des ingrédients ?"
        if usage:
            message += (f"\n\nAttention : il est utilisé dans {usage} recette(s). "
                        "Ces recettes conserveront cet ingrédient, mais il ne sera "
                        "plus proposé dans le menu déroulant, sauf si vous le rajoutez.")
        if not messagebox.askyesno("Confirmer", message):
            return
        ingredients = [n for n in load_ingredients() if n.lower() != name.lower()]
        self.app.ingredient_names = save_ingredients(ingredients)
        set_ingredient_override(name, allergens=[], nutrition={})
        set_ingredient_price(name, None, None)
        if self.manage_window is not None:
            self.manage_window.app.ingredient_names = self.app.ingredient_names
            self.manage_window._populate()
        self.destroy()


class IngredientSpellCheckWindow(tk.Toplevel):
    """Détecte les paires d'ingrédients qui se ressemblent fortement (90 % de
    similarité ou plus — pluriels non fusionnés, fautes de frappe) et propose
    de les fusionner, une par une ou plusieurs à la fois."""

    SIMILARITY_THRESHOLD = 0.90

    def __init__(self, app, manage_window=None):
        super().__init__(app)
        self.app = app
        self.manage_window = manage_window
        self.title("Vérification orthographique des ingrédients")
        self.geometry("580x560")
        self.grab_set()

        ttk.Label(
            self,
            text="Paires d'ingrédients qui se ressemblent à 90 % ou plus\n"
                 "(doublons probables ou fautes de frappe) :",
            font=("Segoe UI", 11, "bold"), justify="center"
        ).pack(pady=10)
        ttk.Label(
            self,
            text="Sélection multiple possible (Ctrl+clic ou Maj+clic) pour\n"
                 "fusionner plusieurs paires d'un coup.",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center"
        ).pack(pady=(0, 5))

        list_frame = ttk.Frame(self)
        list_frame.pack(padx=15, pady=5, fill="both", expand=True)
        self.listbox = tk.Listbox(list_frame, width=64, height=16, selectmode="extended")
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.pairs = []
        self._scan()

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="🔗 Fusionner la sélection",
                   command=self.merge_selected).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="🔄 Relancer l'analyse",
                   command=self._scan).grid(row=0, column=1, padx=5)

        ttk.Label(
            self,
            text="Pour une seule paire, on vous demande laquelle des deux\n"
                 "graphies garder. Pour plusieurs paires à la fois, l'ingrédient\n"
                 "le moins utilisé dans vos recettes est automatiquement fusionné\n"
                 "vers celui utilisé dans le plus de recettes.",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center"
        ).pack(pady=(0, 10))

    def _scan(self):
        self.listbox.delete(0, tk.END)
        self.pairs = find_similar_ingredient_pairs(
            self.app.ingredient_names, threshold=self.SIMILARITY_THRESHOLD
        )
        if not self.pairs:
            self.listbox.insert(tk.END, "Aucun doublon probable détecté. 🎉")
            return
        for name_a, name_b, ratio in self.pairs:
            self.listbox.insert(tk.END, f"{name_a}   ↔   {name_b}     ({int(ratio * 100)} % similaires)")

    def _merge_pair(self, keep, remove):
        ingredients = [n for n in load_ingredients() if n.lower() != remove.lower()]
        self.app.ingredient_names = save_ingredients(ingredients)
        rename_ingredient_everywhere(remove, keep)

    def merge_selected(self):
        sel = self.listbox.curselection()
        if not sel or not self.pairs:
            messagebox.showinfo("Info", "Sélectionnez au moins une paire dans la liste.")
            return
        selected_pairs = [self.pairs[i] for i in sel]

        if len(selected_pairs) == 1:
            name_a, name_b, ratio = selected_pairs[0]
            choice = messagebox.askyesnocancel(
                "Fusionner",
                f"Fusionner « {name_a} » et « {name_b} » ?\n\n"
                f"Oui = tout renommer en « {name_a} »\n"
                f"Non = tout renommer en « {name_b} »\n"
                f"Annuler = ne rien faire"
            )
            if choice is None:
                return
            keep, remove = (name_a, name_b) if choice else (name_b, name_a)
            self._merge_pair(keep, remove)
            self.app.refresh_recipes()
            if self.manage_window is not None:
                self.manage_window.app.ingredient_names = self.app.ingredient_names
                self.manage_window._populate()
            messagebox.showinfo("Fusionné", f"« {remove} » a été fusionné avec « {keep} ».")
        else:
            if not messagebox.askyesno(
                "Confirmer",
                f"Fusionner automatiquement ces {len(selected_pairs)} paires ?\n\n"
                "Pour chaque paire, l'ingrédient le moins utilisé dans vos "
                "recettes sera fusionné vers celui utilisé dans le plus de "
                "recettes (le premier par ordre alphabétique en cas d'égalité)."
            ):
                return
            for name_a, name_b, ratio in selected_pairs:
                usage_a = count_ingredient_usage(name_a)
                usage_b = count_ingredient_usage(name_b)
                if usage_a > usage_b:
                    keep, remove = name_a, name_b
                elif usage_b > usage_a:
                    keep, remove = name_b, name_a
                else:
                    keep, remove = sorted([name_a, name_b], key=ingredient_sort_key)
                self._merge_pair(keep, remove)
            self.app.refresh_recipes()
            if self.manage_window is not None:
                self.manage_window.app.ingredient_names = self.app.ingredient_names
                self.manage_window._populate()
            messagebox.showinfo("Fusionné", f"{len(selected_pairs)} paire(s) fusionnée(s).")

        self._scan()


def build_full_backup_zip(path):
    """Construit une archive ZIP contenant recipes.json, ingredients.json et
    toutes les photos. Utilisé aussi bien par l'export manuel que par les
    sauvegardes automatiques."""
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        if os.path.exists(DATA_FILE):
            zf.write(DATA_FILE, arcname="recipes.json")
        if os.path.exists(INGREDIENTS_FILE):
            zf.write(INGREDIENTS_FILE, arcname="ingredients.json")
        if os.path.isdir(IMAGES_DIR):
            for fname in os.listdir(IMAGES_DIR):
                full = os.path.join(IMAGES_DIR, fname)
                if os.path.isfile(full):
                    zf.write(full, arcname=f"images/{fname}")


def restore_from_zip(path, merge):
    """Restaure des données à partir d'une archive ZIP (export manuel ou
    sauvegarde automatique). merge=True fusionne avec les données actuelles,
    merge=False remplace tout."""
    with zipfile.ZipFile(path, "r") as zf:
        names = zf.namelist()
        imported_recipes = []
        imported_ingredients = []
        if "recipes.json" in names:
            imported_recipes = json.loads(zf.read("recipes.json").decode("utf-8"))
        if "ingredients.json" in names:
            imported_ingredients = json.loads(zf.read("ingredients.json").decode("utf-8"))
        image_entries = [n for n in names if n.startswith("images/") and not n.endswith("/")]

        if not merge:
            if os.path.isdir(IMAGES_DIR):
                for fname in os.listdir(IMAGES_DIR):
                    full = os.path.join(IMAGES_DIR, fname)
                    if os.path.isfile(full):
                        os.remove(full)
            for entry in image_entries:
                fname = os.path.basename(entry)
                with open(os.path.join(IMAGES_DIR, fname), "wb") as out:
                    out.write(zf.read(entry))
            save_recipes(imported_recipes)
            save_ingredients(imported_ingredients)
        else:
            rename_map = {}
            for entry in image_entries:
                old_fname = os.path.basename(entry)
                ext = os.path.splitext(old_fname)[1]
                new_fname = f"{uuid.uuid4().hex}{ext}"
                with open(os.path.join(IMAGES_DIR, new_fname), "wb") as out:
                    out.write(zf.read(entry))
                rename_map[old_fname] = new_fname

            existing_recipes = load_recipes()
            existing_names_lower = {r["name"].strip().lower() for r in existing_recipes}
            for r in imported_recipes:
                img = r.get("image")
                if img and img in rename_map:
                    r["image"] = rename_map[img]
                imgs = r.get("images")
                if imgs:
                    r["images"] = [rename_map.get(f, f) for f in imgs]
                if r.get("name", "").strip().lower() in existing_names_lower:
                    r["name"] = f"{r['name']} (importé)"
                existing_recipes.append(r)
                existing_names_lower.add(r["name"].strip().lower())
            save_recipes(existing_recipes)

            existing_ingredients = load_ingredients()
            save_ingredients(existing_ingredients + imported_ingredients)


# ---------------------------------------------------------------------------
# Paramètres de l'application (dossier de sauvegarde cloud, etc.)
# ---------------------------------------------------------------------------

def load_settings():
    if os.path.exists(SETTINGS_FILE):
        try:
            with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
                if isinstance(data, dict):
                    return data
        except Exception:
            return {}
    return {}


def save_settings(settings):
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(settings, f, ensure_ascii=False, indent=2)


def get_cloud_backup_folder():
    folder = load_settings().get("cloud_backup_folder") or ""
    return folder if folder and os.path.isdir(folder) else ""


def set_cloud_backup_folder(path):
    settings = load_settings()
    if path:
        settings["cloud_backup_folder"] = path
    else:
        settings.pop("cloud_backup_folder", None)
    save_settings(settings)


AUTO_BACKUP_RETENTION = 10  # nombre de sauvegardes automatiques conservées
AUTO_BACKUP_MIN_INTERVAL_HOURS = 24  # fréquence minimale entre deux sauvegardes auto
# Préfixe distinctif des fichiers de sauvegarde, pour ne jamais les confondre
# avec ceux d'un autre programme dans un même dossier (ex. cloud partagé).
AUTO_BACKUP_PREFIX = "sauvegarde_auto_mesrecettes_"


def list_auto_backups():
    """Retourne la liste des sauvegardes automatiques existantes
    (chemin complet), les plus récentes en premier."""
    if not os.path.isdir(BACKUPS_DIR):
        return []
    files = [
        os.path.join(BACKUPS_DIR, f) for f in os.listdir(BACKUPS_DIR)
        if f.startswith(AUTO_BACKUP_PREFIX) and f.endswith(".zip")
    ]
    files.sort(key=os.path.getmtime, reverse=True)
    return files


def create_auto_backup():
    """Crée une nouvelle sauvegarde automatique horodatée, puis supprime les
    plus anciennes au-delà de AUTO_BACKUP_RETENTION. Si un dossier cloud est
    configuré (Google Drive, OneDrive, Dropbox...), une copie y est aussi
    déposée : le client cloud déjà installé sur le PC se charge ensuite de
    l'envoyer en ligne automatiquement."""
    os.makedirs(BACKUPS_DIR, exist_ok=True)
    filename = f"{AUTO_BACKUP_PREFIX}{datetime.now().strftime('%Y-%m-%d_%H%M%S')}.zip"
    path = os.path.join(BACKUPS_DIR, filename)
    build_full_backup_zip(path)

    existing = list_auto_backups()
    for old_path in existing[AUTO_BACKUP_RETENTION:]:
        try:
            os.remove(old_path)
        except OSError:
            pass

    cloud_folder = get_cloud_backup_folder()
    if cloud_folder:
        try:
            cloud_path = os.path.join(cloud_folder, filename)
            shutil.copy2(path, cloud_path)
            cloud_backups = sorted(
                [os.path.join(cloud_folder, f) for f in os.listdir(cloud_folder)
                 if f.startswith(AUTO_BACKUP_PREFIX) and f.endswith(".zip")],
                key=os.path.getmtime, reverse=True
            )
            for old_cloud_path in cloud_backups[AUTO_BACKUP_RETENTION:]:
                try:
                    os.remove(old_cloud_path)
                except OSError:
                    pass
        except Exception:
            pass  # un souci côté dossier cloud ne doit jamais faire échouer la sauvegarde locale

    return path


def migrate_old_backup_filenames():
    """Renomme les sauvegardes créées avant l'ajout du préfixe distinctif
    'mesrecettes' (ex. 'sauvegarde_auto_2026-01-01_120000.zip') vers le
    nouveau format, pour qu'elles restent reconnues par l'application au
    lieu de devenir invisibles."""
    if not os.path.isdir(BACKUPS_DIR):
        return
    old_pattern = re.compile(r"^sauvegarde_auto_(\d{4}-\d{2}-\d{2}_\d{6})\.zip$")
    for fname in os.listdir(BACKUPS_DIR):
        match = old_pattern.match(fname)
        if match:
            new_path = os.path.join(BACKUPS_DIR, f"{AUTO_BACKUP_PREFIX}{match.group(1)}.zip")
            if not os.path.exists(new_path):
                try:
                    os.rename(os.path.join(BACKUPS_DIR, fname), new_path)
                except OSError:
                    pass


def maybe_create_auto_backup():
    """Crée automatiquement une sauvegarde si aucune n'existe encore ou si la
    dernière date de plus de AUTO_BACKUP_MIN_INTERVAL_HOURS. Ne fait jamais
    planter l'application en cas de problème (disque plein, permissions...)."""
    try:
        migrate_old_backup_filenames()
        if not os.path.exists(DATA_FILE):
            return  # rien à sauvegarder pour un tout premier lancement
        existing = list_auto_backups()
        if existing:
            last_mtime = os.path.getmtime(existing[0])
            age_hours = (datetime.now().timestamp() - last_mtime) / 3600
            if age_hours < AUTO_BACKUP_MIN_INTERVAL_HOURS:
                return
        create_auto_backup()
    except Exception:
        pass


class ImportExportWindow(tk.Toplevel):
    """Fenêtre pour exporter ou importer l'ensemble des données (recettes,
    ingrédients et photos) sous forme d'une seule archive ZIP, et pour
    consulter/restaurer les sauvegardes automatiques."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Importer / Exporter les données")
        self.geometry("480x880")
        self.minsize(440, 500)
        self.resizable(True, True)
        self.grab_set()

        ttk.Label(self, text="Sauvegarder ou transférer vos données",
                  font=("Segoe UI", 13, "bold")).pack(pady=15)

        ttk.Label(
            self,
            text="L'export crée un fichier .zip contenant vos recettes,\n"
                 "vos ingrédients et vos photos, pour les sauvegarder\n"
                 "ou les transférer sur un autre ordinateur.",
            justify="center", font=("Segoe UI", 9)
        ).pack(pady=(0, 10))

        ttk.Button(self, text="📤 Exporter toutes mes données (.zip)",
                   width=42, command=self.export_data).pack(pady=6)

        ttk.Label(
            self,
            text="L'import lit un fichier .zip précédemment exporté.\n"
                 "Vous pourrez choisir de fusionner avec vos données\n"
                 "actuelles, ou de tout remplacer.",
            justify="center", font=("Segoe UI", 9)
        ).pack(pady=(10, 10))

        ttk.Button(self, text="📥 Importer des données (.zip)",
                   width=42, command=self.import_data).pack(pady=6)

        ttk.Separator(self, orient="horizontal").pack(fill="x", padx=20, pady=15)

        ttk.Label(self, text="🗄️ Sauvegardes automatiques", font=("Segoe UI", 12, "bold")).pack()
        ttk.Label(
            self,
            text=f"Une sauvegarde est créée automatiquement au démarrage de\n"
                 f"l'application (au maximum une par {AUTO_BACKUP_MIN_INTERVAL_HOURS}h), et les "
                 f"{AUTO_BACKUP_RETENTION} plus\nrécentes sont conservées ici.",
            justify="center", font=("Segoe UI", 9)
        ).pack(pady=(5, 10))

        list_frame = ttk.Frame(self)
        list_frame.pack(padx=15, fill="both", expand=True)
        self.backup_listbox = tk.Listbox(list_frame, height=8)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.backup_listbox.yview)
        self.backup_listbox.configure(yscrollcommand=scrollbar.set)
        self.backup_listbox.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self._populate_backups()

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="💾 Sauvegarder maintenant",
                   command=self.backup_now).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="♻️ Restaurer la sélection",
                   command=self.restore_selected).grid(row=0, column=1, padx=5)

        ttk.Separator(self, orient="horizontal").pack(fill="x", padx=20, pady=15)

        ttk.Label(self, text="☁️ Sauvegarde automatique dans le cloud",
                  font=("Segoe UI", 12, "bold")).pack()
        ttk.Label(
            self,
            text="Choisissez un dossier synchronisé par un client déjà\n"
                 "installé sur ce PC (Google Drive, OneDrive, Dropbox...).\n"
                 "Chaque sauvegarde automatique y sera aussi copiée, et ce\n"
                 "client se chargera de l'envoyer dans le cloud tout seul.",
            justify="center", font=("Segoe UI", 9)
        ).pack(pady=(5, 8))

        self.cloud_folder_label = ttk.Label(self, text="", font=("Segoe UI", 9, "bold"),
                                             foreground="#266", wraplength=400, justify="center")
        self.cloud_folder_label.pack(pady=(0, 8))
        self._refresh_cloud_label()

        cloud_btn_frame = ttk.Frame(self)
        cloud_btn_frame.pack(pady=(0, 15))
        ttk.Button(cloud_btn_frame, text="📁 Choisir un dossier cloud",
                   command=self.choose_cloud_folder).grid(row=0, column=0, padx=5)
        ttk.Button(cloud_btn_frame, text="🚫 Désactiver",
                   command=self.disable_cloud_backup).grid(row=0, column=1, padx=5)

    def _refresh_cloud_label(self):
        folder = get_cloud_backup_folder()
        if folder:
            self.cloud_folder_label.config(text=f"✅ Activé : {folder}")
        else:
            self.cloud_folder_label.config(text="Non configuré pour le moment.", foreground=COLOR_TEXT_MUTED)

    def choose_cloud_folder(self):
        folder = filedialog.askdirectory(
            title="Choisir un dossier synchronisé (Google Drive, OneDrive, Dropbox...)"
        )
        if not folder:
            return
        set_cloud_backup_folder(folder)
        self._refresh_cloud_label()
        if messagebox.askyesno(
            "Dossier configuré",
            f"Dossier cloud configuré :\n{folder}\n\n"
            "Voulez-vous y copier une sauvegarde dès maintenant ?"
        ):
            self.backup_now()

    def disable_cloud_backup(self):
        if not get_cloud_backup_folder():
            return
        set_cloud_backup_folder(None)
        self._refresh_cloud_label()
        messagebox.showinfo("Désactivé", "La sauvegarde automatique dans le cloud est désactivée.")

    def _populate_backups(self):
        self.backup_listbox.delete(0, tk.END)
        self.backups = list_auto_backups()
        for path in self.backups:
            mtime = datetime.fromtimestamp(os.path.getmtime(path))
            size_kb = os.path.getsize(path) / 1024
            self.backup_listbox.insert(
                tk.END, f"{mtime.strftime('%d/%m/%Y à %H:%M')}   ({size_kb:.0f} Ko)"
            )
        if not self.backups:
            self.backup_listbox.insert(tk.END, "Aucune sauvegarde automatique pour le moment.")

    def backup_now(self):
        try:
            create_auto_backup()
        except Exception as e:
            messagebox.showerror("Erreur", f"La sauvegarde a échoué :\n{e}")
            return
        self._populate_backups()
        messagebox.showinfo("Sauvegarde créée", "Une nouvelle sauvegarde automatique a été créée.")

    def restore_selected(self):
        sel = self.backup_listbox.curselection()
        if not sel or not self.backups:
            messagebox.showinfo("Info", "Sélectionnez une sauvegarde dans la liste.")
            return
        path = self.backups[sel[0]]

        mode = messagebox.askyesnocancel(
            "Mode de restauration",
            "Comment restaurer cette sauvegarde ?\n\n"
            "Oui = Fusionner (ajouter aux données actuelles, sans rien supprimer)\n"
            "Non = Remplacer entièrement les données actuelles\n"
            "Annuler = ne rien faire"
        )
        if mode is None:
            return
        try:
            restore_from_zip(path, merge=bool(mode))
        except Exception as e:
            messagebox.showerror("Erreur", f"La restauration a échoué :\n{e}")
            return
        self.app.refresh_recipes()
        self.app.refresh_ingredients()
        messagebox.showinfo("Restauration terminée", "Les données ont été restaurées avec succès.")
        self.destroy()

    def export_data(self):
        path = filedialog.asksaveasfilename(
            title="Exporter mes données",
            defaultextension=".zip",
            filetypes=[("Archive ZIP", "*.zip")],
            initialfile="mes_recettes_export.zip"
        )
        if not path:
            return
        try:
            build_full_backup_zip(path)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Vos données ont été exportées vers :\n{path}")

    def import_data(self):
        path = filedialog.askopenfilename(
            title="Choisir une archive à importer",
            filetypes=[("Archive ZIP", "*.zip")]
        )
        if not path:
            return

        mode = messagebox.askyesnocancel(
            "Mode d'import",
            "Comment importer ces données ?\n\n"
            "Oui = Fusionner (ajouter aux données actuelles, sans rien supprimer)\n"
            "Non = Remplacer entièrement les données actuelles\n"
            "Annuler = ne rien faire"
        )
        if mode is None:
            return

        try:
            restore_from_zip(path, merge=bool(mode))
        except Exception as e:
            messagebox.showerror("Erreur", f"L'import a échoué :\n{e}")
            return

        self.app.refresh_recipes()
        self.app.refresh_ingredients()
        messagebox.showinfo("Import terminé", "Les données ont été importées avec succès.")
        self.destroy()


class ShoppingChecklistWindow(tk.Toplevel):
    """Affiche une liste de courses déjà calculée sous forme de cases à
    cocher, pour pointer les articles au fur et à mesure des courses."""

    def __init__(self, app, grouped_totals, title="Liste de courses"):
        super().__init__(app)
        self.app = app
        self.title(f"☑️ {title}")
        self.geometry("480x600")
        self.grab_set()

        ttk.Label(self, text=f"☑️ {title}", font=("Segoe UI", 14, "bold")).pack(pady=(15, 5))
        ttk.Label(self, text="Cochez chaque article au fur et à mesure de vos courses.",
                  font=("Segoe UI", 9), foreground=COLOR_TEXT_MUTED).pack(pady=(0, 10))

        container = ttk.Frame(self)
        container.pack(fill="both", expand=True, padx=15)
        canvas = tk.Canvas(container, highlightthickness=0)
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        rows_frame = ttk.Frame(canvas)
        rows_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=rows_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.checks = []
        for rayon, items in grouped_totals:
            ttk.Label(rows_frame, text=rayon, font=("Segoe UI", 11, "bold")).pack(
                anchor="w", pady=(12, 3))
            for name, qty, unit in items:
                unit_display = f" {unit}" if unit else ""
                var = tk.BooleanVar()
                lbl_text = f"{name} : {qty}{unit_display}"
                chk = ttk.Checkbutton(rows_frame, text=lbl_text, variable=var,
                                       command=lambda: None)
                chk.pack(anchor="w", padx=10, pady=1)
                self.checks.append((var, chk, lbl_text))
                var.trace_add("write", lambda *args, v=var, c=chk: self._update_style(v, c))

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="☑️ Tout cocher", command=self.check_all).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="⬜ Tout décocher", command=self.uncheck_all).grid(row=0, column=1, padx=5)

        self.progress_label = ttk.Label(self, text="", font=("Segoe UI", 9), foreground=COLOR_TEXT_MUTED)
        self.progress_label.pack(pady=(0, 10))
        self._update_progress()

    def _update_style(self, var, chk):
        style_name = "Checked.TCheckbutton" if var.get() else "TCheckbutton"
        try:
            style = ttk.Style(self)
            style.configure("Checked.TCheckbutton", foreground="#999")
            chk.configure(style=style_name)
        except tk.TclError:
            pass
        self._update_progress()

    def _update_progress(self):
        total = len(self.checks)
        done = sum(1 for var, chk, text in self.checks if var.get())
        self.progress_label.config(text=f"{done} / {total} article(s) coché(s)")

    def check_all(self):
        for var, chk, text in self.checks:
            var.set(True)

    def uncheck_all(self):
        for var, chk, text in self.checks:
            var.set(False)


class AllRecipesWindow(tk.Toplevel):
    """Fenêtre listant toutes les recettes avec sélection + nombre de personnes,
    pour calculer et exporter en PDF la quantité totale d'ingrédients nécessaire."""

    SORT_OPTIONS = RECIPE_SORT_OPTIONS

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Toutes les recettes - Liste de courses")
        self.geometry("1220x720")
        self.minsize(720, 500)
        self.resizable(True, True)
        self.grab_set()
        self.last_result = None  # dernier résultat calculé, pour les exports/impression

        ttk.Label(self, text="Sélectionnez les recettes et le nombre de personnes :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(10, 5))

        top_frame = ttk.Frame(self)
        top_frame.pack(pady=(0, 5), fill="x", padx=15)
        ttk.Label(top_frame, text="🔍 Rechercher :").pack(side="left")
        self.search_entry = ttk.Entry(top_frame, width=20)
        self.search_entry.pack(side="left", padx=5, fill="x", expand=True)
        self.search_entry.bind("<KeyRelease>", lambda e: self._filter_rows())
        ttk.Label(top_frame, text="Trier par :").pack(side="left", padx=(10, 2))
        self.sort_combo = ttk.Combobox(top_frame, values=self.SORT_OPTIONS, state="readonly", width=18)
        self.sort_combo.set(self.SORT_OPTIONS[0])
        self.sort_combo.pack(side="left")
        self.sort_combo.bind("<<ComboboxSelected>>", lambda e: self._apply_sort())

        ingredient_filter_frame = ttk.LabelFrame(self, text="Filtrer par ingrédient")
        ingredient_filter_frame.pack(pady=(0, 8), padx=15, fill="x")
        ingredient_values = sorted(self.app.ingredient_names, key=ingredient_sort_key)

        ttk.Label(ingredient_filter_frame, text="Je veux :").grid(row=0, column=0, sticky="w", padx=5, pady=3)
        self.want_entries = []
        for i in range(2):
            entry = self._make_ingredient_filter_entry(ingredient_filter_frame, ingredient_values)
            entry.grid(row=0, column=1 + i, padx=5, pady=3)
            self.want_entries.append(entry)

        ttk.Label(ingredient_filter_frame, text="Je ne veux pas :").grid(row=1, column=0, sticky="w", padx=5, pady=3)
        self.exclude_entries = []
        for i in range(2):
            entry = self._make_ingredient_filter_entry(ingredient_filter_frame, ingredient_values)
            entry.grid(row=1, column=1 + i, padx=5, pady=3)
            self.exclude_entries.append(entry)

        ttk.Button(ingredient_filter_frame, text="Réinitialiser",
                   command=self._reset_ingredient_filters).grid(row=0, column=3, rowspan=2, padx=8)
        ttk.Label(ingredient_filter_frame, text="Tapez les premières lettres pour filtrer la liste.",
                  font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED).grid(
            row=2, column=0, columnspan=4, sticky="w", padx=5, pady=(0, 3))

        container = ttk.Frame(self)
        container.pack(fill="both", expand=True, padx=10)

        canvas = tk.Canvas(container, height=240, highlightthickness=0)
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        rows_frame = ttk.Frame(canvas)
        rows_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=rows_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.checks = []
        for row_index, recipe in enumerate(self.app.recipes):
            row = ttk.Frame(rows_frame)
            row.grid(row=row_index, column=0, sticky="ew", pady=4)
            var = tk.BooleanVar()
            chk = ttk.Checkbutton(row, text=format_recipe_list_label(recipe),
                                   variable=var, width=110)
            chk.grid(row=0, column=0, sticky="w")
            ttk.Label(row, text="Nb. personnes :").grid(row=0, column=1)
            pers_entry = ttk.Entry(row, width=5)

            # Préremplit à partir d'une sélection faite depuis "Voir une
            # recette précise" (bouton "Ajouter à la liste de courses"),
            # sinon utilise le nombre de personnes par défaut de la recette.
            preselected = self.app.shopping_selection.get(recipe["name"])
            if preselected is not None:
                var.set(True)
                pers_entry.insert(0, str(preselected))
            else:
                pers_entry.insert(0, str(recipe.get("default_persons") or 1))
            pers_entry.grid(row=0, column=2, padx=5)
            self.checks.append((var, recipe, pers_entry, row))

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="Calculer la liste de courses",
                   command=self.compute).grid(row=0, column=0, padx=5, pady=3, columnspan=4)
        ttk.Button(btn_frame, text="📝 Exporter en texte",
                   command=self.export_txt).grid(row=1, column=0, padx=5, pady=3)
        ttk.Button(btn_frame, text="📊 Exporter en Excel",
                   command=self.export_excel).grid(row=1, column=1, padx=5, pady=3)
        ttk.Button(btn_frame, text="📄 Exporter en PDF",
                   command=self.export_pdf).grid(row=1, column=2, padx=5, pady=3)
        ttk.Button(btn_frame, text="🖨️ Imprimer",
                   command=self.print_shopping_list).grid(row=1, column=3, padx=5, pady=3)
        ttk.Button(btn_frame, text="☑️ Mode courses (cocher au fur et à mesure)",
                   command=self.open_checklist).grid(row=2, column=0, columnspan=4, pady=(5, 0))
        ttk.Button(btn_frame, text="🗑 Vider la sélection",
                   command=self.clear_selection).grid(row=3, column=0, columnspan=4, pady=(5, 0))

        self.result_text = tk.Text(self, height=12, width=64)
        self.result_text.pack(pady=10)

    def _apply_sort(self):
        option = self.sort_combo.get()
        reverse = option in ("Ajoutées récemment",)
        ordered = sorted(self.checks, key=lambda t: recipe_sort_key(t[1], option), reverse=reverse)
        for new_index, (var, recipe, pers_entry, row) in enumerate(ordered):
            row.grid(row=new_index, column=0, sticky="ew", pady=4)
        self.checks = ordered
        self._filter_rows()

    def clear_selection(self):
        for var, recipe, pers_entry, row in self.checks:
            var.set(False)
        self.app.shopping_selection.clear()

    def _filter_rows(self):
        search = self.search_entry.get().strip()
        search_key = ingredient_sort_key(search) if search else ""

        known_keys = {ingredient_sort_key(n) for n in self.app.ingredient_names}

        def valid_typed_names(entries):
            names = []
            for e in entries:
                txt = e.get().strip()
                if txt and ingredient_sort_key(txt) in known_keys:
                    names.append(txt)
            return names

        want_names = valid_typed_names(self.want_entries)
        exclude_names = valid_typed_names(self.exclude_entries)
        want_keys = {ingredient_sort_key(n) for n in want_names}
        exclude_keys = {ingredient_sort_key(n) for n in exclude_names}

        for var, recipe, pers_entry, row in self.checks:
            if not recipe_matches_search(recipe, search_key):
                row.grid_remove()
                continue

            recipe_ing_keys = {ingredient_sort_key(ing["name"]) for ing in recipe["ingredients"]}
            if want_keys and not want_keys.issubset(recipe_ing_keys):
                row.grid_remove()
                continue
            if exclude_keys and (exclude_keys & recipe_ing_keys):
                row.grid_remove()
                continue

            row.grid()

    def _reset_ingredient_filters(self):
        for entry in self.want_entries + self.exclude_entries:
            entry.delete(0, tk.END)
        self._filter_rows()

    # ---- Autocomplétion des champs de filtre par ingrédient (même principe
    # que le choix d'ingrédient dans le formulaire de recette) ----

    def _make_ingredient_filter_entry(self, parent, values):
        entry = ttk.Entry(parent, width=22)
        entry.full_values = values
        entry.bind("<KeyRelease>", lambda e: self._on_filter_entry_keyrelease(e, entry))
        entry.bind("<FocusIn>", lambda e: self._on_filter_entry_focus_in(e, entry))
        entry.bind("<FocusOut>", lambda e: self._on_filter_entry_focus_out(e, entry))
        return entry

    def _hide_filter_suggestions(self, entry):
        popup = getattr(entry, "_suggestion_popup", None)
        if popup is not None:
            try:
                popup.destroy()
            except tk.TclError:
                pass
            entry._suggestion_popup = None
            entry._suggestion_listbox = None

    def _show_filter_suggestions(self, entry, filtered):
        self._hide_filter_suggestions(entry)
        if not filtered:
            return
        popup = tk.Toplevel(entry)
        popup.wm_overrideredirect(True)
        try:
            popup.wm_attributes("-topmost", True)
        except tk.TclError:
            pass
        x = entry.winfo_rootx()
        y = entry.winfo_rooty() + entry.winfo_height()
        width = max(entry.winfo_width(), 160)
        height = min(6, len(filtered)) * 20
        popup.wm_geometry(f"{width}x{height}+{x}+{y}")

        listbox = tk.Listbox(popup, height=min(6, len(filtered)), exportselection=False)
        listbox.pack(fill="both", expand=True)
        for v in filtered:
            listbox.insert(tk.END, v)

        def choose(event=None):
            sel = listbox.curselection()
            if sel:
                value = listbox.get(sel[0])
                entry.delete(0, tk.END)
                entry.insert(0, value)
            self._hide_filter_suggestions(entry)
            entry.focus_set()
            self._filter_rows()

        listbox.bind("<ButtonRelease-1>", choose)
        listbox.bind("<Return>", choose)
        entry._suggestion_popup = popup
        entry._suggestion_listbox = listbox

    def _on_filter_entry_keyrelease(self, event, entry):
        if event.keysym == "Down":
            listbox = getattr(entry, "_suggestion_listbox", None)
            if listbox is not None:
                listbox.focus_set()
                listbox.selection_set(0)
            return
        if event.keysym in ("Escape",):
            self._hide_filter_suggestions(entry)
            return
        if event.keysym in ("Return", "Tab", "Shift_L", "Shift_R", "Control_L", "Control_R",
                              "Caps_Lock", "Alt_L", "Alt_R", "Left", "Right"):
            return
        full_values = getattr(entry, "full_values", [])
        typed = entry.get()
        filtered = self._filter_ingredients(full_values, typed)
        if filtered:
            self._show_filter_suggestions(entry, filtered)
        else:
            self._hide_filter_suggestions(entry)

    def _on_filter_entry_focus_in(self, event, entry):
        full_values = getattr(entry, "full_values", [])
        typed = entry.get()
        filtered = self._filter_ingredients(full_values, typed)
        if filtered:
            self._show_filter_suggestions(entry, filtered)

    def _on_filter_entry_focus_out(self, event, entry):
        entry.after(200, lambda: self._hide_filter_suggestions(entry))
        self._filter_rows()

    @staticmethod
    def _filter_ingredients(full_values, typed):
        if not typed:
            return full_values
        typed_key = ingredient_sort_key(typed)
        filtered = [v for v in full_values if ingredient_sort_key(v).startswith(typed_key)]
        if not filtered:
            filtered = [v for v in full_values if typed_key in ingredient_sort_key(v)]
        return filtered

    def compute(self):
        chosen_pairs = []
        chosen_recipes = []
        for var, recipe, pers_entry, row in self.checks:
            if not var.get():
                continue
            try:
                persons = float(pers_entry.get().strip().replace(",", "."))
            except ValueError:
                messagebox.showerror("Erreur", f"Nombre de personnes invalide pour « {recipe['name']} ».")
                return None
            chosen_pairs.append((recipe, persons))
            chosen_recipes.append((recipe["name"], persons))

        if not chosen_recipes:
            messagebox.showinfo("Info", "Sélectionnez au moins une recette.")
            return None

        grouped_totals = compute_grouped_totals(chosen_pairs)

        self.result_text.delete("1.0", tk.END)
        self.result_text.insert(tk.END, "=== Liste de courses totale ===\n")
        for rayon, items in grouped_totals:
            self.result_text.insert(tk.END, f"\n--- {rayon} ---\n")
            for name, qty, unit in items:
                unit_display = f" {unit}" if unit else ""
                self.result_text.insert(tk.END, f"- {name} : {qty}{unit_display}\n")

        self.last_result = (chosen_recipes, grouped_totals)
        return self.last_result

    def export_txt(self):
        result = self.compute()  # recalcule à partir de la sélection actuelle
        if result is None:
            return
        chosen_recipes, grouped_totals = result

        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses en texte",
            defaultextension=".txt",
            filetypes=[("Fichier texte", "*.txt")],
            initialfile="liste_de_courses.txt"
        )
        if not path:
            return
        try:
            write_shopping_list_txt(path, "Liste de courses", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste de courses enregistrée :\n{path}")

    def export_excel(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'export Excel nécessite le module 'openpyxl'.\n"
                "Installez-le avec : pip install openpyxl"
            )
            return

        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result

        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses en Excel",
            defaultextension=".xlsx",
            filetypes=[("Fichier Excel", "*.xlsx")],
            initialfile="liste_de_courses.xlsx"
        )
        if not path:
            return

        try:
            wb = build_shopping_list_workbook(chosen_recipes, grouped_totals)
            wb.save(path)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste de courses enregistrée :\n{path}")

    def export_pdf(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'export PDF nécessite le module 'reportlab'.\n"
                "Installez-le avec : pip install reportlab"
            )
            return

        result = self.compute()  # recalcule à partir de la sélection actuelle
        if result is None:
            return
        chosen_recipes, grouped_totals = result

        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses en PDF",
            defaultextension=".pdf",
            filetypes=[("Fichier PDF", "*.pdf")],
            initialfile="liste_de_courses.pdf"
        )
        if not path:
            return

        try:
            build_shopping_list_pdf(path, "Liste de courses", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste de courses enregistrée :\n{path}")

    def print_shopping_list(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'impression nécessite le module 'reportlab' pour générer la mise en page.\n"
                "Installez-le avec : pip install reportlab"
            )
            return

        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result

        temp_path = get_temp_pdf_path("liste_de_courses")
        try:
            build_shopping_list_pdf(temp_path, "Liste de courses", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"La préparation de l'impression a échoué :\n{e}")
            return

        result_status = print_file(temp_path)
        report_print_result(result_status, temp_path, "la liste de courses")

    def open_checklist(self):
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        ShoppingChecklistWindow(self.app, grouped_totals, title="Liste de courses")


class OneRecipeWindow(tk.Toplevel):
    """Fenêtre pour afficher une recette précise (avec ses photos), avec
    quantités recalculées selon le nombre de personnes choisi."""

    def __init__(self, app, initial_recipe_name=None):
        super().__init__(app)
        self.app = app
        self.title("Voir une recette")
        self.geometry("580x820")
        self.grab_set()
        self._gallery_thumb_refs = []
        self.current_recipe = None
        self.filtered_indices = []

        ttk.Label(self, text="Choisissez une recette :", font=("Segoe UI", 11, "bold")).pack(pady=(10, 5))

        top_frame = ttk.Frame(self)
        top_frame.pack(pady=(0, 5), fill="x", padx=15)
        ttk.Label(top_frame, text="🔍 Rechercher :").pack(side="left")
        self.search_entry = ttk.Entry(top_frame, width=16)
        self.search_entry.pack(side="left", padx=5, fill="x", expand=True)
        self.search_entry.bind("<KeyRelease>", lambda e: self._populate())
        ttk.Label(top_frame, text="Trier :").pack(side="left", padx=(5, 2))
        self.sort_combo = ttk.Combobox(top_frame, values=RECIPE_SORT_OPTIONS, state="readonly", width=16)
        self.sort_combo.set(RECIPE_SORT_OPTIONS[0])
        self.sort_combo.pack(side="left")
        self.sort_combo.bind("<<ComboboxSelected>>", lambda e: self._populate())

        list_frame = ttk.Frame(self)
        list_frame.pack(pady=5, padx=15, fill="both")
        self.listbox = tk.Listbox(list_frame, width=60, height=8)
        list_scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.listbox.yview)
        self.listbox.configure(yscrollcommand=list_scrollbar.set)
        self.listbox.pack(side="left", fill="both", expand=True)
        list_scrollbar.pack(side="right", fill="y")
        self._populate()

        # ---- Galerie de photos ----
        gallery_outer = ttk.Frame(self)
        gallery_outer.pack(fill="x", padx=15, pady=(10, 0))
        self.gallery_canvas = tk.Canvas(gallery_outer, height=140, highlightthickness=0)
        gallery_scrollbar = ttk.Scrollbar(gallery_outer, orient="horizontal",
                                           command=self.gallery_canvas.xview)
        self.gallery_frame = ttk.Frame(self.gallery_canvas)
        self.gallery_frame.bind(
            "<Configure>", lambda e: self.gallery_canvas.configure(scrollregion=self.gallery_canvas.bbox("all"))
        )
        self.gallery_canvas.create_window((0, 0), window=self.gallery_frame, anchor="nw")
        self.gallery_canvas.configure(xscrollcommand=gallery_scrollbar.set)
        self.gallery_canvas.pack(fill="x")
        gallery_scrollbar.pack(fill="x")

        # ---- Nombre de personnes + ajustement rapide ----
        persons_frame = ttk.Frame(self)
        persons_frame.pack(pady=(10, 5))
        ttk.Label(persons_frame, text="Nombre de personnes :").grid(row=0, column=0, columnspan=4, pady=(0, 5))
        self.pers_entry = ttk.Entry(persons_frame, width=8)
        self.pers_entry.insert(0, "1")
        self.pers_entry.grid(row=1, column=0, padx=3)
        ttk.Button(persons_frame, text="−1", width=4,
                   command=lambda: self._adjust_persons(delta=-1)).grid(row=1, column=1, padx=3)
        ttk.Button(persons_frame, text="+1", width=4,
                   command=lambda: self._adjust_persons(delta=1)).grid(row=1, column=2, padx=3)
        ttk.Button(persons_frame, text="÷2", width=4,
                   command=lambda: self._adjust_persons(factor=0.5)).grid(row=1, column=3, padx=3)
        ttk.Button(persons_frame, text="×2", width=4,
                   command=lambda: self._adjust_persons(factor=2)).grid(row=1, column=4, padx=3)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="Afficher la recette", command=self.show_recipe).grid(row=0, column=0, padx=5, pady=3)
        ttk.Button(btn_frame, text="📄 Exporter en PDF",
                   command=self.export_recipe_pdf).grid(row=0, column=1, padx=5, pady=3)
        ttk.Button(btn_frame, text="🖨️ Imprimer",
                   command=self.print_recipe).grid(row=0, column=2, padx=5, pady=3)
        ttk.Button(btn_frame, text="🛒 Ajouter à la liste de courses",
                   command=self.add_to_shopping_list).grid(row=1, column=0, columnspan=2, padx=5, pady=3)
        ttk.Button(btn_frame, text="🍳 J'ai cuisiné ça !",
                   command=self.mark_as_cooked).grid(row=1, column=2, padx=5, pady=3)
        ttk.Button(btn_frame, text="🖥️ Mode cuisine (plein écran)",
                   command=self.open_cooking_mode).grid(row=2, column=0, columnspan=2, padx=5, pady=3)
        ttk.Button(btn_frame, text="📱 QR Code",
                   command=self.show_qr_code).grid(row=2, column=2, padx=5, pady=3)
        ttk.Button(btn_frame, text="⏲️ Minuteurs",
                   command=self.open_timers).grid(row=3, column=0, columnspan=3, padx=5, pady=3)
        ttk.Button(btn_frame, text="📔 Journal de cuisine",
                   command=self.open_cook_log).grid(row=4, column=0, columnspan=3, padx=5, pady=3)

        self.result_text = tk.Text(self, height=16, width=52, wrap="word")
        self.result_text.pack(pady=5, padx=15, fill="both", expand=True)

        if initial_recipe_name:
            for row_index, idx in enumerate(self.filtered_indices):
                if self.app.recipes[idx]["name"] == initial_recipe_name:
                    self.listbox.selection_clear(0, tk.END)
                    self.listbox.selection_set(row_index)
                    self.listbox.see(row_index)
                    self.show_recipe()
                    break

    def _populate(self):
        self.listbox.delete(0, tk.END)
        self.filtered_indices = []
        search = self.search_entry.get().strip()
        search_key = ingredient_sort_key(search) if search else ""
        option = self.sort_combo.get()
        indexed = list(enumerate(self.app.recipes))
        indexed = [pair for pair in indexed if recipe_matches_search(pair[1], search_key)]
        reverse = option in ("Ajoutées récemment",)
        indexed.sort(key=lambda pair: recipe_sort_key(pair[1], option), reverse=reverse)
        for idx, recipe in indexed:
            self.listbox.insert(tk.END, format_recipe_list_label(recipe))
            self.filtered_indices.append(idx)

    def _refresh_gallery(self, recipe):
        for child in self.gallery_frame.winfo_children():
            child.destroy()
        self._gallery_thumb_refs = []

        images = get_recipe_images(recipe)
        if not images:
            ttk.Label(self.gallery_frame, text="(aucune photo)").pack(side="left", padx=10, pady=10)
            return

        for fname in images:
            thumb = load_thumbnail(fname, size=(160, 120))
            cell = ttk.Frame(self.gallery_frame)
            cell.pack(side="left", padx=5, pady=5)
            if thumb is not None:
                self._gallery_thumb_refs.append(thumb)
                ttk.Label(cell, image=thumb).pack()
            else:
                ttk.Label(cell, text="(aperçu indisponible)").pack()

    def _adjust_persons(self, factor=None, delta=None):
        try:
            current = float(self.pers_entry.get().strip().replace(",", "."))
        except ValueError:
            current = 1.0
        if delta is not None:
            current = current + delta
        if factor is not None:
            current = current * factor
        current = max(0.5, current)
        if current == int(current):
            current = int(current)
        self.pers_entry.delete(0, tk.END)
        self.pers_entry.insert(0, str(current))
        if self.current_recipe is not None:
            self._display_recipe(self.current_recipe)

    def show_recipe(self):
        sel = self.listbox.curselection()
        if not sel:
            messagebox.showinfo("Info", "Sélectionnez une recette dans la liste.")
            return
        idx = self.filtered_indices[sel[0]]
        recipe = self.app.recipes[idx]
        is_new_selection = self.current_recipe is not recipe
        self.current_recipe = recipe
        if is_new_selection:
            self.pers_entry.delete(0, tk.END)
            self.pers_entry.insert(0, str(recipe.get("default_persons", 1) or 1))
            record_recipe_view(recipe["name"])
        self._display_recipe(recipe)

    def add_to_shopping_list(self):
        if self.current_recipe is None:
            messagebox.showinfo("Info", "Affichez d'abord une recette avec « Afficher la recette ».")
            return
        try:
            persons = float(self.pers_entry.get().strip().replace(",", "."))
        except ValueError:
            messagebox.showerror("Erreur", "Nombre de personnes invalide.")
            return
        self.app.shopping_selection[self.current_recipe["name"]] = persons
        messagebox.showinfo(
            "Ajouté",
            f"« {self.current_recipe['name']} » ({persons} pers.) sera présélectionnée "
            "la prochaine fois que vous ouvrirez « Voir toutes les recettes »."
        )

    def mark_as_cooked(self):
        if self.current_recipe is None:
            messagebox.showinfo("Info", "Affichez d'abord une recette avec « Afficher la recette ».")
            return
        recipes = load_recipes()
        target_name = self.current_recipe.get("name")
        for r in recipes:
            if r is self.current_recipe or r.get("name") == target_name:
                r["times_cooked"] = r.get("times_cooked", 0) + 1
                cooked_dates = r.get("cooked_dates", [])
                cooked_dates.append(datetime.now().strftime("%Y-%m-%d"))
                r["cooked_dates"] = cooked_dates
                self.current_recipe = r
                break
        save_recipes(recipes)
        self.app.refresh_recipes()

        def _on_log_done(note, photo_filename):
            recipes2 = load_recipes()
            for r in recipes2:
                if r.get("name") == target_name:
                    cook_log = r.get("cook_log", [])
                    cook_log.append({
                        "date": datetime.now().strftime("%Y-%m-%d"),
                        "note": note,
                        "photo": photo_filename,
                    })
                    r["cook_log"] = cook_log
                    self.current_recipe = r
                    break
            save_recipes(recipes2)
            self.app.refresh_recipes()
            messagebox.showinfo("Marqué", f"« {target_name} » a été marquée comme cuisinée aujourd'hui !")

        CookLogEntryDialog(self.app, target_name, _on_log_done)

    def open_cook_log(self):
        if self.current_recipe is None:
            messagebox.showinfo("Info", "Affichez d'abord une recette avec « Afficher la recette ».")
            return
        CookLogWindow(self.app, self.current_recipe)

    def _display_recipe(self, recipe):
        try:
            persons = float(self.pers_entry.get().strip().replace(",", "."))
        except ValueError:
            messagebox.showerror("Erreur", "Nombre de personnes invalide.")
            return

        self._refresh_gallery(recipe)

        self.result_text.delete("1.0", tk.END)
        cat = recipe.get("category", "Autre")
        star = "⭐ " if recipe.get("favorite") else ""
        self.result_text.insert(tk.END, f"=== {star}[{cat}] {recipe['name']} ({persons} pers.) ===\n\n")

        rating = recipe.get("rating", 0)
        if rating:
            self.result_text.insert(tk.END, f"Note : {rating_stars(rating)}\n\n")

        info_bits = []
        if recipe.get("prep_time"):
            info_bits.append(f"Préparation : {recipe['prep_time']} min")
        if recipe.get("cook_time"):
            info_bits.append(f"Cuisson : {recipe['cook_time']} min")
        if recipe.get("difficulty"):
            info_bits.append(f"Difficulté : {recipe['difficulty']}")
        if info_bits:
            self.result_text.insert(tk.END, " | ".join(info_bits) + "\n\n")

        allergens = recipe.get("allergens") or []
        if allergens:
            self.result_text.insert(tk.END, f"⚠ Allergènes : {', '.join(allergens)}\n\n")

        for ing in recipe["ingredients"]:
            qty = round(ing["quantity"] * persons, 2)
            unit = f" {ing['unit']}" if ing["unit"] else ""
            self.result_text.insert(tk.END, f"- {ing['name'].capitalize()} : {qty}{unit}\n")

        cost, cost_known, cost_total = compute_recipe_cost(recipe, persons)
        if cost_known:
            partial = "" if cost_known == cost_total else f" (estimation partielle, {cost_known}/{cost_total} ingrédients avec prix connu)"
            self.result_text.insert(tk.END, f"\n💰 Coût estimé : {cost:.2f} €{partial}\n")

        nutrition, nutri_known, nutri_total = compute_recipe_nutrition(recipe, persons)
        if nutri_known:
            partial = "" if nutri_known == nutri_total else f" (estimation partielle, {nutri_known}/{nutri_total} ingrédients reconnus)"
            self.result_text.insert(
                tk.END,
                f"🥗 Valeurs nutritionnelles estimées{partial} :\n"
                f"   {nutrition['kcal']:.0f} kcal · {nutrition['protein_g']:.0f} g protéines · "
                f"{nutrition['carbs_g']:.0f} g glucides · {nutrition['fat_g']:.0f} g lipides\n"
            )

        description = recipe.get("description", "").strip()
        if description:
            self.result_text.insert(tk.END, f"\n--- Description ---\n{description}\n")
        personal_notes = recipe.get("personal_notes", "").strip()
        if personal_notes:
            self.result_text.insert(tk.END, f"\n--- Notes personnelles ---\n{personal_notes}\n")

    @staticmethod
    def _build_recipe_pdf(path, recipe, persons):
        c = pdf_canvas.Canvas(path, pagesize=A4)
        width, height = A4
        draw_recipe_content(c, recipe, persons, width, height)
        c.save()

    def export_recipe_pdf(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'export PDF nécessite le module 'reportlab'.\n"
                "Installez-le avec : pip install reportlab"
            )
            return
        if self.current_recipe is None:
            messagebox.showinfo("Info", "Affichez d'abord une recette avec « Afficher la recette ».")
            return
        try:
            persons = float(self.pers_entry.get().strip().replace(",", "."))
        except ValueError:
            messagebox.showerror("Erreur", "Nombre de personnes invalide.")
            return

        recipe = self.current_recipe
        path = filedialog.asksaveasfilename(
            title="Exporter la recette en PDF",
            defaultextension=".pdf",
            filetypes=[("Fichier PDF", "*.pdf")],
            initialfile=f"{recipe['name']}.pdf"
        )
        if not path:
            return

        try:
            self._build_recipe_pdf(path, recipe, persons)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Recette exportée :\n{path}")

    def print_recipe(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'impression nécessite le module 'reportlab' pour générer la mise en page.\n"
                "Installez-le avec : pip install reportlab"
            )
            return
        if self.current_recipe is None:
            messagebox.showinfo("Info", "Affichez d'abord une recette avec « Afficher la recette ».")
            return
        try:
            persons = float(self.pers_entry.get().strip().replace(",", "."))
        except ValueError:
            messagebox.showerror("Erreur", "Nombre de personnes invalide.")
            return

        recipe = self.current_recipe
        temp_path = get_temp_pdf_path("recette")
        try:
            self._build_recipe_pdf(temp_path, recipe, persons)
        except Exception as e:
            messagebox.showerror("Erreur", f"La préparation de l'impression a échoué :\n{e}")
            return

        result_status = print_file(temp_path)
        report_print_result(result_status, temp_path, f"« {recipe['name']} »")

    def show_qr_code(self):
        if not QRCODE_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'export en QR code nécessite le module 'qrcode'.\n"
                "Installez-le avec : pip install qrcode"
            )
            return
        if not PIL_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'export en QR code nécessite aussi le module 'Pillow'.\n"
                "Installez-le avec : pip install pillow"
            )
            return
        if self.current_recipe is None:
            messagebox.showinfo("Info", "Affichez d'abord une recette avec « Afficher la recette ».")
            return
        try:
            persons = float(self.pers_entry.get().strip().replace(",", "."))
        except ValueError:
            messagebox.showerror("Erreur", "Nombre de personnes invalide.")
            return
        QRCodeWindow(self.app, self.current_recipe, persons)

    def open_timers(self):
        recipe = self.current_recipe
        label_text = recipe["name"] if recipe else "Minuteur"
        default_minutes = 10
        if recipe is not None:
            try:
                if recipe.get("cook_time"):
                    default_minutes = max(1, round(float(recipe["cook_time"])))
                elif recipe.get("prep_time"):
                    default_minutes = max(1, round(float(recipe["prep_time"])))
            except (TypeError, ValueError):
                pass

        # Cette fenêtre est modale (grab_set), ce qui bloquerait toute autre
        # fenêtre de l'application — y compris la fenêtre des minuteurs, qui
        # doit rester utilisable en même temps qu'on consulte la recette.
        # On relâche donc le grab ici ; les deux fenêtres restent ensuite
        # utilisables librement en parallèle.
        try:
            self.grab_release()
        except tk.TclError:
            pass

        # Une seule fenêtre de minuteurs pour toute l'application : si elle
        # est déjà ouverte, on y ajoute simplement un minuteur de plus au
        # lieu d'en ouvrir une deuxième.
        existing = getattr(self.app, "timers_window", None)
        if existing is not None and existing.winfo_exists():
            existing.add_timer(label_text, default_minutes)
            existing.deiconify()
            existing.lift()
            existing.focus_force()
        else:
            self.app.timers_window = TimersWindow(self.app, label_text, default_minutes)

    def open_cooking_mode(self):
        if self.current_recipe is None:
            messagebox.showinfo("Info", "Affichez d'abord une recette avec « Afficher la recette ».")
            return
        try:
            persons = float(self.pers_entry.get().strip().replace(",", "."))
        except ValueError:
            messagebox.showerror("Erreur", "Nombre de personnes invalide.")
            return
        # CookingModeWindow n'a pas son propre grab_set() : sans relâcher
        # celui de cette fenêtre, le mode cuisine serait inutilisable (même
        # pas fermable) tant que "Voir une recette précise" reste ouverte.
        try:
            self.grab_release()
        except tk.TclError:
            pass
        CookingModeWindow(self.app, self.current_recipe, persons)


class CookingModeWindow(tk.Toplevel):
    """Affichage plein écran, en gros caractères et sans menus, d'une
    recette — pratique à consulter en cuisinant, posé à côté des
    fourneaux."""

    def __init__(self, app, recipe, persons):
        super().__init__(app)
        self.app = app
        self.recipe = recipe
        self.persons = persons
        self.title(f"Mode cuisine — {recipe['name']}")
        self.configure(bg="white")
        self._is_fullscreen = False
        # Une fenêtre maximisée (plutôt qu'un vrai plein écran) s'ouvre
        # instantanément : le vrai plein écran provoque, sur certains
        # systèmes Windows, un blocage de plusieurs secondes le temps que la
        # transition d'affichage se fasse. Le plein écran natif reste
        # disponible via F11 pour qui le souhaite.
        try:
            self.state("zoomed")
        except tk.TclError:
            self.geometry(f"{self.winfo_screenwidth()}x{self.winfo_screenheight()}+0+0")
        self.bind("<Escape>", lambda e: self.destroy())
        self.bind("<F11>", lambda e: self._toggle_fullscreen())

        top_bar = tk.Frame(self, bg="white")
        top_bar.pack(fill="x", pady=10)
        tk.Button(top_bar, text="✕ Fermer (Échap)", font=("Segoe UI", 13),
                  command=self.destroy).pack(side="right", padx=30)
        tk.Label(top_bar, text="F11 : plein écran", font=("Segoe UI", 9),
                 bg="white", fg="#999").pack(side="right", padx=10)

        pers_frame = tk.Frame(top_bar, bg="white")
        pers_frame.pack(side="left", padx=30)
        tk.Button(pers_frame, text="−", font=("Segoe UI", 14, "bold"), width=3,
                  command=lambda: self._adjust(-1)).pack(side="left")
        self.pers_label = tk.Label(pers_frame, text=f"{self._fmt(persons)} pers.",
                                    font=("Segoe UI", 14), bg="white")
        self.pers_label.pack(side="left", padx=10)
        tk.Button(pers_frame, text="+", font=("Segoe UI", 14, "bold"), width=3,
                  command=lambda: self._adjust(1)).pack(side="left")

        outer = tk.Frame(self, bg="white")
        outer.pack(fill="both", expand=True)
        self.canvas = tk.Canvas(outer, bg="white", highlightthickness=0)
        scrollbar = ttk.Scrollbar(outer, orient="vertical", command=self.canvas.yview)
        self.content = tk.Frame(self.canvas, bg="white")
        self.content.bind("<Configure>", lambda e: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.create_window((0, 0), window=self.content, anchor="n")
        self.canvas.configure(yscrollcommand=scrollbar.set)
        self.canvas.pack(side="left", fill="both", expand=True, padx=60)
        scrollbar.pack(side="right", fill="y")

        self._render()

    def _toggle_fullscreen(self):
        self._is_fullscreen = not self._is_fullscreen
        try:
            self.attributes("-fullscreen", self._is_fullscreen)
        except tk.TclError:
            self._is_fullscreen = False

    @staticmethod
    def _fmt(value):
        return int(value) if value == int(value) else value

    def _adjust(self, delta):
        self.persons = max(0.5, self.persons + delta)
        self.pers_label.config(text=f"{self._fmt(self.persons)} pers.")
        self._render()

    def _render(self):
        for child in self.content.winfo_children():
            child.destroy()

        recipe = self.recipe
        star = "⭐ " if recipe.get("favorite") else ""
        tk.Label(self.content, text=f"{star}{recipe['name']}", font=("Segoe UI", 34, "bold"),
                 bg="white", wraplength=1000, justify="center").pack(pady=(10, 5))

        info_bits = []
        if recipe.get("prep_time"):
            info_bits.append(f"Préparation : {recipe['prep_time']} min")
        if recipe.get("cook_time"):
            info_bits.append(f"Cuisson : {recipe['cook_time']} min")
        if recipe.get("difficulty"):
            info_bits.append(f"Difficulté : {recipe['difficulty']}")
        if info_bits:
            tk.Label(self.content, text="   |   ".join(info_bits), font=("Segoe UI", 16),
                     bg="white", fg="#555").pack(pady=(0, 20))

        tk.Label(self.content, text="Ingrédients", font=("Segoe UI", 22, "bold"),
                 bg="white").pack(pady=(10, 8), anchor="w", fill="x")
        for ing in recipe["ingredients"]:
            qty = round(ing["quantity"] * self.persons, 2)
            if qty == int(qty):
                qty = int(qty)
            unit = f" {ing['unit']}" if ing["unit"] else ""
            tk.Label(self.content, text=f"•  {ing['name'].capitalize()} : {qty}{unit}",
                     font=("Segoe UI", 18), bg="white", anchor="w", justify="left",
                     wraplength=1000).pack(fill="x", pady=3, anchor="w")

        description = recipe.get("description", "").strip()
        if description:
            tk.Label(self.content, text="Préparation", font=("Segoe UI", 22, "bold"),
                     bg="white").pack(pady=(25, 8), anchor="w", fill="x")
            tk.Label(self.content, text=description, font=("Segoe UI", 16), bg="white",
                     justify="left", anchor="w", wraplength=1000).pack(fill="x", anchor="w")

        personal_notes = recipe.get("personal_notes", "").strip()
        if personal_notes:
            tk.Label(self.content, text="Notes personnelles", font=("Segoe UI", 20, "bold"),
                     bg="white", fg="#555").pack(pady=(20, 8), anchor="w", fill="x")
            tk.Label(self.content, text=personal_notes, font=("Segoe UI", 14), bg="white",
                     fg="#555", justify="left", anchor="w", wraplength=1000).pack(fill="x", anchor="w")

        tk.Label(self.content, text="", bg="white").pack(pady=30)  # marge basse


class IngredientSearchWindow(tk.Toplevel):
    """Recherche inversée : à partir d'un ingrédient, retrouver toutes les
    recettes qui l'utilisent."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Recherche par ingrédient")
        self.geometry("480x600")
        self.grab_set()

        ttk.Label(self, text="Quel ingrédient recherchez-vous ?",
                  font=("Segoe UI", 11, "bold")).pack(pady=(10, 5))

        search_frame = ttk.Frame(self)
        search_frame.pack(fill="x", padx=15, pady=(0, 5))
        ttk.Label(search_frame, text="🔍").pack(side="left")
        self.search_entry = ttk.Entry(search_frame)
        self.search_entry.pack(side="left", fill="x", expand=True, padx=5)
        self.search_entry.bind("<KeyRelease>", lambda e: self._populate_ingredients())

        list_frame = ttk.Frame(self)
        list_frame.pack(fill="both", padx=15, pady=5)
        self.ing_listbox = tk.Listbox(list_frame, height=10)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=self.ing_listbox.yview)
        self.ing_listbox.configure(yscrollcommand=scrollbar.set)
        self.ing_listbox.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")
        self.ing_listbox.bind("<Double-Button-1>", lambda e: self.search_recipes())
        self._populate_ingredients()

        ttk.Button(self, text="🔍 Voir les recettes qui l'utilisent",
                   command=self.search_recipes).pack(pady=8)

        self.result_text = tk.Text(self, height=14, width=54, wrap="word")
        self.result_text.pack(pady=5, padx=15, fill="both", expand=True)

    def _populate_ingredients(self):
        search = self.search_entry.get().strip()
        search_key = ingredient_sort_key(search) if search else ""
        self.ing_listbox.delete(0, tk.END)
        for name in self.app.ingredient_names:
            if search_key and search_key not in ingredient_sort_key(name):
                continue
            self.ing_listbox.insert(tk.END, name)

    def search_recipes(self):
        sel = self.ing_listbox.curselection()
        if not sel:
            messagebox.showinfo("Info", "Sélectionnez un ingrédient dans la liste.")
            return
        target_name = self.ing_listbox.get(sel[0])
        target_key = ingredient_sort_key(target_name)

        matches = []
        for recipe in self.app.recipes:
            for ing in recipe["ingredients"]:
                if ingredient_sort_key(ing["name"]) == target_key:
                    matches.append((recipe, ing))
                    break

        self.result_text.delete("1.0", tk.END)
        if not matches:
            self.result_text.insert(tk.END, f"Aucune recette n'utilise « {target_name} » pour le moment.\n")
            return

        self.result_text.insert(tk.END, f"Recettes utilisant « {target_name} » ({len(matches)}) :\n\n")
        for recipe, ing in matches:
            star = "⭐ " if recipe.get("favorite") else ""
            cat = recipe.get("category", "Autre")
            qty = ing["quantity"]
            if qty == int(qty):
                qty = int(qty)
            unit = f" {ing['unit']}" if ing["unit"] else ""
            self.result_text.insert(
                tk.END, f"- {star}[{cat}] {recipe['name']} ({qty}{unit} pour 1 personne)\n"
            )


class TimerRow(tk.Frame):
    """Un minuteur réglable et indépendant, affiché comme une ligne à
    l'intérieur de TimersWindow. Quand il arrive à zéro, la ligne clignote
    en rouge et un signal sonore retentit jusqu'à ce que l'utilisateur
    interagisse avec elle (démarrer, réinitialiser, ou simplement cliquer
    dessus)."""

    def __init__(self, parent, timers_window, label="Minuteur", minutes=10):
        super().__init__(parent, background=COLOR_CARD, highlightbackground=COLOR_BORDER,
                          highlightthickness=1)
        self.timers_window = timers_window
        self.remaining_seconds = max(0, int(minutes) * 60)
        self.running = False
        self.finished = False
        self._after_id = None
        self._flash_after_id = None
        self._flash_on = False

        top_row = tk.Frame(self, background=COLOR_CARD)
        top_row.pack(fill="x", padx=8, pady=(8, 2))
        self.name_entry = ttk.Entry(top_row, width=18)
        self.name_entry.insert(0, label)
        self.name_entry.pack(side="left")
        self.display_label = tk.Label(top_row, text=self._format_time(), font=("Segoe UI", 18, "bold"),
                                       background=COLOR_CARD, foreground=COLOR_ACCENT_DARK, width=7)
        self.display_label.pack(side="left", padx=10)
        ttk.Button(top_row, text="🗑", width=3,
                   command=lambda: self.timers_window.remove_timer(self)).pack(side="right")

        bottom_row = tk.Frame(self, background=COLOR_CARD)
        bottom_row.pack(fill="x", padx=8, pady=(0, 8))
        ttk.Label(bottom_row, text="Min :", style="Card.TLabel").pack(side="left")
        self.minutes_entry = ttk.Entry(bottom_row, width=4)
        self.minutes_entry.insert(0, str(minutes))
        self.minutes_entry.pack(side="left", padx=(2, 8))
        ttk.Label(bottom_row, text="Sec :", style="Card.TLabel").pack(side="left")
        self.seconds_entry = ttk.Entry(bottom_row, width=4)
        self.seconds_entry.insert(0, "0")
        self.seconds_entry.pack(side="left", padx=(2, 8))

        self.start_button = ttk.Button(bottom_row, text="▶️", width=3, command=self.start)
        self.start_button.pack(side="left", padx=2)
        self.pause_button = ttk.Button(bottom_row, text="⏸️", width=3, command=self.pause, state="disabled")
        self.pause_button.pack(side="left", padx=2)
        ttk.Button(bottom_row, text="🔄", width=3, command=self.reset).pack(side="left", padx=2)

        # Cliquer n'importe où sur la ligne (ou sur le gros affichage du
        # temps) fait taire l'alarme si le minuteur est terminé.
        self.bind("<Button-1>", lambda e: self._stop_flash())
        self.display_label.bind("<Button-1>", lambda e: self._stop_flash())

        self._refresh_display()

    def _format_time(self):
        mins, secs = divmod(max(0, self.remaining_seconds), 60)
        return f"{mins:02d}:{secs:02d}"

    def _refresh_display(self):
        self.display_label.config(text=self._format_time())

    def start(self):
        if self.finished:
            self._stop_flash()
        if self.running:
            return
        if self.remaining_seconds <= 0:
            try:
                minutes = int(self.minutes_entry.get().strip() or 0)
                seconds = int(self.seconds_entry.get().strip() or 0)
            except ValueError:
                messagebox.showerror("Erreur", "Durée invalide.")
                return
            self.remaining_seconds = max(0, minutes * 60 + seconds)
            if self.remaining_seconds <= 0:
                messagebox.showinfo("Info", "Réglez une durée avant de démarrer.")
                return
        self.running = True
        self.minutes_entry.config(state="disabled")
        self.seconds_entry.config(state="disabled")
        self.start_button.config(state="disabled")
        self.pause_button.config(state="normal")
        self._tick()

    def _tick(self):
        if not self.running:
            return
        if self.remaining_seconds <= 0:
            self._on_finished()
            return
        self.remaining_seconds -= 1
        self._refresh_display()
        self._after_id = self.after(1000, self._tick)

    def pause(self):
        self.running = False
        if self._after_id is not None:
            try:
                self.after_cancel(self._after_id)
            except Exception:
                pass
            self._after_id = None
        self.start_button.config(state="normal")
        self.pause_button.config(state="disabled")

    def reset(self):
        self.running = False
        if self._after_id is not None:
            try:
                self.after_cancel(self._after_id)
            except Exception:
                pass
            self._after_id = None
        self._stop_flash()
        self.minutes_entry.config(state="normal")
        self.seconds_entry.config(state="normal")
        self.start_button.config(state="normal")
        self.pause_button.config(state="disabled")
        try:
            minutes = int(self.minutes_entry.get().strip() or 0)
            seconds = int(self.seconds_entry.get().strip() or 0)
        except ValueError:
            minutes, seconds = 0, 0
        self.remaining_seconds = max(0, minutes * 60 + seconds)
        self._refresh_display()

    def _on_finished(self):
        self.running = False
        self.finished = True
        self._after_id = None
        self.pause_button.config(state="disabled")
        self.start_button.config(state="normal")
        self.minutes_entry.config(state="normal")
        self.seconds_entry.config(state="normal")
        self._refresh_display()
        self._start_flash()

    def _set_children_bg(self, widget, color):
        for child in widget.winfo_children():
            if isinstance(child, tk.Frame):
                child.configure(background=color)
                self._set_children_bg(child, color)
            elif isinstance(child, tk.Label):
                child.configure(background=color)

    def _start_flash(self):
        self._flash_on = False
        self._flash_step()

    def _flash_step(self):
        if not self.finished:
            return
        self._flash_on = not self._flash_on
        color = COLOR_ERROR if self._flash_on else COLOR_CARD
        text_color = "white" if self._flash_on else COLOR_ACCENT_DARK
        self.configure(background=color)
        self._set_children_bg(self, color)
        self.display_label.config(background=color, foreground=text_color)
        try:
            self.bell()
        except tk.TclError:
            pass
        self._flash_after_id = self.after(500, self._flash_step)

    def _stop_flash(self):
        if not self.finished and self._flash_after_id is None:
            return
        self.finished = False
        if self._flash_after_id is not None:
            try:
                self.after_cancel(self._flash_after_id)
            except Exception:
                pass
            self._flash_after_id = None
        self.configure(background=COLOR_CARD)
        self._set_children_bg(self, COLOR_CARD)
        self.display_label.config(background=COLOR_CARD, foreground=COLOR_ACCENT_DARK)

    def cancel(self):
        """Arrête tout minuterie/clignotement en cours (appelé à la
        fermeture de la fenêtre ou à la suppression de cette ligne)."""
        if self._after_id is not None:
            try:
                self.after_cancel(self._after_id)
            except Exception:
                pass
        if self._flash_after_id is not None:
            try:
                self.after_cancel(self._flash_after_id)
            except Exception:
                pass


class CookLogEntryDialog(tk.Toplevel):
    """Petite fenêtre pour ajouter, juste après avoir marqué une recette
    comme cuisinée, une note et/ou une photo optionnelles au journal de
    cuisine de cette recette."""

    def __init__(self, app, recipe_name, on_done):
        super().__init__(app)
        self.app = app
        self.on_done = on_done
        self.photo_path = None
        self.title("📔 Ajouter au journal de cuisine")
        self.geometry("420x400")
        self.resizable(False, False)
        self.grab_set()

        ttk.Label(self, text=f"🍳 « {recipe_name} »", font=("Segoe UI", 12, "bold"),
                  wraplength=380, justify="center").pack(pady=(15, 2))
        ttk.Label(self, text="Comment était-ce ? Une note et/ou une photo\n(facultatif, vous pouvez aussi passer directement).",
                  font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center").pack(pady=(0, 10))

        self.note_text = tk.Text(self, height=7, width=42, wrap="word")
        self.note_text.pack(padx=15, pady=(0, 10))

        photo_frame = ttk.Frame(self)
        photo_frame.pack(pady=(0, 10))
        self.photo_label = ttk.Label(photo_frame, text="Aucune photo choisie", foreground=COLOR_TEXT_MUTED)
        self.photo_label.pack(side="left", padx=(0, 8))
        ttk.Button(photo_frame, text="📷 Choisir une photo", command=self.choose_photo).pack(side="left")

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=15)
        ttk.Button(btn_frame, text="💾 Enregistrer", command=self.save).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="Passer", style="Secondary.TButton", command=self.skip).grid(row=0, column=1, padx=5)

        self.protocol("WM_DELETE_WINDOW", self.skip)

    def choose_photo(self):
        path = filedialog.askopenfilename(
            title="Choisir une photo",
            filetypes=[("Images", "*.jpg *.jpeg *.png *.webp *.gif *.bmp")]
        )
        if path:
            self.photo_path = path
            self.photo_label.config(text=os.path.basename(path), foreground=COLOR_TEXT)

    def save(self):
        note = self.note_text.get("1.0", "end-1c").strip()
        photo_filename = copy_image_to_store(self.photo_path) if self.photo_path else None
        self.on_done(note, photo_filename)
        self.destroy()

    def skip(self):
        self.on_done("", None)
        self.destroy()


class CookLogWindow(tk.Toplevel):
    """Affiche l'historique des fois où une recette a été cuisinée, avec les
    notes et photos éventuellement ajoutées à chaque fois (la plus récente
    en premier)."""

    def __init__(self, app, recipe):
        super().__init__(app)
        self.app = app
        self.recipe = recipe
        self.title(f"📔 Journal de cuisine — {recipe['name']}")
        self.geometry("480x600")
        self.minsize(400, 400)
        self.resizable(True, True)
        self.grab_set()

        ttk.Label(self, text=f"📔 {recipe['name']}", font=("Segoe UI", 13, "bold"),
                  wraplength=440, justify="center").pack(pady=(15, 2))
        times_cooked = recipe.get("times_cooked", 0)
        ttk.Label(self, text=f"Cuisinée {times_cooked} fois au total",
                  font=("Segoe UI", 9), foreground=COLOR_TEXT_MUTED).pack(pady=(0, 10))

        container = ttk.Frame(self)
        container.pack(fill="both", expand=True, padx=15, pady=(0, 15))
        canvas = tk.Canvas(container, highlightthickness=0)
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        rows_frame = ttk.Frame(canvas)
        rows_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=rows_frame, anchor="n")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        canvas.bind("<Enter>", lambda e: canvas.bind_all("<MouseWheel>", _on_mousewheel))
        canvas.bind("<Leave>", lambda e: canvas.unbind_all("<MouseWheel>"))

        self._thumb_refs = []
        cook_log = list(recipe.get("cook_log", []))
        cook_log.sort(key=lambda e: e.get("date", ""), reverse=True)

        if not cook_log:
            ttk.Label(
                rows_frame,
                text="Aucune note enregistrée pour le moment.\n"
                     "Utilisez « 🍳 J'ai cuisiné ça ! » pour en ajouter une.",
                foreground=COLOR_TEXT_MUTED, justify="center"
            ).pack(pady=30)

        for entry in cook_log:
            entry_card = tk.Frame(rows_frame, background=COLOR_CARD, highlightbackground=COLOR_BORDER,
                                   highlightthickness=1)
            entry_card.pack(fill="x", pady=6, padx=2)
            try:
                date_display = datetime.fromisoformat(entry["date"]).strftime("%d/%m/%Y")
            except (KeyError, ValueError, TypeError):
                date_display = entry.get("date", "?")
            ttk.Label(entry_card, text=date_display, font=("Segoe UI", 10, "bold"),
                      style="Card.TLabel", foreground=COLOR_ACCENT_DARK).pack(anchor="w", padx=10, pady=(8, 2))

            photo_filename = entry.get("photo")
            if photo_filename:
                thumb = load_thumbnail(photo_filename, size=(220, 160))
                if thumb is not None:
                    self._thumb_refs.append(thumb)
                    tk.Label(entry_card, image=thumb, background=COLOR_CARD).pack(padx=10, pady=4)

            note = (entry.get("note") or "").strip()
            if note:
                ttk.Label(entry_card, text=note, style="Card.TLabel", wraplength=420,
                          justify="left").pack(anchor="w", padx=10, pady=(0, 8))
            else:
                ttk.Label(entry_card, text="(pas de note)", style="Card.TLabel",
                          foreground=COLOR_TEXT_MUTED, font=("Segoe UI", 8)).pack(anchor="w", padx=10, pady=(0, 8))


class TimersWindow(tk.Toplevel):
    """Fenêtre unique regroupant plusieurs minuteurs indépendants et
    réglables, pour chronométrer différentes étapes d'une recette en même
    temps (ex. un pour les pâtes, un pour la sauce...). Le bouton
    "➕ Ajouter un minuteur" empile un nouveau minuteur sous les précédents."""

    def __init__(self, app, initial_label="Minuteur", initial_minutes=10):
        super().__init__(app)
        self.app = app
        self.title("⏲️ Minuteurs")
        self.geometry("380x560")
        self.minsize(340, 300)
        self.resizable(True, True)
        # Reste visible au premier plan même par-dessus une autre fenêtre
        # maximisée (ex. le mode cuisine) : on veut toujours voir les
        # minuteurs en cours, quoi qu'on affiche par ailleurs.
        try:
            self.attributes("-topmost", True)
        except tk.TclError:
            pass

        ttk.Label(self, text="⏲️ Minuteurs", font=("Segoe UI", 13, "bold")).pack(pady=(15, 5))
        ttk.Label(
            self, text="Réglez chaque minuteur puis ▶️ pour le démarrer.\n"
                       "À la fin, la ligne clignote en rouge avec un signal sonore.",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center"
        ).pack(pady=(0, 8))

        container = ttk.Frame(self)
        container.pack(fill="both", expand=True, padx=10)
        canvas = tk.Canvas(container, highlightthickness=0)
        scrollbar = ttk.Scrollbar(container, orient="vertical", command=canvas.yview)
        self.rows_frame = ttk.Frame(canvas)
        self.rows_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.rows_frame, anchor="n")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.timer_rows = []
        self.add_timer(initial_label, initial_minutes)

        ttk.Button(self, text="➕ Ajouter un minuteur",
                   command=lambda: self.add_timer()).pack(pady=10)

        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def add_timer(self, label="Minuteur", minutes=10):
        row = TimerRow(self.rows_frame, self, label, minutes)
        row.pack(fill="x", pady=6, padx=4)
        self.timer_rows.append(row)
        self.update_idletasks()

    def remove_timer(self, row):
        row.cancel()
        row.destroy()
        if row in self.timer_rows:
            self.timer_rows.remove(row)

    def _on_close(self):
        for row in list(self.timer_rows):
            row.cancel()
        if getattr(self.app, "timers_window", None) is self:
            self.app.timers_window = None
        self.destroy()


class QRCodeWindow(tk.Toplevel):
    """Affiche une recette (nom + ingrédients) sous forme de QR code, à
    scanner avec un téléphone, et permet de l'enregistrer en image PNG."""

    MAX_CHARS = 800  # limite la taille du texte encodé pour rester scannable

    def __init__(self, app, recipe, persons):
        super().__init__(app)
        self.app = app
        self.recipe = recipe
        self.title(f"QR Code — {recipe['name']}")
        self.geometry("420x540")
        self.resizable(False, False)
        self.grab_set()

        text = self._build_text(recipe, persons)

        qr = qrcode.QRCode(border=2)
        qr.add_data(text)
        qr.make(fit=True)
        qr_img = qr.make_image(fill_color="black", back_color="white").convert("RGB")
        qr_img = qr_img.resize((340, 340))
        self._qr_img = qr_img
        self._photo = ImageTk.PhotoImage(qr_img)

        ttk.Label(self, text=f"QR Code — {recipe['name']}",
                  font=("Segoe UI", 12, "bold"), wraplength=380, justify="center").pack(pady=10)
        ttk.Label(self, image=self._photo).pack(pady=5)
        ttk.Label(
            self,
            text="Scannez avec l'appareil photo ou une application de\n"
                 "lecture de QR code pour voir le nom et les ingrédients.",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center"
        ).pack(pady=5)

        ttk.Button(self, text="💾 Enregistrer en image (PNG)",
                   command=self.save_image).pack(pady=10)

        if len(text) >= self.MAX_CHARS:
            ttk.Label(
                self,
                text="⚠️ La recette est longue : le QR code contient un\n"
                     "résumé tronqué (nom + ingrédients uniquement).",
                font=("Segoe UI", 8), foreground=COLOR_ERROR, justify="center"
            ).pack(pady=(0, 10))

    @classmethod
    def _build_text(cls, recipe, persons):
        lines = [recipe["name"], "", f"Ingrédients ({persons} pers.) :"]
        for ing in recipe["ingredients"]:
            qty = round(ing["quantity"] * persons, 2)
            if qty == int(qty):
                qty = int(qty)
            unit = f" {ing['unit']}" if ing["unit"] else ""
            lines.append(f"- {ing['name'].capitalize()} : {qty}{unit}")
        text = "\n".join(lines)
        if len(text) > cls.MAX_CHARS:
            text = text[: cls.MAX_CHARS - 3] + "..."
        return text

    def save_image(self):
        safe_name = re.sub(r'[\\/:*?"<>|]', "_", self.recipe["name"])
        path = filedialog.asksaveasfilename(
            title="Enregistrer le QR code",
            defaultextension=".png",
            filetypes=[("Image PNG", "*.png")],
            initialfile=f"qrcode_{safe_name}.png"
        )
        if not path:
            return
        try:
            self._qr_img.save(path)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'enregistrement a échoué :\n{e}")
            return
        messagebox.showinfo("Enregistré", f"QR code enregistré :\n{path}")


class WhatCanICookWindow(tk.Toplevel):
    """Fenêtre pour indiquer les ingrédients qu'on a sous la main, et voir
    quelles recettes sont réalisables (ou presque)."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Que puis-je cuisiner ?")
        self.geometry("640x660")
        self.grab_set()
        # Pré-coche les ingrédients de base qu'on a presque toujours sous la
        # main (correspondance exacte, insensible à la casse, avec la liste
        # d'ingrédients de l'utilisateur — un ingrédient de base absent de sa
        # liste est simplement ignoré ici).
        available_lower = {n.lower(): n for n in self.app.ingredient_names}
        self.have_names = [
            available_lower[staple.lower()] for staple in PANTRY_STAPLES
            if staple.lower() in available_lower
        ]

        ttk.Label(self, text="Indiquez les ingrédients que vous avez chez vous :",
                  font=("Segoe UI", 11, "bold")).pack(pady=(10, 2))
        ttk.Label(
            self,
            text="Quelques ingrédients de base courants sont déjà cochés ci-contre\n"
                 "(sel, huile, farine...) — retirez ceux que vous n'avez pas.",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center"
        ).pack(pady=(0, 5))

        columns = ttk.Frame(self)
        columns.pack(fill="both", expand=False, padx=15, pady=5)

        left = ttk.Frame(columns)
        left.pack(side="left", fill="both", expand=True, padx=(0, 10))
        ttk.Label(left, text="Tous les ingrédients :").pack()
        search_frame = ttk.Frame(left)
        search_frame.pack(fill="x", pady=3)
        ttk.Label(search_frame, text="🔍").pack(side="left")
        self.search_entry = ttk.Entry(search_frame)
        self.search_entry.pack(side="left", fill="x", expand=True, padx=3)
        self.search_entry.bind("<KeyRelease>", lambda e: self._populate_all())
        self.all_listbox = tk.Listbox(left, height=14)
        self.all_listbox.pack(fill="both", expand=True)
        self.all_listbox.bind("<Double-Button-1>", lambda e: self._add_selected())
        ttk.Button(left, text="➕ Ajouter →", command=self._add_selected).pack(pady=5)

        right = ttk.Frame(columns)
        right.pack(side="left", fill="both", expand=True)
        ttk.Label(right, text="Ce que j'ai :").pack()
        self.have_listbox = tk.Listbox(right, height=16)
        self.have_listbox.pack(fill="both", expand=True, pady=(3, 0))
        self.have_listbox.bind("<Double-Button-1>", lambda e: self._remove_selected())
        ttk.Button(right, text="🗑 Retirer", command=self._remove_selected).pack(pady=5)

        self._populate_all()
        self._populate_have()

        ttk.Button(self, text="🔍 Voir les recettes réalisables",
                   command=self.compute_feasible).pack(pady=8)

        self.result_text = tk.Text(self, height=12, width=70, wrap="word")
        self.result_text.pack(pady=5, padx=15, fill="both", expand=True)

    def _populate_all(self):
        search = self.search_entry.get().strip()
        search_key = ingredient_sort_key(search) if search else ""
        self.all_listbox.delete(0, tk.END)
        for name in self.app.ingredient_names:
            if name in self.have_names:
                continue
            if search_key and search_key not in ingredient_sort_key(name):
                continue
            self.all_listbox.insert(tk.END, name)

    def _populate_have(self):
        self.have_listbox.delete(0, tk.END)
        for name in self.have_names:
            self.have_listbox.insert(tk.END, name)

    def _add_selected(self):
        sel = self.all_listbox.curselection()
        for i in sel:
            name = self.all_listbox.get(i)
            if name not in self.have_names:
                self.have_names.append(name)
        self._populate_have()
        self._populate_all()

    def _remove_selected(self):
        sel = self.have_listbox.curselection()
        names_to_remove = {self.have_listbox.get(i) for i in sel}
        self.have_names = [n for n in self.have_names if n not in names_to_remove]
        self._populate_have()
        self._populate_all()

    def compute_feasible(self):
        have_keys = {ingredient_sort_key(n) for n in self.have_names}
        if not have_keys:
            messagebox.showinfo("Info", "Ajoutez au moins un ingrédient que vous avez.")
            return

        results = []
        for recipe in self.app.recipes:
            seen = set()
            missing = []
            for ing in recipe["ingredients"]:
                key = ingredient_sort_key(ing["name"])
                if key not in have_keys and key not in seen:
                    seen.add(key)
                    missing.append(ing["name"].capitalize())
            results.append((recipe, len(missing), missing))

        results.sort(key=lambda t: (t[1], ingredient_sort_key(t[0]["name"])))

        self.result_text.delete("1.0", tk.END)
        feasible = [r for r in results if r[1] == 0]
        almost = [r for r in results if 0 < r[1] <= 3]

        if feasible:
            self.result_text.insert(tk.END, "✅ Réalisables avec ce que vous avez :\n\n")
            for recipe, missing_count, missing in feasible:
                star = "⭐ " if recipe.get("favorite") else ""
                self.result_text.insert(tk.END, f"- {star}{recipe['name']}\n")
        else:
            self.result_text.insert(tk.END, "Aucune recette n'est réalisable à 100 % avec ces ingrédients.\n")

        if almost:
            self.result_text.insert(tk.END, "\n🟡 Presque (il manque 1 à 3 ingrédients) :\n\n")
            for recipe, missing_count, missing in almost:
                self.result_text.insert(tk.END, f"- {recipe['name']} (manque : {', '.join(missing)})\n")

        if not feasible and not almost:
            self.result_text.insert(tk.END, "\nEssayez d'ajouter d'autres ingrédients à votre sélection.\n")


class WeeklyPlanWindow(tk.Toplevel):
    """Planning des repas de la semaine (petit-déjeuner, déjeuner en 3 temps,
    dîner en 3 temps), avec génération automatique de la liste de courses
    pour l'ensemble des recettes planifiées."""

    MEAL_SLOTS = [
        "Petit-déjeuner",
        "Déjeuner — Entrée", "Déjeuner — Plat", "Déjeuner — Dessert",
        "Dîner — Entrée", "Dîner — Plat", "Dîner — Dessert",
    ]

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Planning de la semaine")
        self.geometry("1080x680")
        self.minsize(600, 400)
        self.resizable(True, True)
        self.grab_set()

        self.plan = load_weekly_plan()  # {jour: {créneau: {'recipe_name':.., 'persons':..}}}
        recipe_names = [r["name"] for r in self.app.recipes]

        ttk.Label(self, text="Planning de la semaine", font=("Segoe UI", 14, "bold")).pack(pady=10)
        ttk.Label(self, text="Vue calendrier : jours en colonnes, repas en lignes.",
                  font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED).pack()

        grid_container = ttk.Frame(self)
        grid_container.pack(fill="both", expand=True, padx=10, pady=10)
        h_scrollbar = ttk.Scrollbar(grid_container, orient="horizontal")
        v_scrollbar = ttk.Scrollbar(grid_container, orient="vertical")
        canvas = tk.Canvas(grid_container, highlightthickness=0,
                            xscrollcommand=h_scrollbar.set, yscrollcommand=v_scrollbar.set)
        h_scrollbar.config(command=canvas.xview)
        v_scrollbar.config(command=canvas.yview)
        v_scrollbar.pack(side="right", fill="y")
        h_scrollbar.pack(side="bottom", fill="x")
        canvas.pack(side="left", fill="both", expand=True)

        calendar_frame = ttk.Frame(canvas)
        calendar_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=calendar_frame, anchor="nw")

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        canvas.bind("<Enter>", lambda e: canvas.bind_all("<MouseWheel>", _on_mousewheel))
        canvas.bind("<Leave>", lambda e: canvas.unbind_all("<MouseWheel>"))

        # En-têtes de colonnes : un jour de la semaine par colonne.
        ttk.Label(calendar_frame, text="", width=17).grid(row=0, column=0, padx=2, pady=2)
        for col, day in enumerate(WEEKDAYS, start=1):
            ttk.Label(calendar_frame, text=day, font=("Segoe UI", 9, "bold"),
                      foreground=COLOR_ACCENT_DARK, anchor="center").grid(
                row=0, column=col, padx=3, pady=(2, 6), sticky="ew")

        self.widgets = {}  # (jour, créneau) -> (combo, pers_entry)
        for row_index, slot in enumerate(self.MEAL_SLOTS, start=1):
            ttk.Label(calendar_frame, text=slot, font=("Segoe UI", 9), anchor="w",
                      width=17, wraplength=120, justify="left").grid(
                row=row_index, column=0, padx=(2, 6), pady=4, sticky="w")
            for col, day in enumerate(WEEKDAYS, start=1):
                cell = tk.Frame(calendar_frame, background=COLOR_CARD, highlightbackground=COLOR_BORDER,
                                 highlightthickness=1)
                cell.grid(row=row_index, column=col, padx=2, pady=2, sticky="nsew")
                day_data = self.plan.get(day) or {}
                slot_data = day_data.get(slot) or {}
                combo = ttk.Combobox(cell, values=["-- Aucune --"] + recipe_names,
                                      state="readonly", width=13)
                combo.set(slot_data.get("recipe_name") or "-- Aucune --")
                combo.pack(padx=3, pady=(3, 1))
                pers_frame = ttk.Frame(cell, style="Card.TFrame")
                pers_frame.pack(padx=3, pady=(0, 3))
                ttk.Label(pers_frame, text="👤", style="Card.TLabel").pack(side="left")
                pers_entry = ttk.Entry(pers_frame, width=3)
                pers_entry.insert(0, str(slot_data.get("persons", 4)))
                pers_entry.pack(side="left")
                self.widgets[(day, slot)] = (combo, pers_entry)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=8)
        ttk.Button(btn_frame, text="💾 Enregistrer le planning",
                   command=self.save_plan).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="🗑 Tout effacer",
                   command=self.clear_plan).grid(row=0, column=1, padx=5)
        ttk.Button(btn_frame, text="📆 Exporter vers un calendrier (.ics)",
                   command=self.export_ics).grid(row=1, column=0, columnspan=2, pady=(5, 0))

        export_frame = ttk.Frame(self)
        export_frame.pack(pady=5)
        ttk.Button(export_frame, text="Calculer la liste de courses de la semaine",
                   command=self.compute).grid(row=0, column=0, columnspan=4, padx=5, pady=3)
        ttk.Button(export_frame, text="📝 Texte", command=self.export_txt).grid(row=1, column=0, padx=5)
        ttk.Button(export_frame, text="📊 Excel", command=self.export_excel).grid(row=1, column=1, padx=5)
        ttk.Button(export_frame, text="📄 PDF", command=self.export_pdf).grid(row=1, column=2, padx=5)
        ttk.Button(export_frame, text="🖨️ Imprimer", command=self.print_list).grid(row=1, column=3, padx=5)
        ttk.Button(export_frame, text="☑️ Mode courses",
                   command=self.open_checklist).grid(row=2, column=0, columnspan=4, pady=(5, 0))

        self.result_text = tk.Text(self, height=8, width=64)
        self.result_text.pack(pady=10, padx=15, fill="both", expand=True)

    def _collect_selection(self):
        new_plan = {}
        pairs = []
        for day in WEEKDAYS:
            for slot in self.MEAL_SLOTS:
                combo, pers_entry = self.widgets[(day, slot)]
                name = combo.get()
                if not name or name == "-- Aucune --":
                    continue
                try:
                    persons = float(pers_entry.get().strip().replace(",", "."))
                except ValueError:
                    messagebox.showerror("Erreur", f"Nombre de personnes invalide pour {day} — {slot}.")
                    return None, None
                recipe = find_recipe_by_name(self.app.recipes, name)
                if recipe is None:
                    continue
                new_plan.setdefault(day, {})[slot] = {"recipe_name": name, "persons": persons}
                pairs.append((recipe, persons))
        return new_plan, pairs

    def save_plan(self):
        new_plan, pairs = self._collect_selection()
        if new_plan is None:
            return
        save_weekly_plan(new_plan)
        self.plan = new_plan
        messagebox.showinfo("Enregistré", "Le planning de la semaine a été enregistré.")

    def clear_plan(self):
        if not messagebox.askyesno("Confirmer", "Effacer tout le planning de la semaine ?"):
            return
        for (day, slot), (combo, pers_entry) in self.widgets.items():
            combo.set("-- Aucune --")
        save_weekly_plan({})
        self.plan = {}

    def export_ics(self):
        new_plan, pairs = self._collect_selection()
        if new_plan is None:
            return
        if not new_plan:
            messagebox.showinfo("Info", "Assignez au moins une recette à un créneau de la semaine.")
            return
        path = filedialog.asksaveasfilename(
            title="Exporter le planning vers un calendrier",
            defaultextension=".ics",
            filetypes=[("Fichier calendrier (.ics)", "*.ics")],
            initialfile="planning_repas.ics"
        )
        if not path:
            return
        try:
            content = build_weekly_plan_ics(new_plan)
            with open(path, "w", encoding="utf-8") as f:
                f.write(content)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo(
            "Export réussi",
            f"Planning exporté :\n{path}\n\n"
            "Importez ce fichier dans Google Agenda, Outlook ou Calendrier "
            "pour voir vos repas s'y répéter chaque semaine."
        )

    def compute(self):
        new_plan, pairs = self._collect_selection()
        if new_plan is None:
            return None
        if not pairs:
            messagebox.showinfo("Info", "Assignez au moins une recette à un créneau de la semaine.")
            return None

        grouped_totals = compute_grouped_totals(pairs)
        day_order = {d: i for i, d in enumerate(WEEKDAYS)}
        slot_order = {s: i for i, s in enumerate(self.MEAL_SLOTS)}
        chosen_recipes = []
        for day, slots in new_plan.items():
            for slot, info in slots.items():
                chosen_recipes.append((f"{day} — {slot} : {info['recipe_name']}", info["persons"]))
        chosen_recipes.sort(
            key=lambda t: (day_order.get(t[0].split(" — ")[0], 99),
                            slot_order.get(t[0].split(" — ")[1].split(" : ")[0], 99))
        )

        self.result_text.delete("1.0", tk.END)
        self.result_text.insert(tk.END, "=== Liste de courses de la semaine ===\n")
        for rayon, items in grouped_totals:
            self.result_text.insert(tk.END, f"\n--- {rayon} ---\n")
            for name, qty, unit in items:
                unit_display = f" {unit}" if unit else ""
                self.result_text.insert(tk.END, f"- {name} : {qty}{unit_display}\n")

        return chosen_recipes, grouped_totals

    def export_txt(self):
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses", defaultextension=".txt",
            filetypes=[("Fichier texte", "*.txt")], initialfile="liste_de_courses_semaine.txt"
        )
        if not path:
            return
        try:
            write_shopping_list_txt(path, "Liste de courses de la semaine", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste enregistrée :\n{path}")

    def export_excel(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Module manquant", "L'export Excel nécessite : pip install openpyxl")
            return
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses", defaultextension=".xlsx",
            filetypes=[("Fichier Excel", "*.xlsx")], initialfile="liste_de_courses_semaine.xlsx"
        )
        if not path:
            return
        try:
            wb = build_shopping_list_workbook(chosen_recipes, grouped_totals)
            wb.save(path)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste enregistrée :\n{path}")

    def export_pdf(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Module manquant", "L'export PDF nécessite : pip install reportlab")
            return
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses", defaultextension=".pdf",
            filetypes=[("Fichier PDF", "*.pdf")], initialfile="liste_de_courses_semaine.pdf"
        )
        if not path:
            return
        try:
            build_shopping_list_pdf(path, "Liste de courses de la semaine", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste enregistrée :\n{path}")

    def print_list(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Module manquant", "L'impression nécessite : pip install reportlab")
            return
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        temp_path = get_temp_pdf_path("planning_semaine")
        try:
            build_shopping_list_pdf(temp_path, "Liste de courses de la semaine", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"La préparation de l'impression a échoué :\n{e}")
            return
        result_status = print_file(temp_path)
        report_print_result(result_status, temp_path, "la liste de courses de la semaine")

    def open_checklist(self):
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        ShoppingChecklistWindow(self.app, grouped_totals, title="Liste de courses de la semaine")


class MenuManagerWindow(tk.Toplevel):
    """Gestion des menus (combinaisons de plusieurs recettes) : création,
    édition, suppression."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Mes menus")
        self.geometry("420x480")
        self.grab_set()

        ttk.Label(self, text="Mes menus enregistrés :", font=("Segoe UI", 11, "bold")).pack(pady=(10, 5))

        self.listbox = tk.Listbox(self, width=40, height=14)
        self.listbox.pack(pady=5, padx=15, fill="both", expand=True)
        self.menus = []
        self._populate()

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=10)
        ttk.Button(btn_frame, text="➕ Nouveau menu", command=self.new_menu).grid(row=0, column=0, padx=5)
        ttk.Button(btn_frame, text="👁 Ouvrir", command=self.open_menu).grid(row=0, column=1, padx=5)
        ttk.Button(btn_frame, text="🗑 Supprimer", command=self.delete_menu).grid(row=0, column=2, padx=5)

    def _populate(self):
        self.listbox.delete(0, tk.END)
        self.menus = load_menus()
        for menu in self.menus:
            self.listbox.insert(tk.END, f"{menu['name']} ({len(menu.get('items', []))} recette(s))")

    def _selected_index(self):
        sel = self.listbox.curselection()
        if not sel:
            messagebox.showinfo("Info", "Sélectionnez un menu dans la liste.")
            return None
        return sel[0]

    def new_menu(self):
        MenuFormWindow(self.app, self, menu_index=None)

    def open_menu(self):
        idx = self._selected_index()
        if idx is None:
            return
        MenuFormWindow(self.app, self, menu_index=idx)

    def delete_menu(self):
        idx = self._selected_index()
        if idx is None:
            return
        menus = load_menus()
        name = menus[idx]["name"]
        if not messagebox.askyesno("Confirmer", f"Supprimer le menu « {name} » ?"):
            return
        menus.pop(idx)
        save_menus(menus)
        self._populate()


class MenuFormWindow(tk.Toplevel):
    """Fenêtre de création/édition d'un menu : nom + liste de recettes avec
    nombre de personnes, export/impression de sa liste de courses."""

    CATEGORY_ORDER = {"Apéro": 0, "Entrée": 1, "Plat": 2, "Sauce": 3,
                       "Dessert": 4, "Boisson": 5, "Autre": 6}

    def __init__(self, app, manager, menu_index=None):
        super().__init__(app)
        self.app = app
        self.manager = manager
        self.menu_index = menu_index
        self.editing = menu_index is not None
        menus = load_menus()
        self.existing_menu = menus[menu_index] if self.editing else None

        self.title("Modifier le menu" if self.editing else "Nouveau menu")
        self.geometry("540x680")
        self.grab_set()

        ttk.Label(self, text="Nom du menu :", font=("Segoe UI", 11, "bold")).pack(pady=(10, 5))
        self.name_entry = ttk.Entry(self, width=40)
        self.name_entry.pack()
        if self.editing:
            self.name_entry.insert(0, self.existing_menu["name"])

        self.items = [dict(it) for it in self.existing_menu.get("items", [])] if self.editing else []

        ttk.Label(self, text="Ajouter une recette au menu :", font=("Segoe UI", 10, "bold")).pack(pady=(15, 5))
        add_frame = ttk.Frame(self)
        add_frame.pack(padx=15, fill="x")
        recipe_names = [r["name"] for r in self.app.recipes]
        self.recipe_combo = ttk.Combobox(add_frame, values=recipe_names, state="readonly", width=26)
        self.recipe_combo.pack(side="left", padx=(0, 5))
        if recipe_names:
            self.recipe_combo.current(0)
        ttk.Label(add_frame, text="pers. :").pack(side="left")
        self.add_persons_entry = ttk.Entry(add_frame, width=5)
        self.add_persons_entry.insert(0, "4")
        self.add_persons_entry.pack(side="left", padx=5)
        ttk.Button(add_frame, text="+ Ajouter", command=self.add_item).pack(side="left", padx=5)

        ttk.Label(self, text="Recettes du menu :", font=("Segoe UI", 10, "bold")).pack(pady=(15, 5))
        self.items_listbox = tk.Listbox(self, width=55, height=7)
        self.items_listbox.pack(padx=15, fill="x")
        self._refresh_items_listbox()
        ttk.Button(self, text="🗑 Retirer du menu", command=self.remove_item).pack(pady=5)

        ttk.Button(self, text="💾 Enregistrer le menu", command=self.save_menu).pack(pady=8)

        export_frame = ttk.Frame(self)
        export_frame.pack(pady=5)
        ttk.Button(export_frame, text="Calculer la liste de courses du menu",
                   command=self.compute).grid(row=0, column=0, columnspan=4, padx=5, pady=3)
        ttk.Button(export_frame, text="📝 Texte", command=self.export_txt).grid(row=1, column=0, padx=5)
        ttk.Button(export_frame, text="📊 Excel", command=self.export_excel).grid(row=1, column=1, padx=5)
        ttk.Button(export_frame, text="📄 PDF", command=self.export_pdf).grid(row=1, column=2, padx=5)
        ttk.Button(export_frame, text="🖨️ Imprimer", command=self.print_list).grid(row=1, column=3, padx=5)
        ttk.Button(export_frame, text="☑️ Mode courses",
                   command=self.open_checklist).grid(row=2, column=0, columnspan=4, pady=(5, 0))

        self.result_text = tk.Text(self, height=8, width=60)
        self.result_text.pack(pady=10, padx=15, fill="both", expand=True)

    def _refresh_items_listbox(self):
        self.items_listbox.delete(0, tk.END)
        for item in self.items:
            recipe = find_recipe_by_name(self.app.recipes, item["recipe_name"])
            cat = recipe.get("category", "Autre") if recipe else "?"
            self.items_listbox.insert(tk.END, f"[{cat}] {item['recipe_name']} ({item['persons']} pers.)")

    def add_item(self):
        name = self.recipe_combo.get()
        if not name:
            return
        try:
            persons = float(self.add_persons_entry.get().strip().replace(",", "."))
        except ValueError:
            messagebox.showerror("Erreur", "Nombre de personnes invalide.")
            return
        self.items.append({"recipe_name": name, "persons": persons})
        self._refresh_items_listbox()

    def remove_item(self):
        sel = self.items_listbox.curselection()
        if not sel:
            messagebox.showinfo("Info", "Sélectionnez une recette du menu à retirer.")
            return
        self.items.pop(sel[0])
        self._refresh_items_listbox()

    def save_menu(self):
        name = self.name_entry.get().strip()
        if not name:
            messagebox.showerror("Erreur", "Merci d'indiquer un nom de menu.")
            return
        if not self.items:
            messagebox.showerror("Erreur", "Ajoutez au moins une recette au menu.")
            return
        menus = load_menus()
        menu_data = {"name": name, "items": self.items}
        if self.editing:
            menus[self.menu_index] = menu_data
        else:
            menus.append(menu_data)
        save_menus(menus)
        self.manager._populate()
        messagebox.showinfo("Enregistré", f"Le menu « {name} » a été enregistré.")

    def _collect_pairs(self):
        pairs = []
        chosen_recipes = []
        sorted_items = sorted(
            self.items,
            key=lambda it: self.CATEGORY_ORDER.get(
                (find_recipe_by_name(self.app.recipes, it["recipe_name"]) or {}).get("category", "Autre"), 4
            )
        )
        for item in sorted_items:
            recipe = find_recipe_by_name(self.app.recipes, item["recipe_name"])
            if recipe is None:
                continue
            persons = item["persons"]
            pairs.append((recipe, persons))
            cat = recipe.get("category", "Autre")
            chosen_recipes.append((f"{cat} — {recipe['name']}", persons))
        return pairs, chosen_recipes

    def compute(self):
        pairs, chosen_recipes = self._collect_pairs()
        if not pairs:
            messagebox.showinfo("Info", "Ajoutez au moins une recette au menu.")
            return None
        grouped_totals = compute_grouped_totals(pairs)
        self.result_text.delete("1.0", tk.END)
        self.result_text.insert(tk.END, "=== Liste de courses du menu ===\n")
        for rayon, items in grouped_totals:
            self.result_text.insert(tk.END, f"\n--- {rayon} ---\n")
            for name, qty, unit in items:
                unit_display = f" {unit}" if unit else ""
                self.result_text.insert(tk.END, f"- {name} : {qty}{unit_display}\n")
        return chosen_recipes, grouped_totals

    def export_txt(self):
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        menu_name = self.name_entry.get().strip() or "menu"
        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses", defaultextension=".txt",
            filetypes=[("Fichier texte", "*.txt")], initialfile=f"{menu_name}.txt"
        )
        if not path:
            return
        try:
            write_shopping_list_txt(path, f"Menu : {menu_name}", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste enregistrée :\n{path}")

    def export_excel(self):
        if not OPENPYXL_AVAILABLE:
            messagebox.showerror("Module manquant", "L'export Excel nécessite : pip install openpyxl")
            return
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        menu_name = self.name_entry.get().strip() or "menu"
        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses", defaultextension=".xlsx",
            filetypes=[("Fichier Excel", "*.xlsx")], initialfile=f"{menu_name}.xlsx"
        )
        if not path:
            return
        try:
            wb = build_shopping_list_workbook(chosen_recipes, grouped_totals)
            wb.save(path)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste enregistrée :\n{path}")

    def export_pdf(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Module manquant", "L'export PDF nécessite : pip install reportlab")
            return
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        menu_name = self.name_entry.get().strip() or "menu"
        path = filedialog.asksaveasfilename(
            title="Enregistrer la liste de courses", defaultextension=".pdf",
            filetypes=[("Fichier PDF", "*.pdf")], initialfile=f"{menu_name}.pdf"
        )
        if not path:
            return
        try:
            build_shopping_list_pdf(path, f"Menu : {menu_name}", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Liste enregistrée :\n{path}")

    def print_list(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror("Module manquant", "L'impression nécessite : pip install reportlab")
            return
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        menu_name = self.name_entry.get().strip() or "menu"
        temp_path = get_temp_pdf_path("menu")
        try:
            build_shopping_list_pdf(temp_path, f"Menu : {menu_name}", chosen_recipes, grouped_totals)
        except Exception as e:
            messagebox.showerror("Erreur", f"La préparation de l'impression a échoué :\n{e}")
            return
        result_status = print_file(temp_path)
        report_print_result(result_status, temp_path, f"le menu « {menu_name} »")

    def open_checklist(self):
        result = self.compute()
        if result is None:
            return
        chosen_recipes, grouped_totals = result
        menu_name = self.name_entry.get().strip() or "menu"
        ShoppingChecklistWindow(self.app, grouped_totals, title=f"Menu : {menu_name}")


class ImportFromUrlWindow(tk.Toplevel):
    """Importe une recette à partir d'un lien internet (fonctionne avec les
    sites utilisant le format de données standard Schema.org Recipe)."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Importer une recette depuis un lien")
        self.geometry("520x280")
        self.resizable(False, False)
        self.grab_set()

        ttk.Label(self, text="🌐 Importer une recette depuis un lien",
                  font=("Segoe UI", 13, "bold")).pack(pady=15)
        ttk.Label(
            self,
            text="Collez l'adresse (URL) d'une page de recette. Cela fonctionne\n"
                 "avec la plupart des grands sites de cuisine (qui utilisent un\n"
                 "format de données standard). Une connexion internet est requise.",
            justify="center", font=("Segoe UI", 9)
        ).pack(pady=(0, 15))

        self.url_entry = ttk.Entry(self, width=55)
        self.url_entry.pack(pady=5, padx=20, fill="x")
        self.url_entry.bind("<Return>", lambda e: self.fetch())

        self.status_label = ttk.Label(self, text="", font=("Segoe UI", 9), foreground=COLOR_TEXT_MUTED)
        self.status_label.pack(pady=(5, 5))

        self.fetch_button = ttk.Button(self, text="🌐 Récupérer la recette", command=self.fetch)
        self.fetch_button.pack(pady=10)

        ttk.Label(
            self,
            text="Après import, vérifiez et complétez la recette si besoin\n"
                 "(le repérage des quantités et unités n'est pas toujours parfait).",
            justify="center", font=("Segoe UI", 8), foreground="#999"
        ).pack(pady=(5, 0))

    def fetch(self):
        url = self.url_entry.get().strip()
        if not url:
            messagebox.showinfo("Info", "Collez d'abord une adresse internet (URL).")
            return
        if not url.lower().startswith(("http://", "https://")):
            url = "https://" + url

        self.fetch_button.config(state="disabled")
        self.status_label.config(text="Récupération en cours...")
        self.update()

        try:
            recipe_data = fetch_recipe_from_url(url)
        except Exception as e:
            self.status_label.config(text="")
            self.fetch_button.config(state="normal")
            messagebox.showerror("Échec de l'import", str(e))
            return

        # Enregistre automatiquement tout ingrédient qui n'existe pas encore
        known_lower = {n.lower() for n in self.app.ingredient_names}
        ingredients_list = load_ingredients()
        changed = False
        for ing in recipe_data["ingredients"]:
            if ing["name"].lower() not in known_lower:
                ingredients_list.append(ing["name"])
                known_lower.add(ing["name"].lower())
                changed = True
        if changed:
            self.app.ingredient_names = save_ingredients(ingredients_list)

        self.status_label.config(text="")
        self.fetch_button.config(state="normal")
        self.destroy()
        RecipeFormWindow(self.app, recipe_index=None, prefill=recipe_data)


class ImportFromPhotoWindow(tk.Toplevel):
    """Importe une recette à partir d'une photo (recette manuscrite ou page
    d'un livre de cuisine), en extrayant le texte par reconnaissance optique
    de caractères (OCR). Contrairement à l'import depuis un lien, le texte
    extrait n'est pas automatiquement organisé en ingrédients/étapes — il
    est présenté à relire et corriger avant de créer la recette."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Importer une recette depuis une photo")
        self.geometry("540x640")
        self.minsize(460, 500)
        self.resizable(True, True)
        self.grab_set()
        self.photo_path = None
        self._preview_ref = None

        ttk.Label(self, text="📷 Importer une recette depuis une photo",
                  font=("Segoe UI", 13, "bold")).pack(pady=(15, 5))
        ttk.Label(
            self,
            text="Prenez en photo (ou scannez) une recette manuscrite ou une\n"
                 "page de livre de cuisine, puis choisissez l'image ici. Le texte\n"
                 "en est extrait automatiquement, mais reste à relire et organiser\n"
                 "vous-même (contrairement à l'import depuis un lien, une photo n'a\n"
                 "pas de structure ingrédients/étapes que l'on puisse deviner).",
            font=("Segoe UI", 8), foreground=COLOR_TEXT_MUTED, justify="center", wraplength=480
        ).pack(pady=(0, 10))

        if not PYTESSERACT_AVAILABLE:
            ttk.Label(
                self,
                text="⚠ Cette fonctionnalité nécessite le module 'pytesseract'\n"
                     "ET le programme Tesseract OCR installé séparément sur ce PC.\n"
                     "Voir le LISEZ-MOI pour les instructions d'installation.",
                foreground=COLOR_ERROR, font=("Segoe UI", 9, "bold"), justify="center"
            ).pack(pady=10)

        self.preview_label = ttk.Label(self, text="Aucune photo choisie", foreground=COLOR_TEXT_MUTED)
        self.preview_label.pack(pady=5)

        btn_frame = ttk.Frame(self)
        btn_frame.pack(pady=5)
        ttk.Button(btn_frame, text="📁 Choisir une photo",
                   command=self.choose_photo).grid(row=0, column=0, padx=5)
        self.extract_button = ttk.Button(btn_frame, text="🔍 Extraire le texte",
                                          command=self.extract_text, state="disabled")
        self.extract_button.grid(row=0, column=1, padx=5)

        ttk.Label(self, text="Texte extrait (modifiable) :", font=("Segoe UI", 9, "bold")).pack(
            pady=(10, 3))
        self.text_box = tk.Text(self, height=14, wrap="word")
        self.text_box.pack(fill="both", expand=True, padx=15, pady=(0, 10))

        ttk.Button(self, text="➡️ Créer la recette avec ce texte",
                   command=self.create_recipe).pack(pady=(0, 15))

    def choose_photo(self):
        path = filedialog.askopenfilename(
            title="Choisir une photo de recette",
            filetypes=[("Images", "*.jpg *.jpeg *.png *.webp *.bmp *.tiff")]
        )
        if not path:
            return
        self.photo_path = path
        self.preview_label.config(text=os.path.basename(path), foreground=COLOR_TEXT)

        if PIL_AVAILABLE:
            try:
                img = Image.open(path)
                img.thumbnail((300, 220))
                self._preview_ref = ImageTk.PhotoImage(img)
                self.preview_label.config(image=self._preview_ref, text="", compound="top")
            except Exception:
                pass

        if PYTESSERACT_AVAILABLE:
            self.extract_button.config(state="normal")

    def extract_text(self):
        if not self.photo_path:
            messagebox.showinfo("Info", "Choisissez d'abord une photo.")
            return
        if not PYTESSERACT_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "Cette fonctionnalité nécessite le module 'pytesseract'\n"
                "(pip install pytesseract) ET le programme Tesseract OCR\n"
                "installé séparément sur ce PC. Voir le LISEZ-MOI."
            )
            return
        try:
            image = Image.open(self.photo_path) if PIL_AVAILABLE else self.photo_path
            extracted = pytesseract.image_to_string(image, lang="fra")
        except Exception as e:
            messagebox.showerror(
                "Échec de l'extraction",
                "La reconnaissance de texte a échoué. Vérifiez que Tesseract OCR "
                f"est bien installé sur ce PC et accessible.\n\nDétail : {e}"
            )
            return
        extracted = extracted.strip()
        self.text_box.delete("1.0", tk.END)
        if extracted:
            self.text_box.insert("1.0", extracted)
        else:
            messagebox.showinfo(
                "Info",
                "Aucun texte n'a pu être extrait de cette photo. Essayez une image "
                "plus nette, mieux cadrée ou mieux éclairée."
            )

    def create_recipe(self):
        raw_text = self.text_box.get("1.0", "end-1c").strip()
        if not raw_text:
            if not messagebox.askyesno(
                "Aucun texte",
                "Aucun texte n'a été extrait ou saisi. Créer quand même une "
                "recette vide (avec juste la photo) ?"
            ):
                return

        images = []
        if self.photo_path:
            copied = copy_image_to_store(self.photo_path)
            if copied:
                images.append(copied)

        prefill = {
            "name": "",
            "description": raw_text[:2056],
            "ingredients": [],
            "prep_time": "",
            "cook_time": "",
            "default_persons": 4,
            "images": images,
        }
        self.destroy()
        RecipeFormWindow(self.app, recipe_index=None, prefill=prefill)


class CookbookExportWindow(tk.Toplevel):
    """Exporte plusieurs recettes réunies en un seul PDF façon livre de
    cuisine (page de sommaire puis une recette par page)."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Exporter le livre de recettes")
        self.geometry("560x600")
        self.grab_set()

        ttk.Label(self, text="📖 Exporter le livre de recettes",
                  font=("Segoe UI", 13, "bold")).pack(pady=(15, 5))
        ttk.Label(
            self,
            text="Sélectionnez les recettes à inclure dans un seul PDF,\n"
                 "façon livre de cuisine.",
            justify="center", font=("Segoe UI", 9)
        ).pack(pady=(0, 10))

        filter_frame = ttk.Frame(self)
        filter_frame.pack(pady=(0, 5), fill="x", padx=15)
        ttk.Label(filter_frame, text="Filtrer par catégorie :").pack(side="left")
        self.category_filter = ttk.Combobox(
            filter_frame, values=["Toutes"] + RecipeFormWindow.CATEGORY_OPTIONS,
            state="readonly", width=15
        )
        self.category_filter.set("Toutes")
        self.category_filter.pack(side="left", padx=5)
        self.category_filter.bind("<<ComboboxSelected>>", lambda e: self._populate())
        ttk.Button(filter_frame, text="Tout cocher", command=self.check_all).pack(side="left", padx=5)
        ttk.Button(filter_frame, text="Tout décocher", command=self.uncheck_all).pack(side="left")

        list_frame = ttk.Frame(self)
        list_frame.pack(padx=15, pady=5, fill="both", expand=True)
        canvas = tk.Canvas(list_frame, highlightthickness=0)
        scrollbar = ttk.Scrollbar(list_frame, orient="vertical", command=canvas.yview)
        self.rows_frame = ttk.Frame(canvas)
        self.rows_frame.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.rows_frame, anchor="nw")
        canvas.configure(yscrollcommand=scrollbar.set)
        canvas.pack(side="left", fill="both", expand=True)
        scrollbar.pack(side="right", fill="y")

        self.checks = []  # (var, recipe)
        self._populate()

        ttk.Button(self, text="📄 Générer le PDF du livre",
                   command=self.export_pdf).pack(pady=15)

    def _populate(self):
        for child in self.rows_frame.winfo_children():
            child.destroy()
        self.checks = []
        category = self.category_filter.get()
        for recipe in self.app.recipes:
            if category != "Toutes" and recipe.get("category", "Autre") != category:
                continue
            var = tk.BooleanVar(value=True)
            ttk.Checkbutton(self.rows_frame, text=format_recipe_list_label(recipe),
                             variable=var).pack(anchor="w", pady=2)
            self.checks.append((var, recipe))

    def check_all(self):
        for var, recipe in self.checks:
            var.set(True)

    def uncheck_all(self):
        for var, recipe in self.checks:
            var.set(False)

    def export_pdf(self):
        if not REPORTLAB_AVAILABLE:
            messagebox.showerror(
                "Module manquant",
                "L'export PDF nécessite le module 'reportlab'.\n"
                "Installez-le avec : pip install reportlab"
            )
            return
        selected = [(recipe, recipe.get("default_persons", 4) or 4)
                    for var, recipe in self.checks if var.get()]
        if not selected:
            messagebox.showinfo("Info", "Sélectionnez au moins une recette.")
            return

        path = filedialog.asksaveasfilename(
            title="Enregistrer le livre de recettes",
            defaultextension=".pdf",
            filetypes=[("Fichier PDF", "*.pdf")],
            initialfile="mon_livre_de_recettes.pdf"
        )
        if not path:
            return
        try:
            build_cookbook_pdf(path, selected)
        except Exception as e:
            messagebox.showerror("Erreur", f"L'export a échoué :\n{e}")
            return
        messagebox.showinfo("Export réussi", f"Livre de recettes enregistré :\n{path}")


class CompareRecipesWindow(tk.Toplevel):
    """Compare deux recettes côte à côte : temps, difficulté, note, et
    ingrédients communs / différents."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Comparer deux recettes")
        self.geometry("720x640")
        self.grab_set()

        recipe_names = [r["name"] for r in self.app.recipes]

        picker_frame = ttk.Frame(self)
        picker_frame.pack(pady=15, padx=15, fill="x")

        ttk.Label(picker_frame, text="Recette A :", font=("Segoe UI", 10, "bold")).grid(
            row=0, column=0, sticky="w", padx=(0, 5))
        self.combo_a = ttk.Combobox(picker_frame, values=recipe_names, state="readonly", width=32)
        self.combo_a.grid(row=0, column=1, padx=5)
        if recipe_names:
            self.combo_a.current(0)

        ttk.Label(picker_frame, text="Recette B :", font=("Segoe UI", 10, "bold")).grid(
            row=1, column=0, sticky="w", padx=(0, 5), pady=(8, 0))
        self.combo_b = ttk.Combobox(picker_frame, values=recipe_names, state="readonly", width=32)
        self.combo_b.grid(row=1, column=1, padx=5, pady=(8, 0))
        if len(recipe_names) > 1:
            self.combo_b.current(1)
        elif recipe_names:
            self.combo_b.current(0)

        ttk.Button(picker_frame, text="⚖️ Comparer", command=self.compare).grid(
            row=0, column=2, rowspan=2, padx=15)

        self.result_text = tk.Text(self, wrap="word")
        self.result_text.pack(fill="both", expand=True, padx=15, pady=(0, 15))

    def compare(self):
        name_a = self.combo_a.get()
        name_b = self.combo_b.get()
        if not name_a or not name_b:
            messagebox.showinfo("Info", "Choisissez une recette dans chaque liste.")
            return
        recipe_a = find_recipe_by_name(self.app.recipes, name_a)
        recipe_b = find_recipe_by_name(self.app.recipes, name_b)
        if recipe_a is None or recipe_b is None:
            return

        self.result_text.delete("1.0", tk.END)
        self.result_text.insert(tk.END, f"{'':20}{name_a:<30}{name_b}\n")
        self.result_text.insert(tk.END, "-" * 80 + "\n")

        def field_line(label, value_a, value_b):
            self.result_text.insert(tk.END, f"{label:<20}{str(value_a):<30}{str(value_b)}\n")

        cat_a = recipe_a.get("category", "Autre")
        cat_b = recipe_b.get("category", "Autre")
        field_line("Catégorie :", cat_a, cat_b)

        fav_a = "⭐ Oui" if recipe_a.get("favorite") else "Non"
        fav_b = "⭐ Oui" if recipe_b.get("favorite") else "Non"
        field_line("Favori :", fav_a, fav_b)

        field_line("Note :", rating_stars(recipe_a.get("rating", 0)), rating_stars(recipe_b.get("rating", 0)))
        field_line("Difficulté :", recipe_a.get("difficulty") or "—", recipe_b.get("difficulty") or "—")

        prep_a = recipe_a.get("prep_time") or "—"
        prep_b = recipe_b.get("prep_time") or "—"
        field_line("Préparation :", f"{prep_a} min" if prep_a != "—" else "—", f"{prep_b} min" if prep_b != "—" else "—")

        cook_a = recipe_a.get("cook_time") or "—"
        cook_b = recipe_b.get("cook_time") or "—"
        field_line("Cuisson :", f"{cook_a} min" if cook_a != "—" else "—", f"{cook_b} min" if cook_b != "—" else "—")

        def total_minutes(r):
            try:
                return float(r.get("prep_time") or 0) + float(r.get("cook_time") or 0)
            except (TypeError, ValueError):
                return 0
        total_a, total_b = total_minutes(recipe_a), total_minutes(recipe_b)
        field_line("Temps total :", f"{total_a:.0f} min" if total_a else "—", f"{total_b:.0f} min" if total_b else "—")

        field_line("Cuisinée :", f"{recipe_a.get('times_cooked', 0)} fois", f"{recipe_b.get('times_cooked', 0)} fois")

        persons_a = recipe_a.get("default_persons", 1) or 1
        persons_b = recipe_b.get("default_persons", 1) or 1
        cost_a, cost_known_a, cost_total_a = compute_recipe_cost(recipe_a, persons_a)
        cost_b, cost_known_b, cost_total_b = compute_recipe_cost(recipe_b, persons_b)
        cost_display_a = f"{cost_a:.2f} € ({persons_a} p.)" if cost_known_a else "—"
        cost_display_b = f"{cost_b:.2f} € ({persons_b} p.)" if cost_known_b else "—"
        field_line("Coût estimé :", cost_display_a, cost_display_b)

        nutri_a, nutri_known_a, nutri_total_a = compute_recipe_nutrition(recipe_a, persons_a)
        nutri_b, nutri_known_b, nutri_total_b = compute_recipe_nutrition(recipe_b, persons_b)
        kcal_display_a = f"{nutri_a['kcal']:.0f} kcal ({persons_a} p.)" if nutri_known_a else "—"
        kcal_display_b = f"{nutri_b['kcal']:.0f} kcal ({persons_b} p.)" if nutri_known_b else "—"
        field_line("Nutrition (kcal) :", kcal_display_a, kcal_display_b)

        ing_names_a = {ing["name"].strip().lower(): ing["name"] for ing in recipe_a["ingredients"]}
        ing_names_b = {ing["name"].strip().lower(): ing["name"] for ing in recipe_b["ingredients"]}
        field_line("Nb. ingrédients :", len(ing_names_a), len(ing_names_b))

        common_keys = set(ing_names_a) & set(ing_names_b)
        only_a_keys = set(ing_names_a) - set(ing_names_b)
        only_b_keys = set(ing_names_b) - set(ing_names_a)

        self.result_text.insert(tk.END, "\n" + "=" * 80 + "\n")
        self.result_text.insert(tk.END, f"\n🟰 Ingrédients communs ({len(common_keys)}) :\n")
        if common_keys:
            for key in sorted(common_keys, key=ingredient_sort_key):
                self.result_text.insert(tk.END, f"  - {ing_names_a[key].capitalize()}\n")
        else:
            self.result_text.insert(tk.END, "  Aucun ingrédient en commun.\n")

        self.result_text.insert(tk.END, f"\n🅰️ Uniquement dans « {name_a} » ({len(only_a_keys)}) :\n")
        if only_a_keys:
            for key in sorted(only_a_keys, key=ingredient_sort_key):
                self.result_text.insert(tk.END, f"  - {ing_names_a[key].capitalize()}\n")
        else:
            self.result_text.insert(tk.END, "  Aucun.\n")

        self.result_text.insert(tk.END, f"\n🅱️ Uniquement dans « {name_b} » ({len(only_b_keys)}) :\n")
        if only_b_keys:
            for key in sorted(only_b_keys, key=ingredient_sort_key):
                self.result_text.insert(tk.END, f"  - {ing_names_b[key].capitalize()}\n")
        else:
            self.result_text.insert(tk.END, "  Aucun.\n")


class StatisticsWindow(tk.Toplevel):
    """Fenêtre affichant des statistiques simples sur les recettes."""

    def __init__(self, app):
        super().__init__(app)
        self.app = app
        self.title("Statistiques")
        self.geometry("560x820")
        self.minsize(480, 500)
        self.resizable(True, True)
        self.grab_set()

        text_frame = ttk.Frame(self)
        text_frame.pack(fill="both", expand=True, padx=15, pady=(15, 5))
        text = tk.Text(text_frame, wrap="word", height=22)
        text_scrollbar = ttk.Scrollbar(text_frame, orient="vertical", command=text.yview)
        text.configure(yscrollcommand=text_scrollbar.set)
        text.pack(side="left", fill="both", expand=True)
        text_scrollbar.pack(side="right", fill="y")

        recipes = self.app.recipes
        total = len(recipes)
        text.insert(tk.END, "=== Statistiques ===\n\n")
        text.insert(tk.END, f"Nombre total de recettes : {total}\n\n")

        cat_counts = {}
        for r in recipes:
            cat = r.get("category", "Autre")
            cat_counts[cat] = cat_counts.get(cat, 0) + 1
        text.insert(tk.END, "Répartition par catégorie :\n")
        for cat in RecipeFormWindow.CATEGORY_OPTIONS:
            if cat in cat_counts:
                text.insert(tk.END, f"  - {cat} : {cat_counts[cat]}\n")
        text.insert(tk.END, "\n")

        diff_counts = {}
        for r in recipes:
            diff = r.get("difficulty") or "Non renseignée"
            diff_counts[diff] = diff_counts.get(diff, 0) + 1
        text.insert(tk.END, "Répartition par difficulté :\n")
        for diff in ["Facile", "Moyen", "Difficile", "Non renseignée"]:
            if diff in diff_counts:
                text.insert(tk.END, f"  - {diff} : {diff_counts[diff]}\n")
        text.insert(tk.END, "\n")

        fav_count = sum(1 for r in recipes if r.get("favorite"))
        text.insert(tk.END, f"Recettes favorites : {fav_count}\n\n")

        rated = [r.get("rating", 0) for r in recipes if r.get("rating")]
        if rated:
            avg_rating = sum(rated) / len(rated)
            text.insert(tk.END, f"Note moyenne (recettes notées) : {avg_rating:.1f} / 5 "
                                 f"({len(rated)} recette(s) notée(s))\n\n")
        else:
            text.insert(tk.END, "Note moyenne : aucune recette notée pour le moment.\n\n")

        best_rated = [r for r in recipes if r.get("rating", 0) == 5]
        if best_rated:
            text.insert(tk.END, "Recette(s) notée(s) 5 étoiles :\n")
            for r in best_rated[:10]:
                text.insert(tk.END, f"  - {r['name']}\n")
            text.insert(tk.END, "\n")

        cooked = sorted(
            (r for r in recipes if r.get("times_cooked", 0) > 0),
            key=lambda r: r.get("times_cooked", 0), reverse=True
        )
        text.insert(tk.END, "Recettes les plus cuisinées :\n")
        if cooked:
            for r in cooked[:10]:
                text.insert(tk.END, f"  - {r['name']} : {r['times_cooked']} fois\n")
        else:
            text.insert(tk.END, "  Aucune recette marquée comme cuisinée pour le moment.\n"
                                 "  (bouton « 🍳 J'ai cuisiné ça ! » dans « Voir une recette précise »)\n")
        text.insert(tk.END, "\n")

        tag_counts = {}
        for r in recipes:
            for tag in r.get("tags", []):
                tag_counts[tag] = tag_counts.get(tag, 0) + 1
        if tag_counts:
            text.insert(tk.END, "Étiquettes les plus utilisées :\n")
            for tag, count in sorted(tag_counts.items(), key=lambda t: -t[1])[:10]:
                text.insert(tk.END, f"  - {tag} : {count}\n")
            text.insert(tk.END, "\n")

        # ---- Recettes oubliées ----
        never_cooked = [r for r in recipes if r.get("times_cooked", 0) == 0]
        text.insert(tk.END, "🕸️ Recettes jamais cuisinées :\n")
        if never_cooked:
            shown = never_cooked[:15]
            for r in shown:
                text.insert(tk.END, f"  - {r['name']}\n")
            remaining = len(never_cooked) - len(shown)
            if remaining > 0:
                text.insert(tk.END, f"  ... et {remaining} autre(s)\n")
        else:
            text.insert(tk.END, "  Toutes vos recettes ont déjà été cuisinées au moins une fois. 👏\n")
        text.insert(tk.END, "\n")

        stale_cutoff_days = 90
        now = datetime.now()
        stale = []
        for r in recipes:
            cooked_dates = r.get("cooked_dates") or []
            if not cooked_dates:
                continue
            try:
                last_date = max(datetime.fromisoformat(d) for d in cooked_dates)
            except ValueError:
                continue
            if (now - last_date).days >= stale_cutoff_days:
                stale.append((r, (now - last_date).days))
        stale.sort(key=lambda t: -t[1])
        text.insert(tk.END, f"🕰️ Pas cuisinées depuis plus de {stale_cutoff_days} jours :\n")
        if stale:
            for r, days in stale[:15]:
                text.insert(tk.END, f"  - {r['name']} (il y a {days} jours)\n")
        else:
            text.insert(tk.END, "  Aucune recette dans ce cas pour le moment.\n")
        text.insert(tk.END, "\n")

        # ---- Coût moyen ----
        costs_per_person = []
        for r in recipes:
            persons = r.get("default_persons", 1) or 1
            cost, known, _ = compute_recipe_cost(r, persons)
            if known > 0:
                costs_per_person.append(cost / persons)
        text.insert(tk.END, "💰 Coût moyen par personne :\n")
        if costs_per_person:
            avg_cost = sum(costs_per_person) / len(costs_per_person)
            without_price = total - len(costs_per_person)
            text.insert(
                tk.END,
                f"  {avg_cost:.2f} € en moyenne, sur {len(costs_per_person)} recette(s) avec au "
                f"moins un prix connu ({without_price} sans prix renseigné)\n"
            )
        else:
            text.insert(tk.END, "  Aucune recette avec un prix renseigné pour le moment.\n"
                                 "  (voir « 💰 Gérer les prix » dans « Gérer les ingrédients »)\n")
        text.insert(tk.END, "\n")

        # ---- Calories moyennes ----
        kcal_per_person = []
        for r in recipes:
            persons = r.get("default_persons", 1) or 1
            nutrition, known, _ = compute_recipe_nutrition(r, persons)
            if known > 0:
                kcal_per_person.append(nutrition["kcal"] / persons)
        text.insert(tk.END, "🥗 Calories moyennes par personne :\n")
        if kcal_per_person:
            avg_kcal = sum(kcal_per_person) / len(kcal_per_person)
            text.insert(
                tk.END,
                f"  {avg_kcal:.0f} kcal en moyenne, sur {len(kcal_per_person)} recette(s) avec des "
                f"ingrédients reconnus dans la base nutritionnelle\n"
            )
        else:
            text.insert(tk.END, "  Aucune recette avec des ingrédients reconnus pour le moment.\n")

        text.config(state="disabled")

        # ---- Graphique d'évolution mensuelle des recettes cuisinées ----
        ttk.Label(self, text="📈 Recettes cuisinées par mois (12 derniers mois)",
                  font=("Segoe UI", 10, "bold")).pack(pady=(8, 2))
        chart_canvas = tk.Canvas(self, height=180, background=COLOR_CARD, highlightthickness=1,
                                  highlightbackground=COLOR_BORDER)
        chart_canvas.pack(fill="x", padx=15, pady=(0, 15))
        self.after(50, lambda: self._draw_monthly_chart(chart_canvas, recipes))

    def _draw_monthly_chart(self, canvas, recipes):
        """Dessine un histogramme simple (sans dépendance externe) du nombre
        de recettes cuisinées par mois, sur les 12 derniers mois."""
        now = datetime.now()
        months = []
        for i in range(11, -1, -1):
            year, month = now.year, now.month - i
            while month <= 0:
                month += 12
                year -= 1
            months.append((year, month))

        counts = {ym: 0 for ym in months}
        for r in recipes:
            for date_str in r.get("cooked_dates") or []:
                try:
                    d = datetime.fromisoformat(date_str)
                except (ValueError, TypeError):
                    continue
                ym = (d.year, d.month)
                if ym in counts:
                    counts[ym] += 1

        canvas.update_idletasks()
        width = max(canvas.winfo_width(), 480)
        height = 180
        margin_bottom, margin_top = 28, 14
        chart_height = height - margin_bottom - margin_top
        max_count = max(max(counts.values(), default=0), 1)
        bar_width = (width - 20) / len(months)

        month_labels_fr = ["jan", "fév", "mar", "avr", "mai", "jun",
                            "jul", "aoû", "sep", "oct", "nov", "déc"]

        canvas.delete("all")
        for i, ym in enumerate(months):
            count = counts[ym]
            bar_height = (count / max_count) * chart_height
            x0 = 10 + i * bar_width + 3
            x1 = 10 + (i + 1) * bar_width - 3
            y1 = height - margin_bottom
            y0 = y1 - bar_height
            canvas.create_rectangle(x0, y0, x1, y1, fill=COLOR_ACCENT, outline="")
            if count > 0:
                canvas.create_text((x0 + x1) / 2, y0 - 8, text=str(count),
                                    font=("Segoe UI", 8), fill=COLOR_TEXT)
            canvas.create_text((x0 + x1) / 2, y1 + 12, text=month_labels_fr[ym[1] - 1],
                                font=("Segoe UI", 7), fill=COLOR_TEXT_MUTED)


if __name__ == "__main__":
    app = App()
    app.mainloop()
